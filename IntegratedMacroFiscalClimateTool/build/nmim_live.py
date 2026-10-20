"""Green investment needs in the Integrated MF-C tool: the National Mitigation Investment Model
(NMIM v0.2) live, for the country picked on the Dashboard.

NMIM's own workbook (InvestmentNeeds/national_model/NMIM_v0.2.xlsx, built by build_workbook_v0_2.py)
computes one country in live formulas. Its engine sheets are copied here, renamed, with the country
taken from the tool's selector instead of NMIM's Settings sheet:

  GI_Inputs   NMIM's assumptions (yellow, named p_...): change them and everything recalculates,
              the fiscal outlook included
  GI_Model    NMIM's Mitigation sheet: 2021-2050 by fuel and sector, CPAT layout (K = 2021)
  GI_Data     NMIM's country data

and the Green investment sheet presents the country detail: investment, incremental investment by
sector, incremental cost under three lenses, who pays (P1-P5), the budget need, GHG and the power
system, with charts. The macro model reads the budget need and power CO2 from GI_Model (macro_live).
"""
import os, re
from copy import copy
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
import xlcharts

HERE = os.path.dirname(os.path.abspath(__file__))
NMIM_XLSX = os.path.join(HERE, "..", "InvestmentNeeds", "national_model", "NMIM_v0.2.xlsx")
SHEETS = {"Inputs": "GI_Inputs", "Mitigation": "GI_Model", "Data": "GI_Data"}
YC = lambda y: get_column_letter(11 + y - 2021)                  # GI_Model year column (K = 2021)

F = "Arial"; INK = "0E1B2A"; ACC = "C05621"; MUT = "5A6B7B"
HDR = PatternFill("solid", fgColor=INK); SEC = PatternFill("solid", fgColor="EAE3F2")
GRY = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="C8D0D6"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def font(size=10, bold=False, color=INK, italic=False):
    return Font(name=F, size=size, bold=bold, color=color, italic=italic)


def _rewrite(v):
    if not (isinstance(v, str) and v.startswith("=")):
        return v
    v = v.replace("Settings!$C$6", "gi_row").replace("Settings!$C$4", "gi_iso").replace("Settings!C5", "gi_name")
    return re.sub(r"(?<![A-Za-z0-9_'])(Inputs|Mitigation|Data)!", lambda m: SHEETS[m.group(1)] + "!", v)


def copy_engine(wb):
    """Copy NMIM's engine sheets and names into wb. Returns {code: row} for GI_Model."""
    src = openpyxl.load_workbook(NMIM_XLSX)
    for old, new in SHEETS.items():
        s, t = src[old], wb.create_sheet(new)
        for row in s.iter_rows():
            for c in row:
                if c.value is None and not c.has_style:
                    continue
                x = t.cell(c.row, c.column, _rewrite(c.value))
                if c.has_style:
                    x.font, x.fill, x.border = copy(c.font), copy(c.fill), copy(c.border)
                    x.number_format, x.alignment = c.number_format, copy(c.alignment)
        for k, d in s.column_dimensions.items():
            t.column_dimensions[k].width = d.width
            t.column_dimensions[k].hidden = d.hidden
        for m in s.merged_cells.ranges:
            t.merge_cells(str(m))
        t.freeze_panes = s.freeze_panes
        t.sheet_properties.tabColor = "BFBFBF"
    for n, d in src.defined_names.items():
        wb.defined_names[n] = DefinedName(n, attr_text=_rewrite("=" + d.attr_text)[1:])
    gm = wb["GI_Model"]
    gm["A1"] = "NMIM v0.2 engine (copied from InvestmentNeeds/national_model/NMIM_v0.2.xlsx): the country picked on the Dashboard"
    gi = wb["GI_Inputs"]
    gi.cell(1, 1).value = "NMIM v0.2 assumptions: change the yellow values; the Green investment sheet and the fiscal outlook recalculate"
    return {gm.cell(r, 8).value: r for r in range(1, gm.max_row + 1) if gm.cell(r, 8).value}


def ref(rows, code, y):
    return f"GI_Model!${YC(y)}${rows[code]}"


def scal(rows, key):
    return f"GI_Model!$J${rows['mit.sc.' + key]}"


def budget_paths(rows, y):
    """Formulas for the incremental budget need's capital and spending parts, share of NMIM GDP, as
    MacroFiscalModels/nmim_bridge.programme: spending = P1 and the public gap of P4 on non-CO2, land
    use and other energy CO2 (annual spending in NMIM); capital = the rest."""
    g = ref(rows, "mit.gdp.lvl", y)
    pub = f"{ref(rows, 'mit.dtyp.bud', y)}/{g}"
    nc, lu, oth = (ref(rows, c, y) for c in ("mit.spd.nc.2", "mit.spd.lu.2", "mit.inv.oth.2"))
    n1, n5, l1 = scal(rows, "ty_nc_p1"), scal(rows, "ty_nc_p5"), scal(rows, "ty_lu_p1")
    spend = (f"MIN({pub},({n1}*{nc}+{l1}*{lu}+p_ty_p4_gap*({nc}*(1-{n1}-{n5})+{lu}*(1-{l1})+{oth}))/{g})")
    return f"IFERROR({pub}-{spend},0)", f"IFERROR({spend},0)"


def power_co2(rows, y, sc):
    return "+".join(f"{ref(rows, f'mit.con.{f}.pow.{sc}', y)}*p_ef_{f}" for f in ("coa", "nga", "oil"))


# ---- the Green investment sheet ----------------------------------------------------------------
SECTORS = [("Power generation", "mit.dinv.gen"), ("Grids and storage", "mit.dinv.grid"), ("Transport", "mit.dinv.trn"),
           ("Buildings", "mit.dinv.bld"), ("Industry", "mit.dinv.ind"), ("Non-CO2 and land use", "mit.dinv.nco"),
           ("Other energy CO2", "mit.dinv.oth"), ("Unmodelled clean investment", "mit.dinv.cal")]
TYPES = [("P1 pure public", "mit.typ.p1.2"), ("P2 regulated utility", "mit.typ.p2.2"),
         ("P3 contracted revenue", "mit.typ.p3.2"), ("P4 viability gap", "mit.typ.p4.2"),
         ("P5 pure private", "mit.typ.p5.2")]


def add_sheet(wb, rows, last_dash):
    ws = wb.create_sheet("Green investment")
    ws.sheet_view.showGridLines = False
    YRS = [2025, 2030, 2035, 2040, 2045, 2050]
    cols = [get_column_letter(4 + j) for j in range(len(YRS))]            # D..I
    ws.column_dimensions["A"].width = 2; ws.column_dimensions["B"].width = 46; ws.column_dimensions["C"].width = 9
    for c in cols:
        ws.column_dimensions[c].width = 9
    ws["B2"] = "Green investment needs: what net zero takes, who pays, what it costs"; ws["B2"].font = font(16, True)
    ws["B3"] = ("The National Mitigation Investment Model (NMIM v0.2) in live formulas for the Dashboard's country: a baseline "
                "(current trends) and a net zero path, 2024-2050, built bottom-up from energy, power and emissions stocks; "
                "investment split by who repays it (P1 pure public ... P5 pure private). The incremental budget need is the "
                "climate programme in the fiscal outlook. Change NMIM's assumptions on GI_Inputs; the engine is GI_Model. "
                "Model output with NMIM's caveats (InvestmentNeeds/national_model/NMIM_Method_v0_2.md).")
    ws["B3"].font = font(9, italic=True, color=MUT); ws["B3"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("B3:I3"); ws.row_dimensions[3].height = 64
    ws["B5"] = "Country"; ws["B5"].font = font(10, True)
    ws["C5"] = "=sel_country"; ws["C5"].font = font(12, True, ACC); ws.merge_cells("C5:I5")
    # the selector for the engine
    ws["AN1"] = "gi_iso"; ws["AO1"] = f"=INDEX(Data_wide!$A$2:$A${last_dash},MATCH(sel_country,Data_wide!$B$2:$B${last_dash},0))"
    ws["AN2"] = "gi_row"; ws["AO2"] = "=MATCH(gi_iso,GI_Data!$A:$A,0)"
    ws["AN3"] = "gi_name"; ws["AO3"] = "=INDEX(GI_Data!$B:$B,gi_row)"
    for n, r in (("gi_iso", 1), ("gi_row", 2), ("gi_name", 3)):
        wb.defined_names[n] = DefinedName(n, attr_text=f"'Green investment'!$AO${r}")
    for c in ("AN", "AO"):
        ws.column_dimensions[c].hidden = True
    ws["B6"] = '=IF(ISNA(gi_row),"Not in NMIM: no investment needs for this country.","")'
    ws["B6"].font = font(10, True, ACC)

    r = 8
    def head(title, unit):
        nonlocal r
        for c in ["B", "C"] + cols:
            ws[f"{c}{r}"].fill = HDR
        ws[f"B{r}"] = title; ws[f"B{r}"].font = font(10, True, "FFFFFF")
        ws[f"C{r}"] = unit; ws[f"C{r}"].font = font(8, True, "FFFFFF")
        for c, y in zip(cols, YRS):
            x = ws[f"{c}{r}"]; x.value = y; x.font = font(9, True, "FFFFFF"); x.alignment = Alignment(horizontal="center")
        r += 1

    def line_(lab, fn, fmt="0.00", bold=False, indent=False):
        nonlocal r
        ws[f"B{r}"] = ("   " if indent else "") + lab; ws[f"B{r}"].font = font(9, bold); ws[f"B{r}"].border = BOX
        for c, y in zip(cols, YRS):
            x = ws[f"{c}{r}"]; x.value = f'=IF(ISNA(gi_row),"",IFERROR({fn(y)},""))'
            x.number_format = fmt; x.font = font(9, bold); x.border = BOX
            if bold:
                x.fill = GRY
        r += 1

    pg = lambda code: (lambda y: f"{ref(rows, code, y)}/{ref(rows, 'mit.gdp.lvl', y)}*100")
    head("Investment", "% GDP")
    line_("Total investment, baseline", pg("mit.inv.tot.1"))
    line_("Total investment, net zero", pg("mit.inv.tot.2"))
    line_("Incremental investment (net zero less baseline)", pg("mit.dinv.tot"), bold=True)
    for lab, code in SECTORS:
        line_(lab, pg(code), indent=True)
    r += 1
    head("Who pays (net zero investment)", "% GDP")
    for lab, code in TYPES:
        line_(lab, pg(code), indent=True)
    line_("Direct budget need, net zero (P1 + public gap of P4)", pg("mit.typ.bud.2"))
    line_("Direct budget need, incremental: the climate programme", pg("mit.dtyp.bud"), bold=True)
    line_("Private capital, net zero (P3 + P5 + private part of P4)", pg("mit.typ.prv.2"))
    r += 1
    head("Incremental cost (capital charge less savings)", "% GDP")
    line_("(a) market cost of capital", pg("mit.ic.a"))
    line_("(b) World Bank rate, 6%", pg("mit.ic.b"))
    line_("(c) 6% and local co-benefits (air pollution)", pg("mit.ic.c"))
    line_("Fuel savings (baseline less net zero)", pg("mit.dfuel"), indent=True)
    r += 1
    head("Emissions and energy", "")
    line_("GHG incl. land use, baseline (MtCO2e)", lambda y: ref(rows, "mit.ghg.tot.1", y), "0.0")
    line_("GHG incl. land use, net zero (MtCO2e)", lambda y: ref(rows, "mit.ghg.tot.2", y), "0.0")
    line_("Fossil power CO2, baseline (Mt)", lambda y: power_co2(rows, y, 1), "0.0")
    line_("Fossil power CO2, net zero (Mt)", lambda y: power_co2(rows, y, 2), "0.0")
    line_("Fossil share of generation, net zero (%)", lambda y: f"{ref(rows, 'mit.gen.fsh.2', y)}*100", "0")
    line_("Solar and wind capacity, net zero (GW)",
          lambda y: f"{ref(rows, 'mit.cap.sol.2', y)}+{ref(rows, 'mit.cap.wnd.2', y)}", "0.0")
    line_("EV share of car sales, net zero (%)", lambda y: f"{ref(rows, 'mit.ev.shs.2', y)}*100", "0")
    line_("GDP (2024 USD bn)", lambda y: ref(rows, "mit.gdp.lvl", y), "#,##0")
    r += 1
    ws[f"B{r}"] = ("Incremental cost: annual capital charge on the incremental capital in service less fuel and O&M savings "
                   "(lens c also less avoided air pollution); negative = a net saving. Budget need: what the budget spends "
                   "(P1, and the 30% public gap of P4); P2 rests on tariffs, a contingent liability where tariffs do not "
                   "recover costs.")
    ws[f"B{r}"].font = font(8, italic=True, color=MUT); ws[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(f"B{r}:I{r}"); ws.row_dimensions[r].height = 40

    # chart data, annual, hidden columns N..
    C0 = 14
    hdr = 8
    data = ([("year", None)] + [(lab, pg(code)) for lab, code in SECTORS] + [(lab, pg(code)) for lab, code in TYPES] +
            [("Baseline", lambda y: ref(rows, "mit.ghg.tot.1", y)), ("Net zero", lambda y: ref(rows, "mit.ghg.tot.2", y)),
             ("Incremental investment", pg("mit.dinv.tot")), ("Budget need (incremental)", pg("mit.dtyp.bud"))])
    DC = {}
    for j, (lab, fn) in enumerate(data):
        c = C0 + j; DC[lab] = c
        ws.cell(hdr, c, lab).font = font(8, True)
        for t, y in enumerate(range(2025, 2051)):
            v = y if fn is None and lab == "year" else (f"=IF(ISNA(gi_row),NA(),IFERROR({fn(y)},NA()))" if fn else None)
            ws.cell(hdr + 1 + t, c, v)
        ws.column_dimensions[get_column_letter(c)].hidden = True
    last = hdr + 26
    cat = (DC["year"], hdr + 1, last)

    lbl = C0 + len(data)                                    # labels: a year every five, for category axes
    for t, y in enumerate(range(2025, 2051)):
        ws.cell(hdr + 1 + t, lbl, str(y) if y % 5 == 0 else "")
    ws.column_dimensions[get_column_letter(lbl)].hidden = True
    pal = ["1F4E79", "4E8542", "C9A227", "C05621", "7A5195", "2E8B88", "8C8C8C", "A3C585"]
    for k, (title, labs) in enumerate([("Incremental investment by sector, % of GDP", [l for l, _ in SECTORS]),
                                       ("Net zero investment by who pays, % of GDP", [l for l, _ in TYPES])]):
        ws.add_chart(xlcharts.area(ws, (lbl, hdr + 1, last), [(DC[l], hdr, last, pal[j % 8]) for j, l in enumerate(labs)],
                                   title, height=7.5, width=14), f"K{5 + 16 * k}")
    ws.add_chart(xlcharts.line(ws, cat, [(DC["Baseline"], hdr, last, "baseline"), (DC["Net zero"], hdr, last, "green")],
                               "GHG incl. land use, MtCO2e", width=14), "K37")
    return ws, DC, hdr, last
