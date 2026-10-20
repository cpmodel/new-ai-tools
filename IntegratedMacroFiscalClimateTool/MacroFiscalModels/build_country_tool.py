"""Build the standalone single-country tool: an Excel workbook that needs no Python to use.

Pick a country and the workbook shows the whole integrated chain for it, with charts: the
investment need (NMIM), feasible carbon pricing (PriceFeasibility), the macro model's GDP and debt
paths (GIMF-lite / DIGNAD-lite / DIGNAR-lite), the LIC-DSF-lite rating, and which settings are data
and which are defaults. Everything is pre-computed by `integrated.py --all`; the workbook only looks
results up, so changing an assumption needs a re-run in Python.

    python3 integrated.py --all          # refresh results/integrated_*.csv (about 10 minutes)
    python3 build_country_tool.py        # -> MacroFiscal_CountryTool_v<VERSION>.xlsx (this folder)

Versioning: the current workbook sits in MacroFiscalModels/; bump VERSION when a changed workbook is
shipped, move the previous file to CountryTool/Archive/ and add a line to CountryTool/VersionNotes.md.
"""
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

VERSION = "0.3"
HERE = os.path.dirname(os.path.abspath(__file__))
RES = os.path.join(HERE, "results")
OUTDIR = HERE                                   # the current version sits in MacroFiscalModels/
DEFAULT_ISO = "KEN"

F = "Arial"; INK = "0E1B2A"; ACC = "C05621"; MUT = "5A6B7B"
HDR = PatternFill("solid", fgColor=INK); TINT = PatternFill("solid", fgColor="F0F3F5")
PICK = PatternFill("solid", fgColor="FFF2CC"); SEC = PatternFill("solid", fgColor="E3EAF2")
thin = Side(style="thin", color="C8D0D6"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
YEARS = list(range(2025, 2051))
DSF_KEYS = ["PV public debt / GDP", "PV ext debt / GDP", "Ext debt service / revenue"]
SETTINGS = [("income", "Income group", ""), ("s", "Public investment efficiency (share of spending that becomes capital)", "0.00"),
            ("Rg", "Return on public capital", "0.00"), ("ig", "Public investment before the programme, share of GDP", "0.000"),
            ("tax", "Tax revenue, share of GDP", "0.000"), ("debt", "Gross public debt, share of GDP", "0.000"),
            ("conc_share", "Concessional share of debt", "0.00"), ("constrained", "Credit-constrained households, share", "0.00"),
            ("rstar", "Real borrowing cost", "0.000"), ("g", "Government consumption, share of GDP (residual)", "0.000"),
            ("grants", "Grants, share of GDP", "0.000"), ("rents", "Fossil rents, share of GDP", "0.000")]


def font(size=10, bold=False, color=INK, italic=False):
    return Font(name=F, size=size, bold=bold, color=color, italic=italic)


def table(ws, df, widths=None):
    for j, c in enumerate(df.columns, 1):
        x = ws.cell(1, j, str(c)); x.font = font(9, True, "FFFFFF"); x.fill = HDR; x.border = BOX
    for i, row in enumerate(df.itertuples(index=False), 2):
        for j, v in enumerate(row, 1):
            if isinstance(v, float) and pd.isna(v):
                v = None
            ws.cell(i, j, v).font = font(9)
    for j, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def main():
    res = pd.read_csv(os.path.join(RES, "integrated_results.csv"))
    paths = pd.read_csv(os.path.join(RES, "integrated_paths.csv"))
    sets = pd.read_csv(os.path.join(RES, "integrated_settings.csv"))
    res = res[res.name.notna()].sort_values("name").reset_index(drop=True)
    wb = Workbook()

    # ---------------------------------------------------------------- data sheets
    wr = wb.active; wr.title = "Results"
    table(wr, res, [8, 24] + [14] * (res.shape[1] - 2))
    RCOL = {c: get_column_letter(j) for j, c in enumerate(res.columns, 1)}
    NR = res.shape[0] + 1

    wp = wb.create_sheet("Paths")
    paths.insert(0, "key", paths.iso3 + "|" + paths.series)
    table(wp, paths, [36, 6, 30] + [8] * len(YEARS))
    NP = paths.shape[0] + 1

    wsx = wb.create_sheet("Settings")
    num = pd.to_numeric(sets.value, errors="coerce")
    sets["value"] = [n if pd.notna(n) else v for n, v in zip(num, sets.value)]   # numbers as numbers, income as text
    sets.insert(0, "key", sets.iso3 + "|" + sets.setting)
    table(wsx, sets, [22, 6, 12, 12, 90])
    NS = sets.shape[0] + 1

    wc = wb.create_sheet("Countries")
    table(wc, res[["name", "iso3"]], [28, 8])

    # ---------------------------------------------------------------- Calc (chart series)
    wk = wb.create_sheet("Calc")
    wk["A1"] = "series"; wk["A1"].font = font(9, True)
    for j, y in enumerate(YEARS, 2):
        wk.cell(1, j, y).font = font(9, True)
    calc_rows = [("pub_cap", "Budget need: capital"), ("pub_spend", "Budget need: spending (no capital)"),
                 ("prv", "Private investment"), ("rev", "Carbon-pricing revenue"), ("sav", "Energy savings"),
                 ("gdp_dev_pct", "GDP vs today's path, %: programme"), ("debt_pct_gdp", "Public debt, % GDP: programme"),
                 ("gdp_dev_base_pct", "GDP vs today's path, %: baseline"), ("debt_base_pct_gdp", "Public debt, % GDP: baseline")]
    for k in DSF_KEYS:
        calc_rows += [(f"dsf_base|{k}", f"{k}: DSF baseline"), (f"dsf_prog|{k}", f"{k}: integrated scenario"),
                      (f"dsf_thr|{k}", f"{k}: threshold")]
    ROW = {}
    for i, (key, label) in enumerate(calc_rows, 2):
        ROW[key] = i
        wk.cell(i, 1, label).font = font(9)
        wk.cell(i, 30, key).font = font(8, color=MUT)                   # series key, column AD
        for j, y in enumerate(YEARS, 2):
            if key.startswith("dsf") and y > 2034:
                continue
            if key.startswith("dsf_thr"):
                src = "Paths!$D$2:$D$%d" % NP                            # threshold stored in the 2025 column
            else:
                src = "Paths!%s$2:%s$%d" % ((get_column_letter(j + 2),) * 2 + (NP,))
            wk.cell(i, j, f'=IFERROR(INDEX({src},MATCH(Country!$F$5&"|"&$AD{i},Paths!$A$2:$A${NP},0)),NA())').number_format = "0.00"
    # differences from the baseline (programme less baseline), computed from the rows above
    diffs = [("d_gap", "Financing gap: budget need less pricing revenue", [("pub_cap", 1), ("pub_spend", 1), ("rev", -1)], 2050),
             ("d_gdp", "GDP: programme less baseline, pp", [("gdp_dev_pct", 1), ("gdp_dev_base_pct", -1)], 2050),
             ("d_debt", "Public debt: programme less baseline, pp of GDP", [("debt_pct_gdp", 1), ("debt_base_pct_gdp", -1)], 2050),
             ("d_dsf", "PV of public debt / GDP: integrated scenario less DSF baseline, pp", [("dsf_prog|PV public debt / GDP", 1),
                                                                                ("dsf_base|PV public debt / GDP", -1)], 2034)]
    i = len(calc_rows) + 2
    for key, label, terms, yend in diffs:
        ROW[key] = i
        wk.cell(i, 1, label).font = font(9)
        for j, y in enumerate(YEARS, 2):
            if y > yend:
                continue
            L = get_column_letter(j)
            expr = "".join(("+" if sgn > 0 else "-") + f"{L}{ROW[k]}" for k, sgn in terms)
            wk.cell(i, j, f"=IFERROR({expr.lstrip('+')},NA())").number_format = "0.00"
        i += 1
    wk.sheet_properties.tabColor = "BFBFBF"

    # ---------------------------------------------------------------- Country (the tool)
    ws = wb.create_sheet("Country", 0)
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGH", [2, 58, 30, 10, 18, 18, 2, 2]):
        ws.column_dimensions[col].width = w
    ws["B2"] = "Macro-fiscal country tool"; ws["B2"].font = font(18, True)
    ws["B3"] = (f"Version {VERSION}. Investment needs (NMIM) -> feasible carbon pricing (PriceFeasibility) -> "
                "GIMF-lite / DIGNAD-lite / DIGNAR-lite -> LIC-DSF-lite. Illustrative calibration: a scenario, not a forecast.")
    ws["B3"].font = font(9, italic=True, color=MUT)
    ws["B5"] = "Country (pick from the list)"; ws["B5"].font = font(11, True)
    ws["C5"] = res.loc[res.iso3 == DEFAULT_ISO, "name"].iloc[0] if (res.iso3 == DEFAULT_ISO).any() else res.name.iloc[0]
    ws["C5"].fill = PICK; ws["C5"].font = font(11, True); ws["C5"].border = BOX
    ws.merge_cells("C5:D5")
    ws["E5"] = "ISO code"; ws["E5"].font = font(8, color=MUT); ws["E5"].alignment = Alignment(horizontal="right")
    ws["F5"] = f"=IFERROR(INDEX(Countries!$B$2:$B${NR},MATCH($C$5,Countries!$A$2:$A${NR},0)),\"\")"
    ws["F5"].font = font(11, True, ACC)
    dv = DataValidation(type="list", formula1=f"=Countries!$A$2:$A${NR}", allow_blank=False)
    ws.add_data_validation(dv); dv.add("C5")
    ISO = "$F$5"                                   # every lookup keys on the ISO code

    def R(field):
        return f'=IFERROR(INDEX(Results!${RCOL[field]}$2:${RCOL[field]}${NR},MATCH({ISO},Results!$A$2:$A${NR},0)),"")'

    r = 7

    def section(title, sub):
        nonlocal r
        for c in "BCDEF":
            ws[f"{c}{r}"].fill = SEC
        ws[f"B{r}"] = title; ws[f"B{r}"].font = font(11, True)
        ws[f"C{r}"] = sub; ws[f"C{r}"].font = font(8, italic=True, color=MUT)
        r += 1

    def line(label, field=None, fmt="0.00", unit="", formula=None):
        nonlocal r
        ws[f"B{r}"] = label; ws[f"B{r}"].font = font(10)
        c = ws[f"C{r}"]; c.value = formula or R(field); c.number_format = fmt; c.font = font(10, True)
        c.alignment = Alignment(horizontal="left" if fmt == "@" else "right")
        ws[f"D{r}"] = unit; ws[f"D{r}"].font = font(8, color=MUT)
        for col in "BCD":
            ws[f"{col}{r}"].border = BOX
        r += 1

    section("The country", "income group, model used and run status")
    line("Income group", "income", "@")
    line("Macro model used", "model", "@")
    line("Why this model", "model_reason", "@")
    line("Run status", "status", "@")
    line("Shared settings taken from defaults (not data)", "assumptions", "0", "of 16")
    r += 1
    section("1. How much: investment need", "NMIM v0.2, incremental to a current-trends baseline")
    line("Budget need 2030, total", "budget_need_2030_pct_gdp", "0.00", "% GDP")
    line("   of which capital (builds public capital)", "budget_capital_2030_pct_gdp", "0.00", "% GDP")
    line("   of which spending (land use, non-CO2: no capital)", "budget_spending_2030_pct_gdp", "0.00", "% GDP")
    line("Private investment 2030", "private_2030_pct_gdp", "0.00", "% GDP")
    line("Budget need, sum 2025-2050", "budget_need_cum_pct_gdp", "0.0", "% of a year's GDP")
    line("Energy savings 2050", "energy_saving_2050_pct_gdp", "0.00", "% GDP")
    r += 1
    section("2. How it can be paid: feasible carbon pricing", "PriceFeasibility, power sector, hybrid mix")
    line("Carbon tax rate", "pricing_tau_t", "0", "USD/t")
    line("Feebate rate", "pricing_tau_f", "0", "USD/t")
    line("Binding feasibility limit", "pricing_binding", "@", "F1 tariff, F2 bottom 40%, F3 generators")
    line("F1: first-year tariff rise", "pricing_F1", "0.0", "%")
    line("F2: net burden on bottom 40%", "pricing_F2", "0.00", "% consumption")
    line("F3: net payment by coal and oil generators", "pricing_F3", "0.0", "% revenue")
    line("Power CO2 abated by pricing", "pricing_abatement_pct", "0.0", "%")
    line("Revenue 2030 (earmarked for the budget need)", "pricing_revenue_2030_pct_gdp", "0.00", "% GDP")
    line("Revenue as share of the budget need, 2025-2050", "pricing_share_of_need_pct", "0", "%")
    line("Financing gap 2030 (need less revenue)", "financing_gap_2030_pct_gdp", "0.00", "% GDP")
    r += 1
    section("3. What it does: macro effects", "the chosen macro model, programme financed as above")
    line("GDP vs today's path, 2030", "gdp_dev_2030_pct", "+0.00;-0.00", "%")
    line("GDP vs today's path, 2050", "gdp_dev_2050_pct", "+0.00;-0.00", "%")
    line("Fossil rents (DIGNAR-lite assumes an 80% fall by 2050)", "rents_pct_gdp", "0.0", "% GDP")
    line("Public debt, 2024", "debt_2024_pct_gdp", "0", "% GDP")
    line("Public debt, 2030", "debt_2030_pct_gdp", "0", "% GDP")
    line("Public debt, 2050", "debt_2050_pct_gdp", "0", "% GDP")
    r += 1
    section("4. Is debt sustainable: LIC-DSF-lite", "low and lower-middle income only; mechanical rating")
    line("Risk of debt distress, baseline", "dsf_rating_base", "@")
    line("Risk of debt distress, with programme", "dsf_rating_programme", "@")
    line("Inputs taken from defaults", "dsf_assumptions", "@")
    r += 1
    section("Settings: data or default", "value used and where it came from")
    for key, label, fmt in SETTINGS:
        k = f'{ISO}&"|{key}"'
        line(label, fmt=fmt or "@",
             formula=f'=IFERROR(INDEX(Settings!$D$2:$D${NS},MATCH({k},Settings!$A$2:$A${NS},0)),"")')
        ws[f"D{r-1}"] = f'=IFERROR(INDEX(Settings!$E$2:$E${NS},MATCH({k},Settings!$A$2:$A${NS},0)),"")'
        ws[f"D{r-1}"].font = font(8, color=MUT)
    ws.column_dimensions["D"].width = 60
    last = r

    # charts
    def chart(title, keys, anchor, ytitle, x_end=2050):
        ch = LineChart()
        ch.title = title; ch.height = 7.5; ch.width = 16
        ch.varyColors = False                      # one colour per series (Excel otherwise colours each year)
        ch.y_axis.title = ytitle; ch.legend.position = "b"
        ncol = 1 + (x_end - 2024)
        for k in keys:
            ref = Reference(wk, min_col=1, max_col=ncol, min_row=ROW[k], max_row=ROW[k])
            ch.add_data(ref, titles_from_data=True, from_rows=True)
        ch.set_categories(Reference(wk, min_col=2, max_col=ncol, min_row=1, max_row=1))
        ch.x_axis.number_format = "0"
        ch.x_axis.delete = False; ch.y_axis.delete = False
        ws.add_chart(ch, anchor)

    # left column: the result; right column: the difference from the baseline (no programme)
    ws["H4"] = "Result"; ws["R4"] = "Difference from baseline (programme less no programme)"
    ws["H4"].font = font(11, True); ws["R4"].font = font(11, True)
    chart("Programme and financing, % of GDP", ["pub_cap", "pub_spend", "prv", "rev"], "H5", "% GDP")
    chart("Financing gap: budget need less pricing revenue, % of GDP", ["d_gap"], "R5", "% GDP")
    chart("GDP vs today's path, %", ["gdp_dev_pct", "gdp_dev_base_pct"], "H21", "%")
    chart("GDP: programme less baseline, percentage points", ["d_gdp"], "R21", "pp")
    chart("Public debt, % of GDP", ["debt_pct_gdp", "debt_base_pct_gdp"], "H37", "% GDP")
    chart("Public debt: programme less baseline, pp of GDP", ["d_debt"], "R37", "pp")
    chart("LIC-DSF-lite: PV of public debt / GDP (2025-2034)",
          ["dsf_base|PV public debt / GDP", "dsf_prog|PV public debt / GDP", "dsf_thr|PV public debt / GDP"],
          "H53", "%", 2034)
    chart("LIC-DSF-lite: PV of public debt, scenario less DSF baseline, pp", ["d_dsf"], "R53", "pp", 2034)
    ws["R69"] = ("DSF difference: the integrated scenario (programme, pricing revenue and, for DIGNAR-lite countries, "
                 "the fall in fossil rents) against the DSF baseline built from the dashboard; the other three "
                 "differences are programme against no programme in the same macro model.")
    ws["R69"].font = font(8, italic=True, color=MUT)

    ws.freeze_panes = "A7"
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # ---------------------------------------------------------------- README
    wm = wb.create_sheet("README", 0)
    wm.sheet_view.showGridLines = False
    wm.column_dimensions["B"].width = 26; wm.column_dimensions["C"].width = 110
    wm["B2"] = "Macro-fiscal country tool"; wm["B2"].font = font(18, True)
    wm["B3"] = f"Version {VERSION}. {res.shape[0]} countries. Built from MacroFiscalModels/results/integrated_*.csv."
    wm["B3"].font = font(9, italic=True, color=MUT)
    rows = [
        ("What it is", "One country at a time, the whole chain from mitigation investment needs to debt sustainability, pre-computed for every country in the National Mitigation Investment Model (NMIM). No Python is needed to use it."),
        ("How to use it", "Go to the Country sheet and pick a country in the yellow cell. Every number and chart updates. If cells look empty, recalculate (F9; Ctrl+Alt+F9 for a full recalculation)."),
        ("The chain", "1. NMIM: incremental mitigation investment by who pays; the budget need is split into capital and spending (land use and non-CO2 build no capital). 2. PriceFeasibility: the most power-sector carbon pricing three political-feasibility limits allow (F1 tariff rise, F2 bottom-40% burden, F3 generators' payments); its tax revenue is earmarked for the budget need. 3. A macro model: DIGNAR-lite for resource-rich non-high-income countries (fossil rents fall 80% by 2050), DIGNAD-lite for low-income and concessional-reliant countries, GIMF-lite otherwise. 4. LIC-DSF-lite for low and lower-middle income countries."),
        ("Inputs", "The Tax and Macro-Fiscal Dashboard panel (debt, taxes, revenue, primary balance, PEFA PI-11, IDS creditor terms, fossil rents, income group), NMIM v0.2 and PriceFeasibility v0.1. The Settings section of the Country sheet says, for each setting, whether it is data or an income-group default."),
        ("Read with care", "Calibrations are illustrative. Energy savings enter GIMF-lite only, so GIMF-lite countries show larger GDP gains than DIGNAD-lite ones. Pricing revenue covers power only. The DSF baseline uses defaults for the external share of debt, exports and capacity; ratings are mechanical, not a DSA. Fossil-rent losses in DIGNAR-lite are the global transition, not the programme."),
        ("Not in it", "CPAT (developed separately). Changing an assumption needs a re-run in Python: python3 MacroFiscalModels/integrated.py --all, then python3 MacroFiscalModels/build_country_tool.py."),
        ("Sheets", "Country: the tool. Results: one row per country. Paths: annual series by country (2025-2050). Settings: calibration value and source by country. Countries: the pick list. Calc: the selected country's series behind the charts."),
        ("Documentation", "MacroFiscalModels/Docs/UserGuide_4_Integrated (how the chain works), MacroFiscalModels_Method_v0_1.md (equations, data, validation), CAVEATS.md and TODO.md."),
    ]
    rr = 5
    for k, v in rows:
        wm[f"B{rr}"] = k; wm[f"B{rr}"].font = font(10, True); wm[f"B{rr}"].alignment = Alignment(vertical="top")
        wm[f"C{rr}"] = v; wm[f"C{rr}"].font = font(10); wm[f"C{rr}"].alignment = Alignment(wrap_text=True, vertical="top")
        wm.row_dimensions[rr].height = 15 * (1 + len(v) // 120); rr += 2

    wb.active = 1
    wb.calculation.fullCalcOnLoad = True
    os.makedirs(OUTDIR, exist_ok=True)
    out = os.path.join(OUTDIR, f"MacroFiscal_CountryTool_v{VERSION}.xlsx")
    wb.save(out)
    print(f"wrote {out}: {res.shape[0]} countries, {len(paths)} path rows")


if __name__ == "__main__":
    main()
