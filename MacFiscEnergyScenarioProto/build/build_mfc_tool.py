#!/usr/bin/env python3
"""Build the Macro-Fiscal-Energy-Scenario Tool (formerly the Integrated MF-C tool).

The tool is the Global Fiscal Indicators dashboard with extra modules spliced in as blocks.
Each module turns a repository model's results into rows of the dashboard's long format
(iso3, country, region, income, is_ccdr, indicator, value, year, source) and brings its own
indicator and source rows (same columns as the dashboard's config). This script writes the
merged panel and config to build/output/merged/ and runs the dashboard's own builder on them,
so the integrated workbook reads, scores and looks exactly like the dashboard.

Inputs : TaxMacroFiscalDashboard/build/{output/global_data_long.csv, config/}
         per module, see MODULES (now: InvestmentNeeds NMIM v0.2 summary; MacroFiscalModels
         integrated run; the DIGNAD-R equations for the live Macro sheet, see macro_live.py)
Output : MacFiscEnergyScenarioProto/Integrated_MFC_Tool_v<version>.xlsx (version in config/tool.json)
         build/output/<module>_long.csv (the module's rows, committed)

To add a module: write an extract function (given the dashboard's long panel) returning long rows, add
config/<name>_indicators.csv and config/<name>_sources.csv, and register it in MODULES.
"""
import os, sys, shutil, subprocess
import numpy as np, pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, ".."))           # MacFiscEnergyScenarioProto/
DASH = os.path.join(ROOT, "..", "TaxMacroFiscalDashboard", "build")   # at the repository root
INV = os.path.join(ROOT, "InvestmentNeeds")
CONFIG = os.path.join(HERE, "config")
OUT = os.path.join(HERE, "output")
MERGED = os.path.join(OUT, "merged")
SHIP = os.path.abspath(os.path.join(HERE, ".."))


# ---------------------------------------------------------------- module: investment needs
NMIM = os.path.join(INV, "national_model", "results", "nmim_summary_v0_2.csv")
SRC = "nmim_summary_v0_2.csv (InvestmentNeeds NMIM v0.2, model output)"
SRC_PIMA = "nmim_summary_v0_2.csv + dashboard PEFA PI-11 (InvestmentNeeds NMIM v0.2, model output)"
# Public investment efficiency gap, as in the v1.1 workbook (IMF 2015, Making Public
# Investment More Efficient): AE 13%, low income 40%, others 27%; where the dashboard has a
# PEFA PI-11 score (1-4) the gap runs linearly from 40% at 1 to 13% at 4.
GAP_AE, GAP_EM, GAP_LIC = 0.13, 0.27, 0.40


def extract_investment_needs(long):
    nm = pd.read_csv(NMIM).set_index("iso")
    g30, g50 = nm.gdp_2030, nm.gdp_2050
    # NMIM's budget need is P1 + a public share of P4 (ty_p4_gap); recover it from the
    # results rather than restating it, and check it is one number for every country.
    p4 = nm.typ4_2030.where(nm.typ4_2030 > 1e-9)
    gap_share = ((nm.bud_2030 - nm.typ1_2030) / p4).dropna()
    if gap_share.max() - gap_share.min() > 1e-6:
        sys.exit("NMIM budget need is not P1 + a fixed share of P4; check ty_p4_gap in nmim_v0_2.py")
    k4 = float(gap_share.median())
    out = []

    def emit(series, code, src=SRC, year=2030):
        for iso, v in series.items():
            if pd.notna(v) and np.isfinite(float(v)):
                out.append((iso, code, float(v), str(year), src))

    pct = lambda x, g=g30: x / g * 100
    # Investment 1: need and cost
    emit(pct(nm.inv2_2030), "nmim_nz_inv_2030_pct_gdp")
    emit(pct(nm.dinv_2030), "nmim_incr_inv_2030_pct_gdp")
    for k in ["gen", "grid", "trn", "bld", "ind", "nco", "oth"]:
        emit(pct(nm[f"dinv_{k}_2030"]), f"nmim_incr_inv_{k}_2030_pct_gdp")
    emit(pct(nm.dinv_2050, g50), "nmim_incr_inv_2050_pct_gdp", year=2050)
    emit(pct(nm.ica_2030), "nmim_incr_cost_a_2030_pct_gdp")
    emit(pct(nm.icb_2030), "nmim_incr_cost_b_2030_pct_gdp")
    emit(pct(nm.icc_2030), "nmim_incr_cost_c_2030_pct_gdp")
    emit(pct(nm.ica_2030 - nm.icb_2030), "nmim_coc_penalty_2030_pct_gdp")
    abt = nm.ghg1_2030 - nm.ghg2_2030                       # MtCO2e; costs in USD bn
    emit((nm.icb_2030 * 1000 / abt).where(abt > 1e-6), "nmim_cost_per_t_b_2030_usd")
    emit(abt / nm.ghg1_2030 * 100, "nmim_abatement_2030_pct")
    emit(pct(nm.pv_ic_b), "nmim_pv_incr_cost_b_pct_gdp", year="2025-2050")

    # Investment 2: public investment needs
    emit(pct(nm.bud_2030), "nmim_budget_need_2030_pct_gdp")
    emit(pct(nm.typ1_2030), "nmim_typ1_2030_pct_gdp")
    emit(pct(k4 * nm.typ4_2030), "nmim_p4_gap_2030_pct_gdp")
    emit(pct(nm.dbud_2030), "nmim_budget_need_incr_2030_pct_gdp")
    pefa = long[long.indicator == "pefa_pi11_pim_score"].set_index("iso3").value
    income = long.drop_duplicates("iso3").set_index("iso3").income
    gap = pd.Series({iso: GAP_LIC - (pefa[iso] - 1) / 3 * (GAP_LIC - GAP_AE) if iso in pefa.index
                     else GAP_AE if nm.grp[iso] == "AE"
                     else GAP_LIC if income.get(iso) == "Low income" else GAP_EM for iso in nm.index})
    emit(pct(nm.bud_2030 / (1 - gap)), "nmim_budget_need_pima_2030_pct_gdp", src=SRC_PIMA)
    emit(pct(nm.typ1_2030 + nm.typ2_2030 + k4 * nm.typ4_2030), "nmim_budget_plus_util_2030_pct_gdp")
    emit(pct(nm.typ1_2030 + nm.typ2_2030 + nm.typ4_2030), "nmim_public_all_2030_pct_gdp")
    emit(pct(nm.prv_2030), "nmim_private_2030_pct_gdp")
    emit(pct(nm.typ3_2030), "nmim_typ3_2030_pct_gdp")
    emit(pct(nm.typ5_2030), "nmim_typ5_2030_pct_gdp")
    emit(pct((1 - k4) * nm.typ4_2030), "nmim_p4_private_2030_pct_gdp")
    emit(nm.prv_2030 / nm.inv2_2030 * 100, "nmim_private_share_2030_pct")
    return pd.DataFrame(out, columns=["iso3", "indicator", "value", "year", "source"])


# ---------------------------------------------------------------- module: macro-fiscal
MFM = os.path.join(ROOT, "MacroFiscalModels", "results")
SRC_MF = "integrated_results.csv (MacroFiscalModels integrated run, model output)"
RISK = {"Low": 1, "Moderate": 2, "High": 3}


def extract_macro_fiscal(long):
    """The integrated run (dashboard -> NMIM -> PriceFeasibility -> GIMF / DIGNAD / DIGNAR-lite ->
    LIC-DSF-lite), pre-computed for every NMIM country by MacroFiscalModels/integrated.py."""
    ir = pd.read_csv(os.path.join(MFM, "integrated_results.csv")).set_index("iso3")
    pa = pd.read_csv(os.path.join(MFM, "integrated_paths.csv")).set_index(["iso3", "series"])
    out = []

    def emit(series, code, year=2030):
        for iso, v in series.items():
            if pd.notna(v) and np.isfinite(float(v)):
                out.append((iso, code, float(v), str(year), SRC_MF))

    emit(ir.financing_gap_2030_pct_gdp, "mf_financing_gap_2030_pct_gdp")
    emit(ir.budget_need_2030_pct_gdp, "mf_budget_need_2030_pct_gdp")
    emit(ir.pricing_revenue_2030_pct_gdp, "mf_pricing_revenue_2030_pct_gdp")
    emit(ir.pricing_share_of_need_pct, "mf_pricing_share_of_need_pct", "2025-2050")
    emit(ir.pricing_tau_t, "mf_carbon_tax_usd_t")
    emit(ir.pricing_abatement_pct, "mf_pricing_abatement_pct")
    emit(ir.gdp_dev_2030_pct, "mf_gdp_dev_2030_pct")
    emit(ir.gdp_dev_2050_pct, "mf_gdp_dev_2050_pct", 2050)
    p50 = lambda s: pa.xs(s, level="series")["2050"]
    emit(((1 + p50("gdp_dev_pct") / 100) / (1 + p50("gdp_dev_base_pct") / 100) - 1) * 100, "mf_gdp_vs_base_2050_pct", 2050)
    emit(ir.debt_2030_pct_gdp, "mf_debt_2030_pct_gdp")
    emit(ir.debt_2050_pct_gdp, "mf_debt_2050_pct_gdp", 2050)
    emit(p50("debt_pct_gdp") - p50("debt_base_pct_gdp"), "mf_debt_vs_base_2050_pp", 2050)
    emit(ir.alt_gdp_dev_2050_pct, "mf_alt_gdp_dev_2050_pct", 2050)
    emit(ir.alt_debt_2050_pct_gdp, "mf_alt_debt_2050_pct_gdp", 2050)
    emit(ir.dsf_ci, "mf_dsf_ci", "2024")
    emit(ir.dsf_rating_base.map(RISK), "mf_dsf_risk_base_code", "2024-2050")
    emit(ir.dsf_rating_programme.map(RISK), "mf_dsf_risk_code", "2024-2050")
    return pd.DataFrame(out, columns=["iso3", "indicator", "value", "year", "source"])


# ---------------------------------------------------------------- module: CDPM-method cross-check
V11 = os.path.join(INV, "Mitigation_Investment_Needs_v1.1.xlsx")
SRC_V11 = "Mitigation_Investment_Needs_v1.1.xlsx (InvestmentNeeds v1.1, CDPM-based method, model output)"


def extract_cdpm_crosscheck(long):
    """The v1.1 workbook: the CDPM-based global method (group totals allocated to countries).
    Reads the values cached in the workbook, so it must have been recalculated and saved."""
    import openpyxl
    wb = openpyxl.load_workbook(V11, read_only=True, data_only=True)
    rows = list(wb["Results_Country"].iter_rows(min_row=4, values_only=True))
    rc = pd.DataFrame(rows[1:], columns=rows[0]).dropna(subset=["ISO3"]).set_index("ISO3")
    if rc["Total investment I"].isna().all():
        sys.exit(f"{V11} has no cached values: open and save it in Excel, or recalculate with LibreOffice.")
    inp = {r[0]: r[2] for r in wb["Inputs"].iter_rows(values_only=True) if r and isinstance(r[0], str)}
    share = rc.Group.map({"AE": inp.get("p_ps_ae"), "China": inp.get("p_ps_chn"), "EMDC": inp.get("p_ps_emdc")})
    out = []
    for code, ser in [("cdpm_total_inv_2030_pct_gdp", rc["I, % GDP"].astype(float) * 100),
                      ("cdpm_incr_inv_2030_pct_gdp", rc["dI, % GDP"].astype(float) * 100),
                      ("cdpm_incr_cost_b_2030_pct_gdp", rc["(b) % GDP"].astype(float) * 100),
                      ("cdpm_public_need_2030_pct_gdp", rc["I, % GDP"].astype(float) * 100 * share.astype(float))]:
        for iso, v in ser.items():
            if pd.notna(v):
                out.append((iso, code, float(v), "2030", SRC_V11))
    return pd.DataFrame(out, columns=["iso3", "indicator", "value", "year", "source"])


# name, extract function, block its blocks follow in reading order, dashboard blocks it replaces
MODULES = [("investment_needs", extract_investment_needs, "Context 3", []),
           ("macro_fiscal", extract_macro_fiscal, "Investment 2", ["Appendix D"]),
           ("cdpm_crosscheck", extract_cdpm_crosscheck, "Appendix E", [])]


# ---------------------------------------------------------------- merge and build
def main():
    os.makedirs(MERGED, exist_ok=True)
    long = pd.read_csv(os.path.join(DASH, "output", "global_data_long.csv"), dtype={"year": str})
    cfg = pd.read_csv(os.path.join(DASH, "config", "indicators.csv")).sort_values("order")
    src = pd.read_csv(os.path.join(DASH, "config", "sources.csv"))
    meta = long.drop_duplicates("iso3").set_index("iso3")[["country", "region", "income", "is_ccdr"]]

    for name, extract, after, replaces in MODULES:
        if replaces:
            gone = set(cfg.code[cfg.block.isin(replaces)])
            cfg = cfg[~cfg.block.isin(replaces)]
            long = long[~long.indicator.isin(gone)]
        rows = extract(long)
        dropped = sorted(set(rows.iso3) - set(meta.index))
        if dropped:
            print(f"{name}: {len(dropped)} economies not in the dashboard, left out: {', '.join(dropped)}")
        rows = rows[rows.iso3.isin(meta.index)].assign(value=lambda d: d.value.round(3))   # as the panel
        rows = rows.join(meta, on="iso3")[long.columns]
        rows.to_csv(os.path.join(OUT, f"{name}_long.csv"), index=False)
        mcfg = pd.read_csv(os.path.join(CONFIG, f"{name}_indicators.csv"))
        missing = set(mcfg.code) - set(rows.indicator)
        if missing:
            sys.exit(f"MODULE/CONFIG MISMATCH ({name}): config names indicators the module never emits: {sorted(missing)}")
        clash = set(mcfg.code) & set(cfg.code)
        if clash:
            sys.exit(f"{name}: indicator codes already in the dashboard: {sorted(clash)}")
        at = cfg.index[cfg.block == after].max()
        pos = list(cfg.index).index(at) + 1
        cfg = pd.concat([cfg.iloc[:pos], mcfg, cfg.iloc[pos:]], ignore_index=True)
        src = pd.concat([src, pd.read_csv(os.path.join(CONFIG, f"{name}_sources.csv"))], ignore_index=True)
        long = pd.concat([long, rows], ignore_index=True)
        print(f"{name}: {len(rows)} rows, {rows.iso3.nunique()} economies, {len(mcfg)} indicators")

    cfg["order"] = range(1, len(cfg) + 1)
    long.to_csv(os.path.join(MERGED, "global_data_long.csv"), index=False)
    cfg.to_csv(os.path.join(MERGED, "indicators.csv"), index=False)
    src.to_csv(os.path.join(MERGED, "sources.csv"), index=False)
    shutil.copy(os.path.join(CONFIG, "tool.json"), os.path.join(MERGED, "tool.json"))

    subprocess.run([sys.executable, os.path.join(DASH, "scripts", "build_dashboard.py"),
                    "--out", MERGED, "--config", MERGED, "--ship", SHIP], check=True)

    # sheets the dashboard builder does not make: the live simple macro model
    import json, macro_live
    tool = json.load(open(os.path.join(CONFIG, "tool.json")))
    macro_live.add(os.path.join(SHIP, f"{tool['file_stem']}_v{tool['version']}.xlsx"), long)


if __name__ == "__main__":
    main()
