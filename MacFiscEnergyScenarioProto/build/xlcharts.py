"""Charts for the Macro-Fiscal-Energy-Scenario Tool, styled one way: fixed colours by role, a numeric year axis
labelled every five years, legend below the plot, no axis titles (the chart title carries the unit),
and hidden cells plotted (chart data often sits in hidden columns)."""
from openpyxl.chart import ScatterChart, AreaChart, Reference, Series
from openpyxl.chart.series import SeriesLabel, StrRef
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RichTextProperties, Font as DFont

# role -> (colour, dash, width in EMU)
STYLE = {
    "policy": ("1F4E79", None, 28575),
    "baseline": ("8C8C8C", "dash", 22225),
    "run": ("C05621", None, 22225),            # the integrated run (pre-computed)
    "run_base": ("E3A77F", "dash", 19050),
    "alt": ("2E8B88", "sysDot", 19050),          # the other model family
    "limit": ("B22222", "dash", 19050),          # thresholds, benchmarks
    "green": ("4E8542", None, 28575),
    "green2": ("A3C585", None, 22225),
    "purple": ("7A5195", None, 22225),
    "gold": ("C9A227", None, 22225),
}


def _txt(size=800, bold=False, color="404040", rot=None):
    cp = CharacterProperties(sz=size, b=bold, solidFill=color, latin=DFont(typeface="Arial"))
    body = RichTextProperties(rot=rot, vert="horz") if rot is not None else RichTextProperties()
    return RichText(bodyPr=body, p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def _axes(ch, yfmt, xmin, xmax, unit):
    x, y = ch.x_axis, ch.y_axis
    x.delete = False; y.delete = False
    x.number_format = "0"; x.txPr = _txt(800); x.title = None; x.majorTickMark = "out"
    if xmin is not None:
        x.scaling.min, x.scaling.max, x.majorUnit = xmin, xmax, unit
    x.majorGridlines = None
    x.crosses = "min"; y.crosses = "min"                 # axes at the bottom and left, also for negative values
    y.number_format = yfmt; y.txPr = _txt(800); y.title = None
    y.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E3E6EA", w=6350))
    x.spPr = GraphicalProperties(ln=LineProperties(solidFill="A6A6A6", w=9525))
    y.spPr = GraphicalProperties(ln=LineProperties(noFill=True))


def _frame(ch, title, height, width):
    ch.title = title
    ch.title.tx.rich.p[0].pPr = ParagraphProperties(defRPr=CharacterProperties(sz=1000, b=True, solidFill="1F2A36",
                                                                               latin=DFont(typeface="Arial")))
    ch.title.overlay = False
    ch.height, ch.width = height, width
    ch.legend.position = "b"; ch.legend.overlay = False; ch.legend.txPr = _txt(800)
    ch.visible_cells_only = False
    ch.display_blanks = "gap"
    ch.plot_area.layout = Layout(manualLayout=ManualLayout(x=0.08, y=0.13, w=0.88, h=0.6, xMode="edge", yMode="edge"))


def line(ws_data, cat, series, title, height=6.6, width=12.5, yfmt="0", skip=5, xmin=None, xmax=None):
    """Lines against years on a numeric axis (a scatter chart), so the year labels fall every `skip`
    years in Excel and LibreOffice alike. cat: (col, first_row, last_row) of the years on ws_data;
    series: (col, header_row, last_row, role), the header cell names the series."""
    ch = ScatterChart(); ch.scatterStyle = "line"; ch.varyColors = False; ch.style = 2
    _frame(ch, title, height, width)
    c, r0, r1 = cat
    xs = Reference(ws_data, min_col=c, max_col=c, min_row=r0, max_row=r1)
    for col, hr, lr, role in series:
        ys = Reference(ws_data, min_col=col, max_col=col, min_row=r0, max_row=r1)
        s = Series(ys, xs, title=None)
        s.tx = SeriesLabel(strRef=StrRef(f"{quote_sheetname(ws_data.title)}!${get_column_letter(col)}${hr}"))
        colr, dash, w = STYLE[role]
        s.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=colr, w=w, prstDash=dash))
        s.smooth = False; s.marker.symbol = "none"
        ch.series.append(s)
    y0 = ws_data.cell(r0, c).value; y1 = ws_data.cell(r1, c).value
    lo = xmin if xmin is not None else (y0 if isinstance(y0, int) else None)
    hi = xmax if xmax is not None else (y1 if isinstance(y1, int) else None)
    _axes(ch, yfmt, lo, hi, skip)
    return ch


def area(ws_data, cat, series, title, height=6.6, width=12.5, yfmt="0.0"):
    """Stacked areas against a category axis. cat: (col, first_row, last_row) of the labels; use a
    label column with a year only every five years, as category axes print every label.
    series: (col, header_row, last_row, colour)."""
    ch = AreaChart(); ch.grouping = "stacked"; ch.varyColors = False; ch.style = 2
    _frame(ch, title, height, width)
    c, r0, r1 = cat
    for col, hr, lr, colr in series:
        ch.add_data(Reference(ws_data, min_col=col, max_col=col, min_row=hr, max_row=lr), titles_from_data=True)
        s = ch.series[-1]
        s.graphicalProperties = GraphicalProperties(solidFill=colr, ln=LineProperties(noFill=True))
    ch.set_categories(Reference(ws_data, min_col=c, max_col=c, min_row=r0, max_row=r1))
    _axes(ch, yfmt, None, None, None)
    ch.x_axis.tickLblSkip = 1
    return ch
