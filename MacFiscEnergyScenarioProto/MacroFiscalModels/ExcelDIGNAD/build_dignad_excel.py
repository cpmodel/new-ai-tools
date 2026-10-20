"""Build DIGNAD-R Excel: a recursive, formula-only simplification of DIGNAD-lite / DIGNAR-lite in one
workbook, across a dozen countries. Proof of concept.

    python3 build_dignad_excel.py              # -> DIGNAD_R_Excel_v<VERSION>.xlsx (this folder)
    python3 build_dignad_excel.py --check      # also recompute in Python (same equations) for the check

Norms (as in the Tax and Macro-Fiscal Dashboard): inputs in Data_wide (one row per country, value
columns then __year columns) and Data_long (iso3, indicator, year, value, source); Spec lists every
input with its block and source; outputs leave in the same long form (sheet Outputs) so a master tool
can read them. Assumptions are named cells (p_...), yellow, on one sheet.

Versioning: bump VERSION, move the old workbook to Archive/, add a line to VersionNotes.md.
"""
import os
import sys
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

VERSION = "0.3"
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(HERE, "..", "..")
PANEL = os.path.join(ROOT, "..", "TaxMacroFiscalDashboard", "build", "output", "global_data_wide.csv")
PATHS = os.path.join(HERE, "..", "results", "integrated_paths.csv")
COUNTRIES = ["KEN", "SLE", "UGA", "LBR", "ETH", "RWA", "MOZ", "GHA", "SEN", "AGO", "COG", "TCD"]
Y0, Y1 = 2024, 2050
T = Y1 - Y0 + 1                                    # rows per country (2024 base + 26 years)

F = "Arial"; INK = "0E1B2A"; ACC = "C05621"; MUT = "5A6B7B"
HDR = PatternFill("solid", fgColor=INK); YEL = PatternFill("solid", fgColor="FFF2CC")
AUG = PatternFill("solid", fgColor="E3EAF2"); SEC = PatternFill("solid", fgColor="E3EAF2")
thin = Side(style="thin", color="C8D0D6"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
INCOME = ["Low income", "Lower middle income", "Upper middle income", "High income"]

# dashboard indicators read (block "Dashboard") and inputs added for this model (block "Augmented")
DASH = [("gross_debt_pct_gdp", "Gross debt, general govt (% GDP)", "IMF WEO"),
        ("tax_revenue_pct_gdp", "Tax revenue (% GDP)", "IMF WoRLD"),
        ("pefa_pi11_pim_score", "PEFA PI-11 public investment management (1-4)", "PEFA"),
        ("ids_multilateral_share_pct", "Multilateral share of PPG external debt (%)", "World Bank IDS"),
        ("ids_private_creditor_rate_pct", "Interest rate, private creditors (%)", "World Bank IDS"),
        ("fossil_rents_total_pct_gdp", "Fossil rents, oil + gas + coal (% GDP)", "WDI")]
# augmented: income-group defaults until data are added to the dashboard (LIC, LMIC, UMIC, HIC)
AUGD = [("public_investment_pct_gdp", "Public investment (% GDP)", (6.0, 5.0, 4.5, 3.5), "Dashboard Appendix C (IMF ICSD, general govt GFCF, to 2019) where available; else default by income group"),
        ("return_public_capital_pct", "Return on public capital (%)", (20, 17, 15, 10), "Default by income group (IMF literature 10-25%)"),
        ("hand_to_mouth_pct", "Hand-to-mouth households (%)", (60, 50, 35, 25), "Default by income group (Findex orders of magnitude); candidate: Findex"),
        ("grants_pct_gdp", "Grants (% GDP)", (2.0, 1.0, 0.0, 0.0), "Default by income group; candidate: OECD CRS / IMF GFS"),
        ("fiscal_rule_speed", "Fiscal rule speed needed to stabilise debt", None, "MacroFiscalModels integrated run: 0.05 default, 0.2 / 0.4 where it had to tighten (DIGNAR-lite starts at 0.2)")]
LONG = [("pub_cap", "Budget need, capital (% GDP)", "NMIM v0.2 via MacroFiscalModels"),
        ("pub_spend", "Budget need, spending: land use, non-CO2 (% GDP)", "NMIM v0.2 via MacroFiscalModels"),
        ("rev", "Feasible power carbon-pricing revenue (% GDP)", "PriceFeasibility v0.1 x NMIM power CO2")]

# assumptions: (name, value, label, source)
PARAMS = [("p_alpha", 0.33, "Private capital share", "Standard"),
          ("p_delta", 0.07, "Private depreciation", "Standard"),
          ("p_delta_g", 0.04, "Public depreciation", "Standard"),
          ("p_lambda_k", 0.30, "Private capital: share of the gap to desired capital closed each year", "Assumption (stands in for DIGNAD's forward-looking investment)"),
          ("p_lambda_f", 0.05, "Fiscal rule: primary-balance adjustment per unit of debt ratio above its start, if Data_wide has no fiscal_rule_speed", "As DIGNAD-lite"),
          ("p_lambda_f_res", 0.20, "Same, resource-rich countries (rents >= threshold)", "As DIGNAR-lite"),
          ("p_res_threshold", 0.10, "Rents threshold for the resource block (share of GDP)", "As the integrated run"),
          ("p_eta", 0.05, "Commercial rate premium per unit of debt ratio above its start", "As DIGNAD-lite"),
          ("p_rk", 0.01, "Concessional real rate", "As DIGNAD-lite"),
          ("p_amort", 0.03, "Repayment rate of new concessional loans", "As DIGNAD-lite"),
          ("p_tr_target", 0.02, "Transfers at the start (government consumption is the residual)", "As calibration.py"),
          ("p_gov_take", 0.50, "Government share of resource rents", "Assumption"),
          ("p_rent_decline", 0.80, "Fall in rents by 2050 on a global net zero path", "Assumption (IEA NZE order of magnitude)"),
          ("p_programme", 1, "Programme scenario includes the NMIM budget need? 1 = yes", "Switch"),
          ("p_pricing", 1, "Pricing revenue earmarked for the programme? 1 = yes", "Switch"),
          ("p_sh_grant", 0.30, "Share of the programme financed by grants", "As DIGNAD-lite"),
          ("p_sh_conc", 0.30, "Share financed by concessional loans (rest: commercial)", "As DIGNAD-lite"),
          ("p_dis_year", 0, "Disaster year (0 = none; e.g. 2030)", "Switch"),
          ("p_eps_k", 0.10, "Disaster: private capital destroyed", "As DIGNAD-lite"),
          ("p_eps_g", 0.10, "Disaster: public capital destroyed", "As DIGNAD-lite"),
          ("p_eps_a", 0.03, "Disaster: TFP loss", "As DIGNAD-lite"),
          ("p_rho_a", 0.60, "Disaster: TFP loss persistence", "As DIGNAD-lite")]
PV = {k: v for k, v, *_ in PARAMS}


def font(size=10, bold=False, color=INK, italic=False):
    return Font(name=F, size=size, bold=bold, color=color, italic=italic)


def header(ws, cols, widths, row=1):
    for j, (c, w) in enumerate(zip(cols, widths), 1):
        x = ws.cell(row, j, c); x.font = font(9, True, "FFFFFF"); x.fill = HDR; x.border = BOX
        ws.column_dimensions[get_column_letter(j)].width = w


def load_inputs():
    w = pd.read_csv(PANEL).set_index("iso3").loc[COUNTRIES]
    wide = w[["country", "region", "income"]].copy()
    for code, _, _ in DASH:
        wide[code] = w[code]
    ir = pd.read_csv(os.path.join(HERE, "..", "results", "integrated_results.csv")).set_index("iso3")
    def rule(iso):
        st = str(ir.loc[iso, "status"]) if iso in ir.index else ""
        if "tightened to" in st:
            return float(st.split("tightened to")[1])
        return 0.2 if ir.loc[iso, "model"] == "DIGNAR-lite" else 0.05
    years = {c: w[c + "__year"] for c, *_ in DASH}
    for code, _, vals, _ in AUGD:
        dflt = pd.Series([vals[INCOME.index(i)] for i in w.income], index=w.index) if vals else None
        if code == "fiscal_rule_speed":
            wide[code] = [rule(c) for c in w.index]; years[code] = "2026 run"
        elif code in w.columns:                    # now in the dashboard: data where it exists
            wide[code] = w[code].fillna(dflt)
            years[code] = w[code + "__year"].where(w[code].notna(), "default")
        else:
            wide[code] = dflt; years[code] = "default"
    p = pd.read_csv(PATHS)
    p = p[p.iso3.isin(COUNTRIES) & p.series.isin([c for c, *_ in LONG])]
    long = p.melt(id_vars=["iso3", "series"], var_name="year", value_name="value").dropna()
    long["year"] = long.year.astype(int)
    long = long.rename(columns={"series": "indicator"})
    long["source"] = long.indicator.map({c: s for c, _, s in LONG})
    return wide.reset_index(), years, long.sort_values(["iso3", "indicator", "year"])


# ------------------------------------------------------------------------------------------- Python check
def python_check(wide, long):
    """The workbook's equations, in Python, to check the formulas (same calibration, same order)."""
    out = {}
    for _, c in wide.iterrows():
        k = INCOME.index(c.income)
        pefa = c.pefa_pi11_pim_score
        gap = 0.40 - (pefa - 1) / 3 * 0.27 if pd.notna(pefa) else (0.40, 0.27, 0.27, 0.13)[k]
        s = 1 - gap
        Rg = c.return_public_capital_pct / 100; ig = c.public_investment_pct_gdp / 100
        rr = c.fossil_rents_total_pct_gdp / 100 if pd.notna(c.fossil_rents_total_pct_gdp) else 0.0
        res = rr >= PV["p_res_threshold"]
        rents0 = rr / (1 - rr) if res else 0.0
        conc = c.ids_multilateral_share_pct / 100 if pd.notna(c.ids_multilateral_share_pct) else (0.70, 0.45, 0.15, 0.0)[k]
        pr = c.ids_private_creditor_rate_pct
        rstar = min(max(pr / 100 - 0.03, 0.03), 0.08) if pd.notna(pr) else (0.04 if k < 2 else 0.03)
        tax = c.tax_revenue_pct_gdp / 100; grants0 = c.grants_pct_gdp / 100
        a, d, dg = PV["p_alpha"], PV["p_delta"], PV["p_delta_g"]
        K0 = a * (1 - tax) / (rstar + d); KG0 = s * ig / dg; psi = Rg * KG0
        A = 1 / (KG0 ** psi * K0 ** a)
        GDP0 = 1 + rents0; D0 = c.gross_debt_pct_gdp / 100 * GDP0; Dk0 = conc * D0
        int0 = PV["p_rk"] * Dk0 + rstar * (D0 - Dk0)
        rev0 = tax + PV["p_gov_take"] * rents0 + grants0
        g = max(0.05, rev0 - ig - int0 - PV["p_tr_target"]); trbar = rev0 - g - ig - int0
        lam_f = c.fiscal_rule_speed if pd.notna(c.fiscal_rule_speed) else (PV["p_lambda_f_res"] if res else PV["p_lambda_f"])
        ser = lambda ind: {r.year: r.value / 100 for r in long[(long.iso3 == c.iso3) & (long.indicator == ind)].itertuples()}
        cap, spd, rv = ser("pub_cap"), ser("pub_spend"), ser("rev")
        sc = {}
        for X, on in (("B", 0), ("P", 1)):
            KG, K, D, Dk, rc, GDPp = KG0, K0, D0, Dk0, rstar, GDP0
            path = []
            for t in range(T):
                yr = Y0 + t
                if t == 0:
                    Y, GDP, tr, ratio = 1.0, GDP0, trbar, D0 / GDP0
                    path.append(dict(GDP=GDP, ratio=ratio, tr=tr, chtm=(1 - tax) * (1 - a) * Y + tr, Y=Y, rents=rents0))
                    continue
                dy = PV["p_dis_year"]
                dis = 1 if dy and yr == dy else 0
                tfp = 1 - PV["p_eps_a"] * PV["p_rho_a"] ** (yr - dy) if dy and yr >= dy else 1
                rents = rents0 * (1 - PV["p_rent_decline"] * min(1, t / (Y1 - Y0)))
                Pc = on * PV["p_programme"] * cap.get(yr, 0); Ps = on * PV["p_programme"] * spd.get(yr, 0)
                R = on * PV["p_programme"] * PV["p_pricing"] * rv.get(yr, 0)
                KGd, Kd = KG * (1 - PV["p_eps_g"] * dis), K * (1 - PV["p_eps_k"] * dis)
                Y = A * tfp * KGd ** psi * Kd ** a
                Kstar = a * (1 - tax) * Y / (rstar + d)
                Kn = Kd + PV["p_lambda_k"] * (Kstar - Kd)
                KGn = (1 - dg) * KGd + s * (ig + Pc)
                GDP = Y + rents
                grants = grants0 + PV["p_sh_grant"] * (Pc + Ps)
                revenue = tax * Y + PV["p_gov_take"] * rents + grants + R
                tr = trbar - lam_f * (D / GDPp - D0 / GDP0)
                prim = g + ig + Pc + Ps + tr - revenue
                interest = PV["p_rk"] * Dk + rc * (D - Dk)
                Dn = D + interest + prim
                Dkn = Dk0 + (Dk - Dk0) * (1 - PV["p_amort"]) + PV["p_sh_conc"] * (Pc + Ps)
                ratio = Dn / GDP
                rc = rstar + PV["p_eta"] * max(0, ratio - D0 / GDP0)
                KG, K, D, Dk, GDPp = KGn, Kn, Dn, Dkn, GDP
                path.append(dict(GDP=GDP, ratio=ratio, tr=tr, chtm=(1 - tax) * (1 - a) * Y + tr, Y=Y, rents=rents))
            sc[X] = path
        out[c.iso3] = sc
    return out


# ------------------------------------------------------------------------------------------- workbook
def build(wide, years, long):
    wb = Workbook()
    names = {}

    def name(n, ref):
        dn = DefinedName(n, attr_text=ref)
        wb.defined_names[n] = dn

    # README ------------------------------------------------------------------
    wm = wb.active; wm.title = "README"; wm.sheet_view.showGridLines = False
    wm.column_dimensions["B"].width = 24; wm.column_dimensions["C"].width = 112
    wm["B2"] = "DIGNAD-R Excel"; wm["B2"].font = font(18, True)
    wm["B3"] = f"Version {VERSION}. Proof of concept: a recursive, formula-only DIGNAD-lite / DIGNAR-lite for {len(COUNTRIES)} countries, 2024-2050."
    wm["B3"].font = font(9, italic=True, color=MUT)
    txt = [
        ("What it is", "A simplified DIGNAD (Debt-Investment-Growth and Natural Disasters) with the DIGNAR resource block, written entirely in Excel formulas, computed for all countries at once. It answers: what does the NMIM public investment need, financed by grants, concessional and commercial debt and earmarked carbon-pricing revenue, do to GDP, public debt, transfers and the poorest households, with and without a global fall in fossil rents?"),
        ("How it differs from DIGNAD-lite", "DIGNAD-lite (Python) solves the whole path at once with forward-looking savers. Here every year follows from the previous one: private capital closes a share of the gap to its desired level each year (p_lambda_k) instead of a forward-looking investment choice, and savers' consumption is not modelled. The supply side (production with public capital and efficiency), the budget, the three kinds of financing, the debt-elastic commercial rate, the fiscal rule, the resource block and the disaster shock follow DIGNAD-lite. See the Check sheet for how close the two are."),
        ("How to use it", "Country sheet: pick a country (yellow cell); see the two scenarios and the charts. Assumptions sheet: change any yellow cell (named p_...), e.g. p_dis_year = 2030 for a disaster, p_pricing = 0 to drop the earmarked revenue, p_sh_grant for more grant financing. Everything recalculates (F9 if not automatic)."),
        ("Scenarios", "Baseline (B): no programme; fossil rents still fall (a global transition). Programme (P): the NMIM incremental budget need (capital builds public capital; land-use and non-CO2 spending does not), financed by p_sh_grant grants, p_sh_conc concessional loans and the rest commercial, less earmarked pricing revenue. P versus B isolates the programme; P versus 2024 is the total."),
        ("Inputs (dashboard norms)", "Data_wide: one row per country, dashboard columns (same codes as the Tax and Macro-Fiscal Dashboard) then the augmented block (shaded; income-group defaults until the data are added to the dashboard), then __year columns. Data_long: annual paths (iso3, indicator, year, value, source). Spec: every input with its block and source."),
        ("Outputs (for a master tool)", "Outputs sheet: results in the dashboard's long form (iso3, country, indicator, year, value, source), codes dig.*. Interface sheet: the contract - what the module needs and what it passes."),
        ("Sheets", "Country (view), Assumptions, Data_wide, Data_long, Spec, Calib (per-country calibration, formulas), Model (one row per country-year, both scenarios, formulas), Outputs, Interface, Check."),
        ("Read with care", "Illustrative calibration; four augmented inputs are income-group defaults. No energy savings, no private investment push, no exchange rate. A proof of concept, not a forecast."),
    ]
    r = 5
    for k, v in txt:
        wm[f"B{r}"] = k; wm[f"B{r}"].font = font(10, True); wm[f"B{r}"].alignment = Alignment(vertical="top")
        wm[f"C{r}"] = v; wm[f"C{r}"].font = font(10); wm[f"C{r}"].alignment = Alignment(wrap_text=True, vertical="top")
        wm.row_dimensions[r].height = 15 * (1 + len(v) // 115); r += 2

    # Assumptions -------------------------------------------------------------
    wa = wb.create_sheet("Assumptions")
    header(wa, ["Name", "Value", "Description", "Source"], [18, 10, 70, 60])
    for i, (n, v, lab, src) in enumerate(PARAMS, 2):
        wa.cell(i, 1, n).font = font(9, True)
        c = wa.cell(i, 2, v); c.fill = YEL; c.font = font(10, True); c.border = BOX
        wa.cell(i, 3, lab).font = font(9); wa.cell(i, 4, src).font = font(9, color=MUT)
        name(n, f"Assumptions!$B${i}")
    wa.freeze_panes = "A2"

    # Data_wide (dashboard form) ----------------------------------------------
    wd = wb.create_sheet("Data_wide")
    dcodes = [c for c, *_ in DASH]; acodes = [c for c, *_ in AUGD]
    cols = ["iso3", "country", "region", "income"] + dcodes + acodes + [c + "__year" for c in dcodes + acodes]
    header(wd, cols, [7, 22, 20, 18] + [14] * (len(cols) - 4))
    for i, row in enumerate(wide.itertuples(index=False), 2):
        rowd = row._asdict()
        for j, ccode in enumerate(cols, 1):
            if ccode.endswith("__year"):
                v = years[ccode[:-6]].get(rowd["iso3"]) if isinstance(years[ccode[:-6]], pd.Series) else years[ccode[:-6]]
                v = None if (isinstance(v, float) and np.isnan(v)) else (int(v) if isinstance(v, float) else v)
            else:
                v = rowd[ccode]
                v = None if (isinstance(v, float) and np.isnan(v)) else v
            x = wd.cell(i, j, v); x.font = font(9); x.border = BOX
            if ccode in acodes:
                x.fill = AUG
    wd.freeze_panes = "B2"
    DW = {c: get_column_letter(j) for j, c in enumerate(cols, 1)}
    NW = len(wide) + 1

    # Data_long ----------------------------------------------------------------
    wl = wb.create_sheet("Data_long")
    header(wl, ["iso3", "indicator", "year", "value", "source"], [7, 14, 7, 10, 40])
    for i, row in enumerate(long.itertuples(index=False), 2):
        for j, v in enumerate([row.iso3, row.indicator, int(row.year), float(row.value), row.source], 1):
            wl.cell(i, j, v).font = font(9)
    NL = len(long) + 1
    name("dl_iso", f"Data_long!$A$2:$A${NL}"); name("dl_ind", f"Data_long!$B$2:$B${NL}")
    name("dl_year", f"Data_long!$C$2:$C${NL}"); name("dl_val", f"Data_long!$D$2:$D${NL}")

    # Spec ---------------------------------------------------------------------
    wsp = wb.create_sheet("Spec")
    header(wsp, ["block", "code", "label", "where", "source"], [12, 30, 52, 12, 80])
    rr = 2
    for code, lab, src in DASH:
        for j, v in enumerate(["Dashboard", code, lab, "Data_wide", src + " (via the Tax and Macro-Fiscal Dashboard panel)"], 1):
            wsp.cell(rr, j, v).font = font(9)
        rr += 1
    for code, lab, _, src in AUGD:
        for j, v in enumerate(["Augmented", code, lab, "Data_wide", src], 1):
            x = wsp.cell(rr, j, v); x.font = font(9); x.fill = AUG
        rr += 1
    for code, lab, src in LONG:
        for j, v in enumerate(["Augmented", code, lab, "Data_long", src], 1):
            x = wsp.cell(rr, j, v); x.font = font(9); x.fill = AUG
        rr += 1

    # Calib (one row per country, formulas) -----------------------------------
    wc = wb.create_sheet("Calib")
    CAL = ["iso3", "country", "income", "k", "s", "Rg", "ig", "tax", "grants0", "conc", "rstar", "rr", "res", "rents0",
           "K0", "KG0", "psi", "A", "GDP0", "D0", "Dk0", "int0", "rev0", "g", "trbar", "lam_f", "ratio0"]
    header(wc, CAL, [7, 20, 18] + [9] * (len(CAL) - 3))
    CC = {c: get_column_letter(j) for j, c in enumerate(CAL, 1)}
    for i in range(2, NW + 1):
        W = lambda code: f"Data_wide!${DW[code]}${i}"
        f = {
            "iso3": f"=Data_wide!$A${i}", "country": f"=Data_wide!$B${i}", "income": f"=Data_wide!$D${i}",
            "k": f'=MATCH(C{i},{{"Low income","Lower middle income","Upper middle income","High income"}},0)',
            "s": f'=1-IF(ISNUMBER({W("pefa_pi11_pim_score")}),0.4-({W("pefa_pi11_pim_score")}-1)/3*0.27,CHOOSE(D{i},0.4,0.27,0.27,0.13))',
            "Rg": f"={W('return_public_capital_pct')}/100", "ig": f"={W('public_investment_pct_gdp')}/100",
            "tax": f"={W('tax_revenue_pct_gdp')}/100", "grants0": f"={W('grants_pct_gdp')}/100",
            "conc": f'=IF(ISNUMBER({W("ids_multilateral_share_pct")}),{W("ids_multilateral_share_pct")}/100,CHOOSE(D{i},0.7,0.45,0.15,0))',
            "rstar": f'=IF(ISNUMBER({W("ids_private_creditor_rate_pct")}),MIN(MAX({W("ids_private_creditor_rate_pct")}/100-0.03,0.03),0.08),IF(D{i}<3,0.04,0.03))',
            "rr": f'=IF(ISNUMBER({W("fossil_rents_total_pct_gdp")}),{W("fossil_rents_total_pct_gdp")}/100,0)',
            "res": f"=IF({CC['rr']}{i}>=p_res_threshold,1,0)",
            "rents0": f"=IF({CC['res']}{i}=1,{CC['rr']}{i}/(1-{CC['rr']}{i}),0)",
            "K0": f"=p_alpha*(1-{CC['tax']}{i})/({CC['rstar']}{i}+p_delta)",
            "KG0": f"={CC['s']}{i}*{CC['ig']}{i}/p_delta_g",
            "psi": f"={CC['Rg']}{i}*{CC['KG0']}{i}",
            "A": f"=1/({CC['KG0']}{i}^{CC['psi']}{i}*{CC['K0']}{i}^p_alpha)",
            "GDP0": f"=1+{CC['rents0']}{i}",
            "D0": f"={W('gross_debt_pct_gdp')}/100*{CC['GDP0']}{i}",
            "Dk0": f"={CC['conc']}{i}*{CC['D0']}{i}",
            "int0": f"=p_rk*{CC['Dk0']}{i}+{CC['rstar']}{i}*({CC['D0']}{i}-{CC['Dk0']}{i})",
            "rev0": f"={CC['tax']}{i}+p_gov_take*{CC['rents0']}{i}+{CC['grants0']}{i}",
            "g": f"=MAX(0.05,{CC['rev0']}{i}-{CC['ig']}{i}-{CC['int0']}{i}-p_tr_target)",
            "trbar": f"={CC['rev0']}{i}-{CC['g']}{i}-{CC['ig']}{i}-{CC['int0']}{i}",
            "lam_f": f"=IF(ISNUMBER({W('fiscal_rule_speed')}),{W('fiscal_rule_speed')},IF({CC['res']}{i}=1,p_lambda_f_res,p_lambda_f))",
            "ratio0": f"={CC['D0']}{i}/{CC['GDP0']}{i}",
        }
        for c in CAL:
            x = wc.cell(i, CAL.index(c) + 1, f[c]); x.font = font(9)
            if c not in ("iso3", "country", "income", "k", "res"):
                x.number_format = "0.0000"
    wc.freeze_panes = "B2"

    # Model (long panel; both scenarios) --------------------------------------
    wmod = wb.create_sheet("Model")
    BASE = ["iso3", "year", "t", "crow", "dis", "tfp", "rents", "cap", "spend", "rev"]
    SV = ["KG", "K", "Y", "Kstar", "GDP", "P", "grants", "revenue", "tr", "prim", "interest", "D", "Dk", "ratio", "rc", "chtm"]
    DEV = ["gdp_dev_prog", "gdp_vs_2024_P", "gdp_vs_2024_B", "debt_diff_pp", "chtm_dev_prog", "tr_diff"]
    MC = BASE + [f"B_{v}" for v in SV] + [f"P_{v}" for v in SV] + DEV
    header(wmod, MC, [7, 6, 4, 5] + [9] * (len(MC) - 4))
    M = {c: get_column_letter(j) for j, c in enumerate(MC, 1)}
    cal = lambda c, r: f"INDEX(Calib!${CC[c]}$1:${CC[c]}${NW},$D{r})"
    row = 2
    for n, iso in enumerate(wide.iso3):
        for t in range(T):
            r, p = row, row - 1
            yr = Y0 + t
            v = {"iso3": iso, "year": yr, "t": t, "crow": n + 2,
                 "dis": f"=IF(AND(p_dis_year>0,B{r}=p_dis_year),1,0)",
                 "tfp": f"=IF(AND(p_dis_year>0,B{r}>=p_dis_year),1-p_eps_a*p_rho_a^(B{r}-p_dis_year),1)",
                 "rents": f"={cal('rents0', r)}*(1-p_rent_decline*MIN(1,C{r}/{Y1 - Y0}))",
                 "cap": f'=IF(C{r}=0,0,SUMIFS(dl_val,dl_iso,$A{r},dl_ind,"pub_cap",dl_year,$B{r})/100)',
                 "spend": f'=IF(C{r}=0,0,SUMIFS(dl_val,dl_iso,$A{r},dl_ind,"pub_spend",dl_year,$B{r})/100)',
                 "rev": f'=IF(C{r}=0,0,SUMIFS(dl_val,dl_iso,$A{r},dl_ind,"rev",dl_year,$B{r})/100)'}
            for X, on in (("B", "0"), ("P", "p_programme")):
                c = lambda var, rr=r: f"{M[X + '_' + var]}{rr}"
                pr = lambda var: f"{M[X + '_' + var]}{p}"
                if t == 0:
                    v.update({f"{X}_KG": f"={cal('KG0', r)}", f"{X}_K": f"={cal('K0', r)}", f"{X}_Y": "=1",
                              f"{X}_Kstar": f"={cal('K0', r)}", f"{X}_GDP": f"={cal('GDP0', r)}", f"{X}_P": "=0",
                              f"{X}_grants": f"={cal('grants0', r)}", f"{X}_revenue": f"={cal('rev0', r)}",
                              f"{X}_tr": f"={cal('trbar', r)}", f"{X}_prim": f"=-{cal('int0', r)}",
                              f"{X}_interest": f"={cal('int0', r)}", f"{X}_D": f"={cal('D0', r)}",
                              f"{X}_Dk": f"={cal('Dk0', r)}", f"{X}_ratio": f"={c('D')}/{c('GDP')}",
                              f"{X}_rc": f"={cal('rstar', r)}",
                              f"{X}_chtm": f"=(1-{cal('tax', r)})*(1-p_alpha)*{c('Y')}+{c('tr')}"})
                    continue
                KGd = f"{pr('KG')}*(1-p_eps_g*$E{r})"; Kd = f"{pr('K')}*(1-p_eps_k*$E{r})"
                v.update({
                    f"{X}_Y": f"={cal('A', r)}*$F{r}*({KGd})^{cal('psi', r)}*({Kd})^p_alpha",
                    f"{X}_Kstar": f"=p_alpha*(1-{cal('tax', r)})*{c('Y')}/({cal('rstar', r)}+p_delta)",
                    f"{X}_K": f"={Kd}+p_lambda_k*({c('Kstar')}-{Kd})",
                    f"{X}_P": f"={on}*($H{r}+$I{r})",
                    f"{X}_KG": f"=(1-p_delta_g)*{KGd}+{cal('s', r)}*({cal('ig', r)}+{on}*$H{r})",
                    f"{X}_GDP": f"={c('Y')}+$G{r}",
                    f"{X}_grants": f"={cal('grants0', r)}+p_sh_grant*{c('P')}",
                    f"{X}_revenue": f"={cal('tax', r)}*{c('Y')}+p_gov_take*$G{r}+{c('grants')}+{on}*p_pricing*$J{r}",
                    f"{X}_tr": f"={cal('trbar', r)}-{cal('lam_f', r)}*({pr('D')}/{pr('GDP')}-{cal('ratio0', r)})",
                    f"{X}_prim": f"={cal('g', r)}+{cal('ig', r)}+{c('P')}+{c('tr')}-{c('revenue')}",
                    f"{X}_interest": f"=p_rk*{pr('Dk')}+{pr('rc')}*({pr('D')}-{pr('Dk')})",
                    f"{X}_D": f"={pr('D')}+{c('interest')}+{c('prim')}",
                    f"{X}_Dk": f"={cal('Dk0', r)}+({pr('Dk')}-{cal('Dk0', r)})*(1-p_amort)+p_sh_conc*{c('P')}",
                    f"{X}_ratio": f"={c('D')}/{c('GDP')}",
                    f"{X}_rc": f"={cal('rstar', r)}+p_eta*MAX(0,{c('ratio')}-{cal('ratio0', r)})",
                    f"{X}_chtm": f"=(1-{cal('tax', r)})*(1-p_alpha)*{c('Y')}+{c('tr')}"})
            first = row - t
            v.update({
                "gdp_dev_prog": f"=({M['P_GDP']}{r}/{M['B_GDP']}{r}-1)*100",
                "gdp_vs_2024_P": f"=({M['P_GDP']}{r}/{M['P_GDP']}{first}-1)*100",
                "gdp_vs_2024_B": f"=({M['B_GDP']}{r}/{M['B_GDP']}{first}-1)*100",
                "debt_diff_pp": f"=({M['P_ratio']}{r}-{M['B_ratio']}{r})*100",
                "chtm_dev_prog": f"=({M['P_chtm']}{r}/{M['B_chtm']}{r}-1)*100",
                "tr_diff": f"=({M['P_tr']}{r}-{M['B_tr']}{r})*100"})
            for c_ in MC:
                x = wmod.cell(r, MC.index(c_) + 1, v[c_]); x.font = font(8)
                if c_ not in ("iso3", "year", "t", "crow", "dis"):
                    x.number_format = "0.0000"
            row += 1
    NM = row - 1
    wmod.freeze_panes = "E2"

    # Outputs (dashboard long form) ------------------------------------------
    wo = wb.create_sheet("Outputs")
    header(wo, ["iso3", "country", "indicator", "year", "value", "source"], [7, 20, 34, 7, 10, 30])
    OUTS = [("dig.gdp_dev_programme_pct", "gdp_dev_prog", 2030), ("dig.gdp_dev_programme_pct", "gdp_dev_prog", 2050),
            ("dig.gdp_vs_2024_programme_pct", "gdp_vs_2024_P", 2050), ("dig.gdp_vs_2024_baseline_pct", "gdp_vs_2024_B", 2050),
            ("dig.debt_programme_pct_gdp", "P_ratio", 2030), ("dig.debt_programme_pct_gdp", "P_ratio", 2050),
            ("dig.debt_baseline_pct_gdp", "B_ratio", 2050), ("dig.debt_due_to_programme_pp", "debt_diff_pp", 2050),
            ("dig.htm_consumption_dev_programme_pct", "chtm_dev_prog", 2050), ("dig.transfers_change_pp_gdp", "tr_diff", 2050),
            ("dig.rents_pct_gdp", "rents", 2050)]
    pct100 = {"P_ratio", "B_ratio", "rents"}
    ro = 2
    for n, iso in enumerate(wide.iso3):
        for code, col, yr in OUTS:
            mrow = 2 + n * T + (yr - Y0)
            val = f"=Model!{M[col]}{mrow}*100" if col in pct100 else f"=Model!{M[col]}{mrow}"
            for j, v in enumerate([iso, f"=Calib!B{n + 2}", code, yr, val, f"DIGNAD-R Excel v{VERSION} (model output)"], 1):
                x = wo.cell(ro, j, v); x.font = font(9)
            wo.cell(ro, 5).number_format = "0.00"
            ro += 1
    wo.freeze_panes = "A2"

    # Interface ---------------------------------------------------------------
    wi = wb.create_sheet("Interface")
    header(wi, ["direction", "code", "label", "unit", "form", "notes"], [10, 34, 52, 10, 12, 60])
    rows = [("in", c, l, "%", "Data_wide", "dashboard column, same code") for c, l, _ in DASH]
    rows += [("in", c, l, "%" if c != "fiscal_rule_speed" else "x", "Data_wide", "augmented; default until in the dashboard" if c != "fiscal_rule_speed" else "augmented; passed in from the integrated run (module-to-module example)") for c, l, *_ in AUGD]
    rows += [("in", c, l, "% GDP", "Data_long", "annual 2025-2050") for c, l, _ in LONG]
    labels = {"dig.gdp_dev_programme_pct": "GDP, programme vs baseline", "dig.gdp_vs_2024_programme_pct": "GDP vs 2024, programme",
              "dig.gdp_vs_2024_baseline_pct": "GDP vs 2024, baseline (incl. rent fall)", "dig.debt_programme_pct_gdp": "Public debt, programme",
              "dig.debt_baseline_pct_gdp": "Public debt, baseline", "dig.debt_due_to_programme_pp": "Debt added by the programme",
              "dig.htm_consumption_dev_programme_pct": "Hand-to-mouth consumption, programme vs baseline",
              "dig.transfers_change_pp_gdp": "Transfers, programme vs baseline (fiscal rule)", "dig.rents_pct_gdp": "Fossil rents (resource block)"}
    seen = set()
    for code, col, yr in OUTS:
        if code in seen:
            continue
        seen.add(code)
        unit = "pp GDP" if "pp" in code else ("% GDP" if "pct_gdp" in code else "%")
        rows.append(("out", code, labels[code], unit, "Outputs (long)", "years " + ", ".join(str(y) for c_, _, y in OUTS if c_ == code)))
    for i, rw in enumerate(rows, 2):
        for j, v in enumerate(rw, 1):
            x = wi.cell(i, j, v); x.font = font(9)
            if rw[0] == "out":
                x.fill = AUG

    # Check: Excel (recursive) against Python DIGNAD-lite / DIGNAR-lite (perfect foresight) -----------
    wch = wb.create_sheet("Check")
    ir = pd.read_csv(os.path.join(HERE, "..", "results", "integrated_results.csv")).set_index("iso3")
    header(wch, ["iso3", "model in the integrated run", "GDP vs 2024, 2050: Python", "GDP vs 2024, 2050: Excel",
                 "Debt 2050: Python", "Debt 2050: Excel", "comparable?"], [7, 24, 14, 14, 12, 12, 40])
    for i, iso in enumerate(wide.iso3, 2):
        mdl = ir.loc[iso, "model"]
        comp = mdl in ("DIGNAD-lite", "DIGNAR-lite")
        def oref(code, yr):
            return (f'=SUMIFS(Outputs!$E:$E,Outputs!$A:$A,$A{i},Outputs!$C:$C,"{code}",Outputs!$D:$D,{yr})')
        vals = [iso, mdl, float(ir.loc[iso, "gdp_dev_2050_pct"]) if comp else None, oref("dig.gdp_vs_2024_programme_pct", 2050),
                float(ir.loc[iso, "debt_2050_pct_gdp"]) if comp else None, oref("dig.debt_programme_pct_gdp", 2050),
                "yes" if comp else "no: GIMF-lite in the integrated run (adds energy savings, private investment)"]
        for j, v in enumerate(vals, 1):
            x = wch.cell(i, j, v); x.font = font(9)
            if j in (3, 4, 5, 6):
                x.number_format = "0.00"
    n = len(wide) + 3
    wch.cell(n, 1, ("Python: MacroFiscalModels/integrated.py (perfect foresight, forward-looking savers, consumption tax and "
                    "transfers both respond, fiscal rule tightened where needed). Excel: this workbook (recursive; private "
                    "capital closes p_lambda_k of the gap each year; one fiscal instrument). The Excel formulas themselves "
                    "were checked against a Python copy of the same recursive equations (build_dignad_excel.py --check): "
                    "equal to rounding.")).font = font(9, italic=True, color=MUT)

    # Country view ------------------------------------------------------------
    wv = wb.create_sheet("Country", 1); wv.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGH", [2, 46, 14, 14, 14, 2, 2, 2]):
        wv.column_dimensions[col].width = w
    wv["B2"] = "DIGNAD-R Excel: one country"; wv["B2"].font = font(18, True)
    wv["B3"] = f"Version {VERSION}. Recursive DIGNAD-lite / DIGNAR-lite, all formulas. Illustrative calibration."
    wv["B3"].font = font(9, italic=True, color=MUT)
    wv["B5"] = "Country (pick)"; wv["B5"].font = font(11, True)
    wv["C5"] = wide.country.iloc[0]; wv["C5"].fill = YEL; wv["C5"].font = font(11, True); wv.merge_cells("C5:D5")
    dv = DataValidation(type="list", formula1=f"=Calib!$B$2:$B${NW}"); wv.add_data_validation(dv); dv.add("C5")
    wv["E5"] = f'=INDEX(Calib!$A$2:$A${NW},MATCH(C5,Calib!$B$2:$B${NW},0))'; wv["E5"].font = font(11, True, ACC)
    wv["B6"] = "First row of the country in Model"; wv["B6"].font = font(8, color=MUT)
    wv["C6"] = f'=MATCH(E5,Model!$A$2:$A${NM},0)+1'; wv["C6"].font = font(8, color=MUT)
    rr = 8
    def kv(label, formula, fmt="0.00"):
        nonlocal rr
        wv[f"B{rr}"] = label; wv[f"B{rr}"].font = font(10); wv[f"B{rr}"].border = BOX
        x = wv[f"C{rr}"]; x.value = formula; x.number_format = fmt; x.font = font(10, True); x.border = BOX
        rr += 1
    crow = f'MATCH($E$5,Calib!$A$2:$A${NW},0)+1'
    C_ = lambda c: f'=INDEX(Calib!${CC[c]}$1:${CC[c]}${NW},{crow})'
    def mval(col, yr, scale=1):
        return f"=INDEX(Model!${M[col]}$1:${M[col]}${NM},$C$6+{yr - Y0})*{scale}"
    for c in "BCDE":
        wv[f"{c}{rr}"].fill = SEC
    wv[f"B{rr}"] = "Calibration"; wv[f"B{rr}"].font = font(11, True); rr += 1
    kv("Public investment efficiency s", C_("s")); kv("Return on public capital", C_("Rg"))
    kv("Resource block on? (1 = DIGNAR)", C_("res"), "0"); kv("Rents at start, share of non-resource GDP", C_("rents0"), "0.000")
    kv("Real commercial rate", C_("rstar"), "0.000"); kv("Concessional share of debt", C_("conc"))
    rr += 1
    for c in "BCDE":
        wv[f"{c}{rr}"].fill = SEC
    wv[f"B{rr}"] = "Results"; wv[f"C{rr}"] = "2030"; wv[f"D{rr}"] = "2050"
    for c in "BCD":
        wv[f"{c}{rr}"].font = font(11, True)
    rr += 1
    def kv2(label, col, scale=1, fmt="0.00"):
        nonlocal rr
        wv[f"B{rr}"] = label; wv[f"B{rr}"].font = font(10); wv[f"B{rr}"].border = BOX
        for cc, yr in (("C", 2030), ("D", 2050)):
            x = wv[f"{cc}{rr}"]; x.value = mval(col, yr, scale); x.number_format = fmt; x.font = font(10, True); x.border = BOX
        rr += 1
    kv2("GDP, programme vs baseline (%)", "gdp_dev_prog"); kv2("GDP vs 2024, programme (%)", "gdp_vs_2024_P")
    kv2("GDP vs 2024, baseline (%)", "gdp_vs_2024_B"); kv2("Public debt, programme (% GDP)", "P_ratio", 100, "0.0")
    kv2("Public debt, baseline (% GDP)", "B_ratio", 100, "0.0"); kv2("Debt added by the programme (pp)", "debt_diff_pp", 1, "0.0")
    kv2("Hand-to-mouth consumption, programme vs baseline (%)", "chtm_dev_prog")
    kv2("Transfers, programme vs baseline (pp GDP)", "tr_diff"); kv2("Fossil rents (% GDP)", "rents", 100, "0.0")

    # series for charts (sheet Calc)
    wk = wb.create_sheet("Calc")
    wk["A1"] = "series"
    for j, yr in enumerate(range(Y0, Y1 + 1), 2):
        wk.cell(1, j, yr)
    series = [("GDP, baseline (2024 = 100)", "B_GDP", "gdp"), ("GDP, programme (2024 = 100)", "P_GDP", "gdp"),
              ("Debt, baseline (% GDP)", "B_ratio", 100), ("Debt, programme (% GDP)", "P_ratio", 100),
              ("Transfers, baseline (% GDP)", "B_tr", 100), ("Transfers, programme (% GDP)", "P_tr", 100),
              ("Programme spending (% GDP)", "P_P", 100), ("Pricing revenue (% GDP)", "rev", 100), ("Fossil rents (% GDP)", "rents", 100)]
    for i, (lab, col, sc) in enumerate(series, 2):
        wk.cell(i, 1, lab)
        for j in range(T):
            ref = f"INDEX(Model!${M[col]}$1:${M[col]}${NM},Country!$C$6+{j})"
            if sc == "gdp":
                f = f"={ref}/INDEX(Model!${M[col]}$1:${M[col]}${NM},Country!$C$6)*100"
            else:
                f = f"={ref}*{sc}"
            wk.cell(i, j + 2, f).number_format = "0.00"
    wk.sheet_properties.tabColor = "BFBFBF"

    def chart(title, rows_, anchor, ytitle):
        ch = LineChart(); ch.title = title; ch.height = 7; ch.width = 16; ch.y_axis.title = ytitle; ch.legend.position = "b"
        ch.varyColors = False                      # one colour per series (Excel otherwise colours each year)
        for i in rows_:
            ch.add_data(Reference(wk, min_col=1, max_col=T + 1, min_row=i, max_row=i), titles_from_data=True, from_rows=True)
        ch.set_categories(Reference(wk, min_col=2, max_col=T + 1, min_row=1, max_row=1))
        ch.x_axis.delete = False; ch.y_axis.delete = False
        wv.add_chart(ch, anchor)
    chart("GDP (2024 = 100)", [2, 3], "G5", "index")
    chart("Public debt, % of GDP", [4, 5], "G20", "% GDP")
    chart("Transfers (fiscal rule), % of GDP", [6, 7], "G35", "% GDP")
    chart("Programme, pricing revenue and rents, % of GDP", [8, 9, 10], "G50", "% GDP")
    wv.page_setup.orientation = "landscape"; wv.page_setup.fitToWidth = 1; wv.page_setup.fitToHeight = 0
    wv.sheet_properties.pageSetUpPr.fitToPage = True

    wb.move_sheet("Assumptions", -(len(wb.sheetnames) - 3))
    wb.active = 1
    wb.calculation.fullCalcOnLoad = True
    out = os.path.join(HERE, f"DIGNAD_R_Excel_v{VERSION}.xlsx")
    wb.save(out)
    return out, M, NM


if __name__ == "__main__":
    wide, years, long = load_inputs()
    out, M, NM = build(wide, years, long)
    print("wrote", out)
    if "--check" in sys.argv:
        res = python_check(wide, long)
        rows = []
        for iso, sc in res.items():
            for t in (6, 26):
                b, p_ = sc["B"][t], sc["P"][t]
                rows.append([iso, Y0 + t, 100 * (p_["GDP"] / b["GDP"] - 1), 100 * (p_["GDP"] / sc["P"][0]["GDP"] - 1),
                             100 * p_["ratio"], 100 * b["ratio"]])
        df = pd.DataFrame(rows, columns=["iso3", "year", "gdp_dev_prog", "gdp_vs_2024_P", "debt_P", "debt_B"])
        df.to_csv(os.path.join(HERE, "check_python.csv"), index=False, float_format="%.6g")
        print(df.round(2).to_string(index=False))
