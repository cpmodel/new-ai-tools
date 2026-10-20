#!/usr/bin/env python3
# =============================================================================
# One-at-a-time sensitivity for the mitigation investment needs workbook.
# Copies the workbook, changes one input, recalculates with LibreOffice, and
# reads the incremental cost by lens. Writes the table into a "Sensitivity"
# sheet of the main workbook as VALUES (re-run this script after changing inputs),
# then recalculates the main workbook so its cached values are restored.
#
# Run:  python InvestmentNeeds/build/run_sensitivity.py [path/to/workbook.xlsx]
# Needs LibreOffice (soffice) on the PATH. Takes about a minute per case.
# =============================================================================
import os, sys, shutil, subprocess, tempfile
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

HERE = os.path.dirname(os.path.abspath(__file__))
WB = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "..", "Mitigation_Investment_Needs_v1.1.xlsx")
WB = os.path.abspath(WB)

CASES = [  # label, {input key: value}
 ("Base case (v1.1 defaults)", {}),
 ("v1.0 calibrated benefits for power and transport", {"p_eng": 0}),
 ("Grids credited with generation benefits", {"p_grid_credit": 1}),
 ("EV premium USD 2,500", {"p_ev_prem": 2500}),
 ("EV premium USD 10,000", {"p_ev_prem": 10000}),
 ("Renewable capex -25%", {"p_capex_pv": 570, "p_capex_wind": 870}),
 ("Renewable capex +25%", {"p_capex_pv": 950, "p_capex_wind": 1450}),
 ("Calibrated savings b -50% (buildings, industry)", {"p_b": 0.0665}),
 ("Calibrated savings b +50% (buildings, industry)", {"p_b": 0.1995}),
 ("Co-benefits c -50% (buildings, industry, AFOLU)", {"p_c": 0.0645}),
 ("EMDC default WACC 15%", {"p_r_emdc": 0.15}),
 ("Rates compared in real terms", {"p_basis": 1}),
 ("Fuels/CCS abatement cost USD 50/t", {"p_cpt_fuel": 50}),
 ("Fuels/CCS abatement cost USD 200/t", {"p_cpt_fuel": 200}),
 ("Incremental shares halved", {"p_inc_ae": 0.19, "p_inc_chn": 0.19, "p_inc_emdc": 0.12}),
]
ROWS = {"ica": "INCREMENTAL COST (a) market cost of capital",
        "icb": "INCREMENTAL COST (b) World Bank rate",
        "icc": "INCREMENTAL COST (c) WB rate + local co-benefits",
        "dI": "Incremental investment need, dI",
        "abt": "Abatement from dI (MtCO2e/yr)"}

def soffice_recalc(path):
    # headless recalculation via a Basic macro-free route: convert xlsx -> xlsx
    out = tempfile.mkdtemp()
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", out, path],
                   check=True, capture_output=True, timeout=600)
    shutil.move(os.path.join(out, os.path.basename(path)), path)

def recalc(path):
    rc = os.environ.get("RECALC_PY")
    if rc:
        subprocess.run([sys.executable, rc, path, "300"], check=True, capture_output=True)
    else:
        soffice_recalc(path)

def run_case(changes, tmpdir):
    f = os.path.join(tmpdir, "case.xlsx")
    wb = openpyxl.load_workbook(WB)
    ws = wb["Inputs"]
    for row in ws.iter_rows(min_row=5):
        if row[0].value in changes:
            row[2].value = changes[row[0].value]
    wb.save(f)
    recalc(f)
    s = openpyxl.load_workbook(f, data_only=True)["Results_Summary"]
    got = {r[0]: r[1:5] for r in s.iter_rows(min_row=5, max_row=30, max_col=5, values_only=True) if r[0]}
    return {k: got[v] for k, v in ROWS.items()}

tmp = tempfile.mkdtemp()
res = []
for label, ch in CASES:
    r = run_case(ch, tmp)
    res.append((label, ch, r))
    print(f"{label:52s} (a) {r['ica'][3]:8.1f}  (b) {r['icb'][3]:8.1f}  (c) {r['icc'][3]:8.1f}", flush=True)

wb = openpyxl.load_workbook(WB)
if "Sensitivity" in wb.sheetnames:
    del wb["Sensitivity"]
ws = wb.create_sheet("Sensitivity", 3)
F = "Arial"
ws["A1"] = "Sensitivity: incremental cost, 2030 (USD bn a year), one input changed at a time"
ws["A1"].font = Font(name=F, size=13, bold=True, color="1F3864")
ws["A2"] = ("STATIC VALUES written by build/run_sensitivity.py from the inputs as they stood when it ran. "
            "Re-run the script after changing inputs. Positive = net cost; negative = net saving.")
ws["A2"].font = Font(name=F, size=8, italic=True, color="595959")
hdr = ["Case", "Inputs changed", "World (a) market", "World (b) WB 6%", "World (c) WB + co-ben.",
       "EMDC (a)", "EMDC (b)", "EMDC (c)", "World dI", "Abatement (Mt/yr)", "World (b) USD/t"]
for j, h in enumerate(hdr):
    c = ws.cell(row=4, column=j + 1, value=h)
    c.font = Font(name=F, size=9, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864")
    c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
ws.row_dimensions[4].height = 36
for i, (label, ch, r) in enumerate(res):
    vals = [label, ", ".join(f"{k}={v}" for k, v in ch.items()) or "-",
            r["ica"][3], r["icb"][3], r["icc"][3], r["ica"][2], r["icb"][2], r["icc"][2],
            r["dI"][3], r["abt"][3], r["icb"][3] / r["abt"][3] * 1000 if r["abt"][3] else None]
    for j, v in enumerate(vals):
        c = ws.cell(row=5 + i, column=j + 1, value=v)
        c.font = Font(name=F, size=9, bold=(i == 0))
        if j >= 2: c.number_format = '#,##0;(#,##0);"-"'
for col, w in zip("ABCDEFGHIJK", (46, 40, 11, 11, 12, 10, 10, 10, 10, 11, 11)):
    ws.column_dimensions[col].width = w
wb.save(WB)
recalc(WB)
print("wrote Sensitivity sheet to", WB)
