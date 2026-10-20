#!/usr/bin/env python3
# =============================================================================
# NMIM v0.2 - workbook builder (CPAT norms)
# =============================================================================
# Writes national_model/NMIM_v0.2.xlsx:
#   ReadMe        purpose, use, colour legend, version log
#   Dashboard     selected country: investment, incremental cost, emissions (+ chart)
#   Settings      country selector (C4)
#   Inputs        parameters in MTInputs column structure (B..H); names p_<name>
#   MODULES->     divider
#   Mitigation    CPAT Mitigation layout, live formulas for the selected country,
#                 rendered from the same templates the Python engine evaluates
#   Results_All   all countries, from the Python engine (values)
#   Typology      selected country: who pays - type shares by sector and investment
#                 by type (live formulas reading Mitigation), power market facts
#   Check         selected country: workbook vs Python on key rows
#   DATA->        divider
#   Data          one row per country: every @field the formulas read
#   PowerMarket   one row per country: power market structure, private participation,
#                 cost recovery, RISE (data/power_market.csv)
#   PyKey         Python results for the key rows (values; read by Check)
#
# Run after nmim_v0_2.py:  python InvestmentNeeds/national_model/build_workbook_v0_2.py
# =============================================================================
import os, re
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.series import SeriesLabel

import nmim_v0_2 as M

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, f"NMIM_v{M.VERSION}.xlsx")
DEFAULT_ISO = "IND"

# CPAT colours (NORMS s2)
TITLE, SECTION, INPUT, CODE, BASE, UNUSED, REVIEW = "00B050", "92D050", "EBF1DE", "DCE6F1", "DDD9C4", "F2F2F2", "FFF2CC"
fill = lambda c: PatternFill("solid", fgColor=c)
F = "Arial"
f9, f9b = Font(name=F, size=9), Font(name=F, size=9, bold=True)
fw = Font(name=F, size=14, bold=True, color="FFFFFF")
fwb = Font(name=F, size=9, bold=True, color="FFFFFF")
fnote = Font(name=F, size=8, italic=True, color="595959")
fblue = Font(name=F, size=9, color="0000FF")
wrap = Alignment(wrap_text=True, vertical="top")

wb = Workbook()
wsR = wb.active; wsR.title = "ReadMe"
wsD = wb.create_sheet("Dashboard")
wsS = wb.create_sheet("Settings")
wsI = wb.create_sheet("Inputs")
wb.create_sheet("MODULES->")
wsM = wb.create_sheet("Mitigation")
wsA = wb.create_sheet("Results_All")
wsY = wb.create_sheet("Typology")
wsC = wb.create_sheet("Check")
wb.create_sheet("DATA->")
wsT = wb.create_sheet("Data")
wsP = wb.create_sheet("PowerMarket")
wsK = wb.create_sheet("PyKey")

def band(ws, row, text, color, last_col, font=f9b):
    for c in range(1, last_col + 1):
        ws.cell(row=row, column=c).fill = fill(color)
    x = ws.cell(row=row, column=1 if color == TITLE else 1, value=text); x.font = font

# -----------------------------------------------------------------------------
# DATA sheet
# -----------------------------------------------------------------------------
D, CAL = M.load_calibrated()
fields = set()
for _, _, _, t in M.SCALARS: fields |= set(re.findall(r"@(\w+)", t))
for r in M.ROWS:
    for t in [r["base"], r["proj"]] + (r["hist"] or []):
        fields |= set(re.findall(r"@(\w+)", t))
fields = ["name"] + sorted(fields - {"name"})
wsT["A1"] = "Data: one row per country (CPAT 2024 panel, Ember, WEO, CBAM tool, v1.1 WACC). Built by extract_data_v0_2.py."
wsT["A1"].font = f9b
wsT.cell(row=3, column=1, value="countrycode").font = f9b
for j, fl in enumerate(fields):
    c = wsT.cell(row=3, column=2 + j, value=fl); c.font = f9b; c.fill = fill(INPUT)
FCOL = {fl: CL(2 + j) for j, fl in enumerate(fields)}
isos = sorted(D)
for n, iso in enumerate(isos):
    wsT.cell(row=4 + n, column=1, value=iso).font = f9
    for j, fl in enumerate(fields):
        v = D[iso].get(fl)
        wsT.cell(row=4 + n, column=2 + j, value=(float(v) if isinstance(v, (int, float)) and fl != "name" else v)).font = f9
DATA_LAST = 3 + len(isos)
wsT.freeze_panes = "B4"

# -----------------------------------------------------------------------------
# SETTINGS
# -----------------------------------------------------------------------------
band(wsS, 1, "NMIM - Settings", TITLE, 6, fw)
wsS["B4"] = "Country (ISO3)"; wsS["C4"] = DEFAULT_ISO
wsS["C4"].fill = fill(INPUT); wsS["C4"].font = Font(name=F, size=11, bold=True, color="0000FF")
wsS["B5"] = "Country name"; wsS["C5"] = f'=INDEX(Data!$B:$B,MATCH($C$4,Data!$A:$A,0))'
wsS["B6"] = "Data row"; wsS["C6"] = f'=MATCH($C$4,Data!$A:$A,0)'
wsS["B8"] = "Choose a country in C4 (ISO3 code from the Data sheet). Every formula on Mitigation recalculates for it."
wsS["B8"].font = fnote
dv = DataValidation(type="list", formula1=f"=Data!$A$4:$A${DATA_LAST}", allow_blank=False)
wsS.add_data_validation(dv); dv.add("C4")
for c in ("B4", "B5", "B6", "C5", "C6"): wsS[c].font = f9b if c.startswith("B") else f9
wsS.column_dimensions["B"].width = 18; wsS.column_dimensions["C"].width = 30

# -----------------------------------------------------------------------------
# INPUTS (MTInputs column structure)
# -----------------------------------------------------------------------------
band(wsI, 1, "NMIM - Inputs (MTInputs column structure)", TITLE, 10, fw)
band(wsI, 2, "Parameters; calculations read column F by name (defined names p_<NameOfParameter>)", SECTION, 10)
hdr = ["", "Parameters", "unit", "MT values", "Dashboard values", "Used for calculation", "Defaults", "NameOfParameter", "Source", "Fit"]
for j, h in enumerate(hdr):
    c = wsI.cell(row=7, column=1 + j, value=h); c.font = f9b
for n, (name, val, unit, desc, src, fit) in enumerate(M.PARAMS):
    r = 8 + n
    wsI.cell(row=r, column=2, value=desc); wsI.cell(row=r, column=3, value=unit)
    wsI.cell(row=r, column=6, value=f"=G{r}")
    g = wsI.cell(row=r, column=7, value=val); g.font = fblue
    wsI.cell(row=r, column=8, value=name); wsI.cell(row=r, column=9, value=src); wsI.cell(row=r, column=10, value=fit)
    for c in range(2, 9): wsI.cell(row=r, column=c).fill = fill(INPUT)
    for c in range(2, 11):
        if c != 7: wsI.cell(row=r, column=c).font = f9
    wb.defined_names["p_" + name] = DefinedName("p_" + name, attr_text=f"Inputs!$F${r}")
wsI.column_dimensions["D"].hidden = True; wsI.column_dimensions["E"].hidden = True
for col, w in zip("BCFGHIJ", (48, 14, 12, 12, 16, 60, 5)): wsI.column_dimensions[col].width = w
wsI.freeze_panes = "B8"

# -----------------------------------------------------------------------------
# MITIGATION (CPAT layout)
# -----------------------------------------------------------------------------
YCOL = {y: CL(11 + (y - 2021)) for y in range(2021, 2051)}       # K=2021 .. AN=2050
LASTC = 11 + (2050 - 2021)
H1 = ["", "Fuel", "Sector", "Description", "Unit", "Source", "Input Code", "Output Code", "Note", "Helper"]
for j, h in enumerate(H1): wsM.cell(row=1, column=1 + j, value=h).font = f9b
for y, col in YCOL.items():
    c = wsM[f"{col}1"]; c.value = y; c.font = f9b
band(wsM, 2, "NMIM v0.2 - National Mitigation Investment Model (selected country: see Settings)", TITLE, LASTC, fwb)

row = 3
SROW, RROW = {}, {}
def section(title):
    global row
    band(wsM, row, title, SECTION, LASTC); row += 1

# section 1: country scalars in column J
section("1. Country scalars (column J)")
CTY_ROW = row
wsM.cell(row=row, column=4, value="Data row of the selected country")
wsM.cell(row=row, column=10, value="=Settings!$C$6"); wsM.cell(row=row, column=8, value="mit.sc.row")
row += 1
for k, desc, unit, t in M.SCALARS:
    SROW[k] = row; row += 1
for r in M.ROWS:
    pass
# lay out time rows by section
order = []
cur = None
for r in M.ROWS:
    if r["sec"] != cur:
        order.append(("SEC", r["sec"])); cur = r["sec"]
    order.append(("ROW", r))
rr = row
for kind, x in order:
    if kind == "SEC": rr += 1
    else: RROW[x["code"]] = rr; rr += 1

def xl(t, col=None, prev=None):
    s = t
    s = re.sub(r"LAG\(([\w.]+),\s*([^)]+)\)",
               lambda m: f"IF(YEAR-2024>({m.group(2)}),INDEX($O${RROW[m.group(1)]}:$AN${RROW[m.group(1)]},YEAR-2024-({m.group(2)})),0)", s)
    s = re.sub(r"\[([\w.]+)@-1\]", lambda m: f"{prev}{RROW[m.group(1)]}", s)
    s = re.sub(r"\[([\w.]+)@N\]", lambda m: f"$N${RROW[m.group(1)]}", s)
    s = re.sub(r"\[([\w.]+)\]", lambda m: f"{col}{RROW[m.group(1)]}", s)
    s = re.sub(r"\{(\w+)\}", lambda m: "p_" + m.group(1), s)
    s = re.sub(r"#(\w+)", lambda m: f"$J${SROW[m.group(1)]}", s)
    s = re.sub(r"@(\w+)", lambda m: f"INDEX(Data!${FCOL[m.group(1)]}:${FCOL[m.group(1)]},$J${CTY_ROW})", s)
    s = s.replace("YEAR", f"{col}$1" if col else "N$1")
    return "=" + s

for k, desc, unit, t in M.SCALARS:
    r = SROW[k]
    wsM.cell(row=r, column=4, value=desc); wsM.cell(row=r, column=5, value=unit)
    wsM.cell(row=r, column=6, value="Data / Inputs"); wsM.cell(row=r, column=7, value=t.replace("@", "").replace("#", "")[:60])
    wsM.cell(row=r, column=8, value=f"mit.sc.{k}")
    wsM.cell(row=r, column=10, value=xl(t, col="N"))
    wsM.cell(row=r, column=7).fill = fill(CODE); wsM.cell(row=r, column=8).fill = fill(CODE)

row = rr - (rr - row)  # (no-op; rows already assigned)
r0 = max(SROW.values()) + 1
cur = None
rr = r0
secno = 1
for kind, x in order:
    if kind == "SEC":
        band(wsM, rr, x, SECTION, LASTC); rr += 1; secno += 1; continue
    r = RROW[x["code"]]
    wsM.cell(row=r, column=1, value=secno)
    wsM.cell(row=r, column=2, value=x["fuel"]); wsM.cell(row=r, column=3, value=x["sector"])
    wsM.cell(row=r, column=4, value=x["desc"]); wsM.cell(row=r, column=5, value=x["unit"])
    wsM.cell(row=r, column=6, value="CPAT data" if x["hist"] else "Model")
    wsM.cell(row=r, column=8, value=x["code"])
    wsM.cell(row=r, column=7).fill = fill(CODE); wsM.cell(row=r, column=8).fill = fill(CODE)
    if x["hist"]:
        for y, t in zip(M.HIST, x["hist"]):
            wsM[f"{YCOL[y]}{r}"] = xl(t, col=YCOL[y])
    wsM[f"N{r}"] = xl(x["base"], col="N")
    wsM[f"N{r}"].fill = fill(BASE)
    for y in range(2025, 2051):
        wsM[f"{YCOL[y]}{r}"] = xl(x["proj"], col=YCOL[y], prev=YCOL[y - 1])
    if x["code"].startswith("mit.ic."):
        for c in range(4, LASTC + 1): wsM.cell(row=r, column=c).fill = fill(REVIEW)
    rr = r + 1
for rowc in wsM.iter_rows(min_row=3, max_row=rr):
    for c in rowc:
        if c.font != f9b and not (c.fill and c.fill.fgColor and c.fill.fgColor.rgb in ("00" + SECTION, "00" + TITLE)):
            c.font = f9
        if c.column >= 11: c.number_format = "#,##0.000"
for col, w in zip("ABCDEFGHIJ", (4, 5, 5, 44, 9, 10, 22, 22, 6, 11)): wsM.column_dimensions[col].width = w
for y, col in YCOL.items(): wsM.column_dimensions[col].width = 9
wsM.freeze_panes = "K3"
M_LAST = rr

# -----------------------------------------------------------------------------
# RESULTS_ALL (Python, values) and PyKey
# -----------------------------------------------------------------------------
sm = pd.read_csv(os.path.join(M.RES, "nmim_summary_v0_2.csv"))
res = pd.read_csv(os.path.join(M.RES, "nmim_results_v0_2.csv"))
band(wsA, 1, "Results for all countries (Python engine, values), USD bn and MtCO2e", TITLE, 22, fw)
wsA["A2"] = ("Incremental cost is the annual cost in that year of all incremental capital in service (capital charge) "
             "plus the change in O&M, less fuel savings (less air pollution avoided in lens c). Positive = net cost. "
             "PV = present value 2025-2050 at the lens rate.")
wsA["A2"].font = fnote
cols = [("iso", "ISO3"), ("name", "Country"), ("grp", "Group"), ("wacc", "Market WACC"),
        ("gdp_2030", "GDP 2030"), ("inv1_2030", "Investment 2030, baseline"), ("inv2_2030", "Investment 2030, net zero"),
        ("dinv_2030", "Incremental investment 2030"), ("ica_2030", "Incr. cost 2030 (a) market"),
        ("icb_2030", "Incr. cost 2030 (b) 6%"), ("icc_2030", "Incr. cost 2030 (c) 6% + co-ben."),
        ("ghg1_2030", "GHG 2030 baseline"), ("ghg2_2030", "GHG 2030 net zero"), ("ghg2_2050", "GHG 2050 net zero"),
        ("cum_inv2", "Cumulative investment 2025-50, net zero"), ("cum_dinv", "Cumulative incremental 2025-50"),
        ("pv_ic_a", "PV incr. cost (a)"), ("pv_ic_b", "PV incr. cost (b)"), ("pv_ic_c", "PV incr. cost (c)"),
        ("cum_abt", "Cumulative abatement 2025-50"), ("co2_model_2024", "Energy CO2 2024, model"),
        ("co2_cpat_2024", "Energy CO2 2024, CPAT"),
        ("pms", "Power market (1 VIU, 2 single buyer, 3 wholesale)"), ("cr_ecy", "Electricity cost recovery"),
        ("typ1_2030", "2030 net zero: P1 pure public"), ("typ2_2030", "P2 regulated utility"),
        ("typ3_2030", "P3 contracted revenue"), ("typ4_2030", "P4 viability gap"), ("typ5_2030", "P5 pure private"),
        ("bud_2030", "Direct budget need 2030 (P1 + gap of P4)"), ("dbud_2030", "Incremental budget need 2030"),
        ("prv_2030", "Private capital 2030")]
for j, (k, h) in enumerate(cols):
    c = wsA.cell(row=4, column=1 + j, value=h); c.font = fwb; c.fill = fill(TITLE)
    c.alignment = Alignment(wrap_text=True, vertical="center")
wsA.row_dimensions[4].height = 40
sm = sm.sort_values("gdp_2030", ascending=False)
for n, (_, s) in enumerate(sm.iterrows()):
    for j, (k, h) in enumerate(cols):
        v = s[k]
        c = wsA.cell(row=5 + n, column=1 + j, value=(round(float(v), 4) if isinstance(v, float) else v))
        c.font = f9
        if j >= 3: c.number_format = ("0.0%" if k == "wacc" else "0.00" if k == "cr_ecy" else "0" if k == "pms"
                                      else "#,##0;(#,##0);-")
wsA.freeze_panes = "C5"; wsA.auto_filter.ref = f"A4:{CL(len(cols))}{4 + len(sm)}"
for j in range(len(cols)): wsA.column_dimensions[CL(1 + j)].width = 22 if j == 1 else 11

KEYCODES = ["mit.gdp.lvl", "mit.act.ecy.res", "mit.gen.tot.2", "mit.gen.fos.1", "mit.inv.gen.2", "mit.inv.grid.2",
            "mit.inv.trn.2", "mit.inv.bld.2", "mit.inv.ind.2", "mit.inv.tot.1", "mit.inv.tot.2", "mit.ghg.tot.1",
            "mit.ghg.tot.2", "mit.ic.a", "mit.ic.b", "mit.ic.c", "mit.typ.p1.2", "mit.typ.p2.2", "mit.typ.p3.2",
            "mit.typ.p4.2", "mit.typ.p5.2", "mit.typ.bud.2"]
wsK["A1"] = "Python results for the key rows (values; Check sheet reads these)"
wsK.cell(row=2, column=1, value="key"); wsK.cell(row=2, column=2, value="iso"); wsK.cell(row=2, column=3, value="code")
KY = [2024, 2025, 2030, 2040, 2050]
for j, y in enumerate(KY): wsK.cell(row=2, column=4 + j, value=y)
kr = 3
for code in KEYCODES:
    sub = res[res.code.str.endswith("." + code) & (res.code.str.len() == len(code) + 4)]
    for _, x in sub.iterrows():
        wsK.cell(row=kr, column=1, value=f"{x.iso}|{code}"); wsK.cell(row=kr, column=2, value=x.iso)
        wsK.cell(row=kr, column=3, value=code)
        for j, y in enumerate(KY): wsK.cell(row=kr, column=4 + j, value=float(x[str(y)]))
        kr += 1
K_LAST = kr - 1

# -----------------------------------------------------------------------------
# CHECK
# -----------------------------------------------------------------------------
band(wsC, 1, "Check: workbook (live formulas) vs Python engine, selected country", TITLE, 4 + 3 * len(KY), fw)
wsC["A2"] = "Differences should be 0 to rounding. Max abs difference in B3."; wsC["A2"].font = fnote
wsC["A3"] = "Max abs diff"; wsC["A3"].font = f9b
hdrs = ["Code"] + [f"WB {y}" for y in KY] + [f"Py {y}" for y in KY] + [f"Diff {y}" for y in KY]
for j, h in enumerate(hdrs):
    c = wsC.cell(row=5, column=1 + j, value=h); c.font = fwb; c.fill = fill(TITLE)
for n, code in enumerate(KEYCODES):
    r = 6 + n
    wsC.cell(row=r, column=1, value=code).font = f9
    for j, y in enumerate(KY):
        wsC.cell(row=r, column=2 + j, value=f"=Mitigation!{YCOL[y]}{RROW[code]}")
        wsC.cell(row=r, column=2 + len(KY) + j,
                 value=f'=INDEX(PyKey!{CL(4 + j)}$3:{CL(4 + j)}${K_LAST},MATCH(Settings!$C$4&"|{code}",PyKey!$A$3:$A${K_LAST},0))')
        a, b = CL(2 + j), CL(2 + len(KY) + j)
        wsC.cell(row=r, column=2 + 2 * len(KY) + j, value=f"={a}{r}-{b}{r}")
        for cc in (2 + j, 2 + len(KY) + j, 2 + 2 * len(KY) + j):
            wsC.cell(row=r, column=cc).number_format = "#,##0.000"; wsC.cell(row=r, column=cc).font = f9
last = 6 + len(KEYCODES) - 1
d0, d1 = CL(2 + 2 * len(KY)), CL(1 + 3 * len(KY))
wsC["B3"] = f"=MAX(MAX({d0}6:{d1}{last}),-MIN({d0}6:{d1}{last}))"; wsC["B3"].fill = fill(REVIEW)
wsC.column_dimensions["A"].width = 18

# -----------------------------------------------------------------------------
# DASHBOARD (selected country)
# -----------------------------------------------------------------------------
band(wsD, 1, "NMIM - Dashboard (selected country)", TITLE, 9, fw)
wsD["B3"] = "Country"; wsD["C3"] = "=Settings!C5"; wsD["C3"].font = Font(name=F, size=12, bold=True)
DROWS = [("GDP (2024 USD bn)", "mit.gdp.lvl"),
         ("Investment, baseline", "mit.inv.tot.1"), ("Investment, net zero", "mit.inv.tot.2"),
         ("Incremental investment", "mit.dinv.tot"),
         ("  of which power generation", "mit.dinv.gen"), ("  grids and storage", "mit.dinv.grid"),
         ("  transport", "mit.dinv.trn"), ("  buildings", "mit.dinv.bld"), ("  industry", "mit.dinv.ind"),
         ("  non-CO2 and land use", "mit.dinv.nco"),
         ("Incremental cost (a) market WACC", "mit.ic.a"), ("Incremental cost (b) World Bank 6%", "mit.ic.b"),
         ("Incremental cost (c) 6% + local co-benefits", "mit.ic.c"),
         ("GHG, baseline (MtCO2e)", "mit.ghg.tot.1"), ("GHG, net zero (MtCO2e)", "mit.ghg.tot.2"),
         ("Net zero investment by type (see Typology)", "mit.inv.tot.2")] + [
         (f"  {nm}", f"mit.typ.p{k}.2") for k, nm in M.TYPES] + [
         ("Direct budget need (P1 + gap of P4)", "mit.typ.bud.2")]
DY = [2025, 2030, 2035, 2040, 2045, 2050]
wsD["B5"] = "USD bn a year unless stated"; wsD["B5"].font = fnote
for j, y in enumerate(DY):
    c = wsD.cell(row=6, column=3 + j, value=y); c.font = fwb; c.fill = fill(TITLE)
wsD.cell(row=6, column=2, value="Item").font = fwb; wsD.cell(row=6, column=2).fill = fill(TITLE)
for n, (lab, code) in enumerate(DROWS):
    r = 7 + n
    wsD.cell(row=r, column=2, value=lab).font = f9b if not lab.startswith("  ") else f9
    for j, y in enumerate(DY):
        c = wsD.cell(row=r, column=3 + j, value=f"=Mitigation!{YCOL[y]}{RROW[code]}")
        c.number_format = "#,##0.0;(#,##0.0);-"; c.font = f9
        if code.startswith("mit.ic."): c.fill = fill(REVIEW)
wsD.column_dimensions["B"].width = 40
ch = LineChart(); ch.title = "GHG emissions, baseline and net zero (MtCO2e)"; ch.height, ch.width = 7.5, 16
for code, nm in (("mit.ghg.tot.1", "Baseline"), ("mit.ghg.tot.2", "Net zero")):
    ref = Reference(wsM, min_col=14, max_col=40, min_row=RROW[code], max_row=RROW[code])
    ch.add_data(ref, from_rows=True, titles_from_data=False)
ch.series[0].tx = SeriesLabel(v="Baseline"); ch.series[1].tx = SeriesLabel(v="Net zero")
ch.set_categories(Reference(wsM, min_col=14, max_col=40, min_row=1, max_row=1))
wsD.add_chart(ch, "B31")

# -----------------------------------------------------------------------------
# POWERMARKET (data, values)
# -----------------------------------------------------------------------------
pm = pd.read_csv(os.path.join(M.DATA, "power_market.csv"))
PMH = [("iso", "ISO3"), ("pms", "Market structure (1 VIU, 2 single buyer, 3 wholesale)"), ("pms_src", "Source"),
       ("pms_note", "Note"), ("reg_indep", "Independent regulator"), ("unbundled", "Generation unbundled"),
       ("privatised", "Privatisation"), ("cost_reflective_psrt", "Cost-reflective tariffs (PSRT)"),
       ("cr_ecy", "Cost recovery: retail price / supply cost (CPAT)"), ("ecy_sub_pct_gdp", "Explicit electricity subsidy, % GDP"),
       ("ffs_exp_pct_gdp", "Explicit fossil subsidies, % GDP"), ("ppi_elec_bn_yr", "PPI private electricity, USD bn a year 2015-24"),
       ("ppi_elec_pct_gdp", "PPI private electricity, % GDP"), ("ppi_gen_n_1524", "PPI generation projects 2015-24"),
       ("ppi_vre_share", "Solar and wind share of PPI generation"), ("ppi_merchant_share", "Merchant share of PPI generation"),
       ("ppi_private_dist", "Private distribution (PPI)"), ("ppi_private_trans", "Private transmission (PPI)"),
       ("rise_re", "RISE renewable energy score"), ("rise_lpf", "RISE levelling the playing field")]
band(wsP, 1, "PowerMarket: power sector structure and finance by country (data/power_market.csv)", TITLE, len(PMH), fw)
wsP["A2"] = ("Structure: Power Sector Reform Tracker 2013 (Urpelainen and Yang; 3 if wholesale market, 2 if IPPs, else 1), "
             "curated for countries PSRT does not cover and where it is out of date, raised to 2 where World Bank PPI shows "
             "two or more private generation projects 2014-2024. Cost recovery from the IMF CPAT 2024 panel. RISE 2023.")
wsP["A2"].font = fnote
for j, (k, h) in enumerate(PMH):
    c = wsP.cell(row=3, column=1 + j, value=h); c.font = fwb; c.fill = fill(TITLE)
    c.alignment = Alignment(wrap_text=True, vertical="center")
wsP.row_dimensions[3].height = 52
for n, x in enumerate(pm.itertuples()):
    for j, (k, h) in enumerate(PMH):
        v = getattr(x, k)
        c = wsP.cell(row=4 + n, column=1 + j, value=None if (isinstance(v, float) and pd.isna(v)) else
                     (round(float(v), 4) if isinstance(v, float) else v))
        c.font = f9
PM_LAST = 3 + len(pm)
PMC = {k: CL(1 + j) for j, (k, h) in enumerate(PMH)}
wsP.freeze_panes = "B4"; wsP.auto_filter.ref = f"A3:{CL(len(PMH))}{PM_LAST}"
for j in range(len(PMH)): wsP.column_dimensions[CL(1 + j)].width = 40 if PMH[j][0] == "pms_note" else 12

# -----------------------------------------------------------------------------
# TYPOLOGY (selected country, live formulas)
# -----------------------------------------------------------------------------
band(wsY, 1, "Typology - who pays for net zero investment (selected country)", TITLE, 10, fw)
wsY["B2"] = "=Settings!C5"; wsY["B2"].font = Font(name=F, size=12, bold=True)
wsY["B3"] = ("P1 pure public (taxes repay it) | P2 regulated utility (tariffs on a regulated asset base) | P3 contracted "
             "revenue (PPA, CfD, concession) | P4 viability gap (market revenue + public top-up) | P5 pure private (own "
             "savings or sales)"); wsY["B3"].font = fnote
pmlook = lambda k: f'=INDEX(PowerMarket!${PMC[k]}$4:${PMC[k]}${PM_LAST},MATCH(Settings!$C$4,PowerMarket!$A$4:$A${PM_LAST},0))'
FACTS = [("Power market structure (1 VIU, 2 single buyer, 3 wholesale)", pmlook("pms"), "0"),
         ("Source", pmlook("pms_src") + '&""', "@"), ("Note", pmlook("pms_note") + '&""', "@"),
         ("Independent regulator / unbundled / privatised (0/1, PSRT)",
          "=" + "&\" / \"&".join(pmlook(k)[1:] for k in ("reg_indep", "unbundled", "privatised")), "@"),
         ("Electricity cost recovery (retail price / supply cost)", pmlook("cr_ecy"), "0.00"),
         ("Private electricity investment, PPI, USD bn a year 2015-24", pmlook("ppi_elec_bn_yr"), "#,##0.00"),
         ("RISE renewable energy score (0-100)", pmlook("rise_re"), "0"),
         ("Utility pressure index (tariff under-recovery)", f"=Mitigation!$J${SROW['ty_u']}", "0.00"),
         ("Capital-cost pressure index (WACC above 6%)", f"=Mitigation!$J${SROW['ty_h']}", "0.00"),
         ("Subsidy pressure index (explicit fossil subsidies)", f"=Mitigation!$J${SROW['ty_v']}", "0.00"),
         ("Top-up share of market end-use investment", f"=Mitigation!$J${SROW['ty_m']}", "0.00")]
band(wsY, 5, "Country context", SECTION, 10)
for n, (lab, fml, fmt) in enumerate(FACTS):
    r = 6 + n
    wsY.cell(row=r, column=2, value=lab).font = f9
    c = wsY.cell(row=r, column=3, value=fml); c.font = f9; c.number_format = fmt
r = 6 + len(FACTS) + 1
band(wsY, r, "Type shares by sector (Mitigation section 1)", SECTION, 10); r += 1
for j, h in enumerate(["Sector", "P1", "P2", "P3", "P4", "P5", "Sum"]):
    c = wsY.cell(row=r, column=2 + j, value=h); c.font = fwb; c.fill = fill(TITLE)
r += 1
J = lambda k: f"Mitigation!$J${SROW[k]}"
SHR = [("Power generation", [f"={J(f'ty_gen_p{k}')}" for k in range(1, 6)]),
       ("Grids and storage", [f"={J(f'ty_grid_p{k}')}" for k in range(1, 6)]),
       ("Transport (EV premium, chargers)", [f"={J(f'ty_trn_p{k}')}" for k in range(1, 6)]),
       ("Buildings", [f"={J(f'ty_bld_p{k}')}" for k in range(1, 6)]),
       ("Industry: fuel switch", ["0", "0", "0", f"={J('ty_ih_p4')}", f"=1-{J('ty_ih_p4')}"]),
       ("Industry: CCS on CBAM goods", ["0", "0", "0", "1", "0"]),
       ("Non-CO2", [f"={J('ty_nc_p1')}", "0", "0", f"=1-{J('ty_nc_p1')}-{J('ty_nc_p5')}", f"={J('ty_nc_p5')}"]),
       ("Land use", [f"={J('ty_lu_p1')}", "0", "0", f"=1-{J('ty_lu_p1')}", "0"]),
       ("Other energy CO2 (aviation, shipping)", ["0", "0", "0", "1", "0"]),
       ("Unmodelled clean investment", [f"={J(f'ty_cal_p{k}')}" for k in range(1, 6)])]
for lab, cells in SHR:
    wsY.cell(row=r, column=2, value=lab).font = f9
    for j, v in enumerate(cells):
        c = wsY.cell(row=r, column=3 + j, value=v if v.startswith("=") else float(v)); c.number_format = "0%"; c.font = f9
    c = wsY.cell(row=r, column=8, value=f"=SUM(C{r}:G{r})"); c.number_format = "0%"; c.font = f9
    r += 1
r += 1
band(wsY, r, "Investment by type, USD bn a year (Mitigation section 12)", SECTION, 10); r += 1
TY = [2025, 2030, 2035, 2040, 2045, 2050]
for j, h in enumerate(["Item"] + TY):
    c = wsY.cell(row=r, column=2 + j, value=h); c.font = fwb; c.fill = fill(TITLE)
r += 1
TROWS = ([(f"Net zero: {nm}", f"mit.typ.p{k}.2") for k, nm in M.TYPES] +
         [("Net zero: total", "mit.inv.tot.2")] +
         [(f"Incremental: {nm}", f"mit.dtyp.p{k}") for k, nm in M.TYPES] +
         [("Direct budget need, net zero (P1 + gap of P4)", "mit.typ.bud.2"),
          ("Direct budget need, incremental", "mit.dtyp.bud"),
          ("Revenue resting on a public tariff or buyer (P2 + P3)", "mit.typ.cnt.2"),
          ("Private capital (P3 + P5 + private part of P4)", "mit.typ.prv.2")])
TY0 = r
for lab, code in TROWS:
    wsY.cell(row=r, column=2, value=lab).font = f9b if "total" in lab or "budget" in lab else f9
    for j, y in enumerate(TY):
        c = wsY.cell(row=r, column=3 + j, value=f"=Mitigation!{YCOL[y]}{RROW[code]}")
        c.number_format = "#,##0.0;(#,##0.0);-"; c.font = f9
        if "budget" in lab: c.fill = fill(REVIEW)
    r += 1
wsY.column_dimensions["B"].width = 56; wsY.column_dimensions["C"].width = 26
for col in "DEFGH": wsY.column_dimensions[col].width = 10
from openpyxl.chart import BarChart
bc = BarChart(); bc.type = "col"; bc.grouping = "stacked"; bc.overlap = 100
bc.title = "Net zero investment by type (USD bn)"; bc.height, bc.width = 7.5, 16
for k in range(5):
    bc.add_data(Reference(wsY, min_col=3, max_col=8, min_row=TY0 + k, max_row=TY0 + k), from_rows=True, titles_from_data=False)
    bc.series[k].tx = SeriesLabel(v=M.TYPES[k][1])
bc.set_categories(Reference(wsY, min_col=3, max_col=8, min_row=TY0 - 1, max_row=TY0 - 1))
wsY.add_chart(bc, f"J5")

# -----------------------------------------------------------------------------
# README
# -----------------------------------------------------------------------------
band(wsR, 1, f"NMIM v{M.VERSION} - National Mitigation Investment Model (simplified)", TITLE, 4, fw)
wsR.column_dimensions["B"].width = 130
TXT = [
 ("What it is", "A simplified national model of net zero mitigation investment for 186 countries, annual 2024-2050, "
  "built to CPAT norms. For each country it projects baseline (current trends) and net zero energy use, power "
  "capacity, investment by sector, emissions, fuel bills and air pollution, and the incremental cost of the net zero "
  "path under three lenses: (a) market WACC, (b) World Bank 6%, (c) 6% with local co-benefits."),
 ("How to use", "Choose a country on Settings!C4. Mitigation recalculates in live formulas (CPAT layout: K = 2021 ... "
  "N = 2024 base year ... AN = 2050). Dashboard summarises it. Inputs holds every parameter (MTInputs column "
  "structure; edit column G). Results_All holds every country from the Python engine; Check compares the two."),
 ("Method", "See NMIM_Method_v0_2.md. Baseline end use follows CPAT's activity equation with CPAT income-group "
  "elasticities and prices fixed in real terms; net zero applies sector transformations (EV sales, heat "
  "electrification, industrial abatement, fossil-free power by group target year); power capacity is a stock model "
  "with lives and learning; clean investment not modelled bottom-up is calibrated to IEA 2024 by group and added to "
  "both scenarios; incremental cost sums capital charges on incremental capital in service."),
 ("Consistency", "Codes and data follow CPAT (IMF CPAT 2024 panel; mit.<indicator>.<fuel>.<sector>.<scenario>). "
  "CBAM goods process CO2 comes from the CBAM Policy Tool's Country data. Market WACC is the v1.1 investment-needs "
  "workbook's (observed, Damodaran or IHLEG default)."),
 ("Colours", "Title 00B050; section 92D050; inputs EBF1DE (blue text); codes DCE6F1; base year DDD9C4; "
  "results to review FFF2CC."),
 ("Typology", "v0.2 splits every investment line into five types by who pays (P1 pure public, P2 regulated "
  "utility, P3 contracted revenue, P4 viability gap, P5 pure private). Shares depend on the country's power market "
  "structure (PowerMarket sheet: PSRT, PPI, curated), electricity cost recovery (CPAT), market WACC, explicit fossil "
  "subsidies and income group (Mitigation section 1 and 12; Typology sheet). Direct budget need = P1 + the public "
  "gap share of P4."),
 ("Caveats", "Simplified: no price response, no power optimisation, car-equivalent vehicles, assumed transformation "
  "dates; several formulas exceed the 100-character norm (to be split). Parameters marked † are assumptions."),
 ("Version log", f"v0.2 | 2026-10-20 | Investment typology (who pays) by country and sector; power market data "
  "(PowerMarket); industry and non-CO2 spending split into their parts; results otherwise unchanged from v0.1.\n"
  "v0.1 | 2026-10-20 | First build: 186 countries, 11 end-use rows, power stock model, IEA 2024 calibration of "
  "investment levels; three lenses; workbook = Python on key rows (Check)."),
]
r = 3
for h, t in TXT:
    wsR.cell(row=r, column=1, value=h).font = f9b
    c = wsR.cell(row=r, column=2, value=t); c.font = f9; c.alignment = wrap
    wsR.row_dimensions[r].height = max(14, 13 * (len(t) // 150 + 1)); r += 1
wsR.column_dimensions["A"].width = 14

for ws in (wsR, wsD, wsS, wsI, wsM, wsA, wsY, wsC, wsP):
    ws.sheet_view.zoomScale = 90
wb.save(OUT)
print("wrote", OUT, "Mitigation rows", M_LAST, "data fields", len(fields))
