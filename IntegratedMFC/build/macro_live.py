"""Live simple macro model for the Integrated MF-C tool: DIGNAD-R for the country picked on the
Dashboard, in formulas.

DIGNAD-R is the recursive, formula-only DIGNAD-lite / DIGNAR-lite of
MacroFiscalModels/ExcelDIGNAD/build_dignad_excel.py (equations in ExcelDIGNAD/README.md). Its
assumptions, augmented-input defaults, input loader and Python copy of the equations are imported
from there, so the two stay one model; only the layout differs: DIGNAD-R Excel computes 12
countries at once, this sheet computes the selected one, for any of the integrated run's countries.

Sheets added to the built workbook:
  Macro              view: the integrated run (pre-computed, Python) beside the live simple model
  Macro_assumptions  named yellow cells p_... (change them; everything recalculates)
  Macro_model        calibration and the 2024-2050 recursion, baseline (B) and programme (P)
  Macro_inputs       one row per country: dashboard inputs, augmented defaults, NMIM programme and
                     pricing paths, and the integrated run's GDP and debt paths

    python3 macro_live.py --check ISO3 ...   recompute the given countries in Python (same equations)
"""
import os, sys
import numpy as np, pandas as pd, openpyxl
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
MFM = os.path.join(ROOT, "MacroFiscalModels")
sys.path.insert(0, os.path.join(MFM, "ExcelDIGNAD"))
import build_dignad_excel as dr                                 # noqa: E402  (the DIGNAD-R model)

Y0, Y1, T = dr.Y0, dr.Y1, dr.T
YEARS = list(range(Y0, Y1 + 1))
F = "Arial"; INK = "0E1B2A"; ACC = "C05621"; MUT = "5A6B7B"
HDR = PatternFill("solid", fgColor=INK); YEL = PatternFill("solid", fgColor="FFF2CC")
SEC = PatternFill("solid", fgColor="EAE3F2"); AUG = PatternFill("solid", fgColor="E3EAF2")
thin = Side(style="thin", color="C8D0D6"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def font(size=10, bold=False, color=INK, italic=False):
    return Font(name=F, size=size, bold=bold, color=color, italic=italic)


def load(countries=None):
    """DIGNAD-R inputs for every country of the integrated run that the dashboard has."""
    ir = pd.read_csv(os.path.join(MFM, "results", "integrated_results.csv")).set_index("iso3")
    panel = pd.read_csv(dr.PANEL).set_index("iso3")
    isos = [i for i in ir.index if i in panel.index and panel.income.get(i) in dr.INCOME]
    dr.COUNTRIES = countries or isos
    wide, years, long = dr.load_inputs()
    return wide, long, ir


def add(path, panel_long):
    wide, long, ir = load()
    names = panel_long.drop_duplicates("iso3").set_index("iso3").country   # the Dashboard's names
    wide["country"] = wide.iso3.map(names).fillna(wide.country)
    pa = pd.read_csv(os.path.join(MFM, "results", "integrated_paths.csv")).set_index(["iso3", "series"])

    wb = openpyxl.load_workbook(path)

    def name(n, ref):
        wb.defined_names[n] = DefinedName(n, attr_text=ref)

    # ---- Macro_inputs: one row per country ------------------------------------------------
    wi = wb.create_sheet("Macro_inputs")
    dcodes = [c for c, *_ in dr.DASH]; acodes = [c for c, *_ in dr.AUGD]
    paths = [("cap", "pub_cap", "Budget need, capital (% GDP)"),
             ("spend", "pub_spend", "Budget need, spending (% GDP)"),
             ("rev", "rev", "Pricing revenue (% GDP)"),
             ("pyg", "gdp_dev_pct", "Integrated run: GDP vs today, programme (%)"),
             ("pygb", "gdp_dev_base_pct", "Integrated run: GDP vs today, baseline (%)"),
             ("pyd", "debt_pct_gdp", "Integrated run: debt, programme (% GDP)"),
             ("pydb", "debt_base_pct_gdp", "Integrated run: debt, baseline (% GDP)")]
    cols = (["iso3", "country", "income", "model", "model_reason", "status"] + dcodes + acodes +
            [f"{p}_{y}" for p, *_ in paths for y in YEARS])
    for j, c in enumerate(cols, 1):
        x = wi.cell(1, j, c); x.font = font(8, True, "FFFFFF"); x.fill = HDR
    lv = long.set_index(["iso3", "indicator", "year"]).value
    for i, row in enumerate(wide.itertuples(index=False), 2):
        rd = row._asdict(); iso = rd["iso3"]
        vals = [iso, rd["country"], rd["income"], ir.model.get(iso), ir.model_reason.get(iso), ir.status.get(iso)]
        vals += [rd[c] for c in dcodes + acodes]
        for p, src, _ in paths:
            for y in YEARS:
                if p in ("cap", "spend", "rev"):                    # as DIGNAD-R: missing year = 0
                    v = 0.0 if y == Y0 else float(lv.get((iso, src, y), 0.0))
                elif y == Y0:                                       # integrated run starts from 2024
                    v = 0.0 if p in ("pyg", "pygb") else ir.debt_2024_pct_gdp.get(iso)
                else:
                    v = pa[str(y)].get((iso, src)) if (iso, src) in pa.index else None
                vals.append(v)
        for j, v in enumerate(vals, 1):
            v = None if (isinstance(v, float) and np.isnan(v)) else v
            wi.cell(i, j, v).font = font(8)
    NI = len(wide) + 1
    IC = {c: get_column_letter(j) for j, c in enumerate(cols, 1)}
    wi.freeze_panes = "C2"; wi.sheet_properties.tabColor = "BFBFBF"
    rng = lambda c: f"Macro_inputs!${IC[c]}$2:${IC[c]}${NI}"

    # ---- Macro_assumptions ----------------------------------------------------------------
    wa = wb.create_sheet("Macro_assumptions")
    for j, (h, w) in enumerate([("Name", 18), ("Value", 10), ("Description", 72), ("Source", 60)], 1):
        x = wa.cell(1, j, h); x.font = font(9, True, "FFFFFF"); x.fill = HDR
        wa.column_dimensions[get_column_letter(j)].width = w
    for i, (n, v, lab, src) in enumerate(dr.PARAMS, 2):
        wa.cell(i, 1, n).font = font(9, True)
        c = wa.cell(i, 2, v); c.fill = YEL; c.font = font(10, True, "0000FF"); c.border = BOX
        wa.cell(i, 3, lab).font = font(9); wa.cell(i, 4, src).font = font(9, color=MUT)
        name(n, f"Macro_assumptions!$B${i}")
    wa.freeze_panes = "A2"

    # ---- Macro_model: inputs and calibration for the selected country -------------------
    wm = wb.create_sheet("Macro_model")
    wm.column_dimensions["A"].width = 30; wm.column_dimensions["B"].width = 12
    wm["A1"] = "DIGNAD-R for the country picked on the Dashboard (recursive DIGNAD-lite / DIGNAR-lite)"
    wm["A1"].font = font(11, True)
    r = 3
    def cell(label, nm, formula, fmt="0.0000"):
        nonlocal r
        wm.cell(r, 1, label).font = font(9)
        x = wm.cell(r, 2, formula); x.font = font(9); x.number_format = fmt
        name(nm, f"Macro_model!$B${r}")
        r += 1
    cell("Row of the country in Macro_inputs", "m_row",
         f"=MATCH(Dashboard!$C$4,{rng('country')},0)", "0")
    # INDEX on an empty cell returns 0, not blank: test for blank first (the dashboard's rule)
    for c in dcodes + acodes:
        ix = f"INDEX({rng(c)},m_row)"
        cell(f"input: {c}", f"i_{c}", f'=IF({ix}="","",{ix})')
    I = lambda c: f"i_{c}"
    cal = [
        ("k", "income group (1 LIC ... 4 HIC)",
         f'=MATCH(INDEX({rng("income")},m_row),{{"Low income","Lower middle income","Upper middle income","High income"}},0)'),
        ("s", "public investment efficiency",
         f'=1-IF(ISNUMBER({I("pefa_pi11_pim_score")}),0.4-({I("pefa_pi11_pim_score")}-1)/3*0.27,CHOOSE(c_k,0.4,0.27,0.27,0.13))'),
        ("Rg", "return on public capital", f"={I('return_public_capital_pct')}/100"),
        ("ig", "public investment, start", f"={I('public_investment_pct_gdp')}/100"),
        ("tax", "tax revenue", f"={I('tax_revenue_pct_gdp')}/100"),
        ("grants0", "grants, start", f"={I('grants_pct_gdp')}/100"),
        ("conc", "concessional share of debt",
         f'=IF(ISNUMBER({I("ids_multilateral_share_pct")}),{I("ids_multilateral_share_pct")}/100,CHOOSE(c_k,0.7,0.45,0.15,0))'),
        ("rstar", "real commercial rate",
         f'=IF(ISNUMBER({I("ids_private_creditor_rate_pct")}),MIN(MAX({I("ids_private_creditor_rate_pct")}/100-0.03,0.03),0.08),IF(c_k<3,0.04,0.03))'),
        ("rr", "fossil rents, share of GDP",
         f'=IF(ISNUMBER({I("fossil_rents_total_pct_gdp")}),{I("fossil_rents_total_pct_gdp")}/100,0)'),
        ("res", "resource block on (1 = DIGNAR)", "=IF(c_rr>=p_res_threshold,1,0)"),
        ("rents0", "rents, share of non-resource GDP", "=IF(c_res=1,c_rr/(1-c_rr),0)"),
        ("K0", "private capital, start", "=p_alpha*(1-c_tax)/(c_rstar+p_delta)"),
        ("KG0", "public capital, start", "=c_s*c_ig/p_delta_g"),
        ("psi", "output elasticity of public capital", "=c_Rg*c_KG0"),
        ("A", "productivity (Y 2024 = 1)", "=1/(c_KG0^c_psi*c_K0^p_alpha)"),
        ("GDP0", "GDP, start", "=1+c_rents0"),
        ("D0", "public debt, start", f"={I('gross_debt_pct_gdp')}/100*c_GDP0"),
        ("Dk0", "concessional debt, start", "=c_conc*c_D0"),
        ("int0", "interest, start", "=p_rk*c_Dk0+c_rstar*(c_D0-c_Dk0)"),
        ("rev0", "revenue, start", "=c_tax+p_gov_take*c_rents0+c_grants0"),
        ("g", "government consumption", "=MAX(0.05,c_rev0-c_ig-c_int0-p_tr_target)"),
        ("trbar", "transfers, start (balances 2024)", "=c_rev0-c_g-c_ig-c_int0"),
        ("lam_f", "fiscal rule speed",
         f"=IF(ISNUMBER({I('fiscal_rule_speed')}),{I('fiscal_rule_speed')},IF(c_res=1,p_lambda_f_res,p_lambda_f))"),
        ("ratio0", "debt ratio, start", "=c_D0/c_GDP0"),
    ]
    r += 1
    wm.cell(r, 1, "Calibration (as DIGNAD-R)").font = font(10, True); r += 1
    for k, lab, f in cal:
        cell(f"{k}: {lab}", f"c_{k}", f)

    # ---- Macro_model: the recursion -----------------------------------------------------
    r += 2
    H0 = r
    BASE = ["year", "t", "dis", "tfp", "rents", "cap", "spend", "rev"]
    SV = ["KG", "K", "Y", "Kstar", "GDP", "P", "grants", "revenue", "tr", "prim", "interest", "D", "Dk", "ratio", "rc", "chtm"]
    DEV = ["gdp_dev_prog", "gdp_vs_2024_P", "gdp_vs_2024_B", "debt_diff_pp", "chtm_dev_prog", "tr_diff"]
    CH = ["ch_gdp_P", "ch_gdp_B", "ch_gdp_py", "ch_gdp_pyb", "ch_debt_P", "ch_debt_B", "ch_debt_py", "ch_debt_pyb",
          "ch_tr_P", "ch_tr_B", "ch_prog", "ch_rev", "ch_rents"]
    CHL = {"ch_gdp_P": "Simple model, programme", "ch_gdp_B": "Simple model, baseline",
           "ch_gdp_py": "Integrated run, programme", "ch_gdp_pyb": "Integrated run, baseline",
           "ch_debt_P": "Simple model, programme", "ch_debt_B": "Simple model, baseline",
           "ch_debt_py": "Integrated run, programme", "ch_debt_pyb": "Integrated run, baseline",
           "ch_tr_P": "Transfers, programme", "ch_tr_B": "Transfers, baseline",
           "ch_prog": "Programme spending", "ch_rev": "Pricing revenue", "ch_rents": "Fossil rents"}
    MC = BASE + [f"B_{v}" for v in SV] + [f"P_{v}" for v in SV] + DEV + CH
    M = {c: get_column_letter(j + 2) for j, c in enumerate(MC)}         # model starts in column B
    for c in MC:
        x = wm[f"{M[c]}{H0}"]; x.value = CHL.get(c, c); x.font = font(8, True, "FFFFFF"); x.fill = HDR
        x.alignment = Alignment(wrap_text=True)
    wm.cell(H0, 1, "Model, 2024-2050 (B = baseline, P = programme)").font = font(9, True)
    R0 = H0 + 1
    py = lambda p, rr: f'INDEX(Macro_inputs!${IC[p + "_" + str(Y0)]}$2:${IC[p + "_" + str(Y1)]}${NI},m_row,$C{rr}+1)'
    for t in range(T):
        rr, p = R0 + t, R0 + t - 1
        v = {"year": Y0 + t, "t": t,
             "dis": f"=IF(AND(p_dis_year>0,$B{rr}=p_dis_year),1,0)",
             "tfp": f"=IF(AND(p_dis_year>0,$B{rr}>=p_dis_year),1-p_eps_a*p_rho_a^($B{rr}-p_dis_year),1)",
             "rents": f"=c_rents0*(1-p_rent_decline*MIN(1,$C{rr}/{Y1 - Y0}))",
             "cap": f"={py('cap', rr)}/100", "spend": f"={py('spend', rr)}/100", "rev": f"={py('rev', rr)}/100"}
        for X, on in (("B", "0"), ("P", "p_programme")):
            c = lambda var: f"{M[X + '_' + var]}{rr}"
            pr = lambda var: f"{M[X + '_' + var]}{p}"
            if t == 0:
                v.update({f"{X}_KG": "=c_KG0", f"{X}_K": "=c_K0", f"{X}_Y": "=1", f"{X}_Kstar": "=c_K0",
                          f"{X}_GDP": "=c_GDP0", f"{X}_P": "=0", f"{X}_grants": "=c_grants0",
                          f"{X}_revenue": "=c_rev0", f"{X}_tr": "=c_trbar", f"{X}_prim": "=-c_int0",
                          f"{X}_interest": "=c_int0", f"{X}_D": "=c_D0", f"{X}_Dk": "=c_Dk0",
                          f"{X}_ratio": f"={c('D')}/{c('GDP')}", f"{X}_rc": "=c_rstar",
                          f"{X}_chtm": f"=(1-c_tax)*(1-p_alpha)*{c('Y')}+{c('tr')}"})
                continue
            KGd = f"{pr('KG')}*(1-p_eps_g*${M['dis']}{rr})"; Kd = f"{pr('K')}*(1-p_eps_k*${M['dis']}{rr})"
            v.update({
                f"{X}_Y": f"=c_A*${M['tfp']}{rr}*({KGd})^c_psi*({Kd})^p_alpha",
                f"{X}_Kstar": f"=p_alpha*(1-c_tax)*{c('Y')}/(c_rstar+p_delta)",
                f"{X}_K": f"={Kd}+p_lambda_k*({c('Kstar')}-{Kd})",
                f"{X}_P": f"={on}*(${M['cap']}{rr}+${M['spend']}{rr})",
                f"{X}_KG": f"=(1-p_delta_g)*{KGd}+c_s*(c_ig+{on}*${M['cap']}{rr})",
                f"{X}_GDP": f"={c('Y')}+${M['rents']}{rr}",
                f"{X}_grants": f"=c_grants0+p_sh_grant*{c('P')}",
                f"{X}_revenue": f"=c_tax*{c('Y')}+p_gov_take*${M['rents']}{rr}+{c('grants')}+{on}*p_pricing*${M['rev']}{rr}",
                f"{X}_tr": f"=c_trbar-c_lam_f*({pr('D')}/{pr('GDP')}-c_ratio0)",
                f"{X}_prim": f"=c_g+c_ig+{c('P')}+{c('tr')}-{c('revenue')}",
                f"{X}_interest": f"=p_rk*{pr('Dk')}+{pr('rc')}*({pr('D')}-{pr('Dk')})",
                f"{X}_D": f"={pr('D')}+{c('interest')}+{c('prim')}",
                f"{X}_Dk": f"=c_Dk0+({pr('Dk')}-c_Dk0)*(1-p_amort)+p_sh_conc*{c('P')}",
                f"{X}_ratio": f"={c('D')}/{c('GDP')}",
                f"{X}_rc": f"=c_rstar+p_eta*MAX(0,{c('ratio')}-c_ratio0)",
                f"{X}_chtm": f"=(1-c_tax)*(1-p_alpha)*{c('Y')}+{c('tr')}"})
        m = lambda col, row=rr: f"{M[col]}{row}"
        v.update({
            "gdp_dev_prog": f"=({m('P_GDP')}/{m('B_GDP')}-1)*100",
            "gdp_vs_2024_P": f"=({m('P_GDP')}/{m('P_GDP', R0)}-1)*100",
            "gdp_vs_2024_B": f"=({m('B_GDP')}/{m('B_GDP', R0)}-1)*100",
            "debt_diff_pp": f"=({m('P_ratio')}-{m('B_ratio')})*100",
            "chtm_dev_prog": f"=({m('P_chtm')}/{m('B_chtm')}-1)*100",
            "tr_diff": f"=({m('P_tr')}-{m('B_tr')})*100",
            # chart series: GDP as 2024 = 100; the integrated run's 'today's path' starts at 100 too
            "ch_gdp_P": f"={m('P_GDP')}/{m('P_GDP', R0)}*100", "ch_gdp_B": f"={m('B_GDP')}/{m('B_GDP', R0)}*100",
            "ch_gdp_py": f'=IFERROR(IF({py("pyg", rr)}="",NA(),100+{py("pyg", rr)}),NA())',
            "ch_gdp_pyb": f'=IFERROR(IF({py("pygb", rr)}="",NA(),100+{py("pygb", rr)}),NA())',
            "ch_debt_P": f"={m('P_ratio')}*100", "ch_debt_B": f"={m('B_ratio')}*100",
            "ch_debt_py": f'=IFERROR(IF({py("pyd", rr)}="",NA(),{py("pyd", rr)}),NA())',
            "ch_debt_pyb": f'=IFERROR(IF({py("pydb", rr)}="",NA(),{py("pydb", rr)}),NA())',
            "ch_tr_P": f"={m('P_tr')}*100", "ch_tr_B": f"={m('B_tr')}*100",
            "ch_prog": f"={m('P_P')}*100", "ch_rev": f"=p_programme*p_pricing*{m('rev')}*100",
            "ch_rents": f"={m('rents')}*100"})
        for col in MC:
            x = wm[f"{M[col]}{rr}"]; x.value = v[col]; x.font = font(8)
            if col not in ("year", "t", "dis"):
                x.number_format = "0.0000"
    RL = R0 + T - 1
    wm.freeze_panes = wm[f"C{R0}"]
    wm.sheet_properties.tabColor = "BFBFBF"
    at = lambda col, yr: f"Macro_model!${M[col]}${R0 + yr - Y0}"

    # ---- Macro: the view ------------------------------------------------------------------
    wv = wb.create_sheet("Macro", wb.sheetnames.index("Dashboard") + 1)
    wv.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGH", [2, 52, 12, 12, 3, 12, 12, 2]):
        wv.column_dimensions[col].width = w
    wv["B2"] = "Macro-fiscal: what the programme does to growth and debt"; wv["B2"].font = font(16, True)
    wv["B3"] = ("Two models side by side. The integrated run is pre-computed in Python (MacroFiscalModels): "
                "GIMF-lite, DIGNAD-lite or DIGNAR-lite by country, forward-looking. The simple model is "
                "DIGNAD-R, the same DIGNAD-lite / DIGNAR-lite equations made recursive and written in formulas: "
                "change the yellow cells on Macro_assumptions and it recalculates. Illustrative calibration.")
    wv["B3"].font = font(9, italic=True, color=MUT); wv["B3"].alignment = Alignment(wrap_text=True, vertical="top")
    wv.merge_cells("B3:G3"); wv.row_dimensions[3].height = 52
    wv["B5"] = "Country (pick it on the Dashboard sheet)"; wv["B5"].font = font(10, True)
    wv["C5"] = "=Dashboard!$C$4"; wv["C5"].font = font(12, True, ACC); wv.merge_cells("C5:G5")
    wv["B6"] = ('=IF(ISNA(m_row),"Not covered: no NMIM / integrated-run data for this country.",'
                '"Integrated run used "&INDEX(' + rng("model") + ',m_row)&" ("&INDEX(' + rng("model_reason") +
                ',m_row)&"); status: "&INDEX(' + rng("status") + ',m_row)&".")')
    wv["B6"].font = font(9, True, MUT); wv.merge_cells("B6:G6")
    r = 8
    for c, txt in zip("BCDFG", ["Results", "Integrated run", "", "Simple model (live)", ""]):
        x = wv[f"{c}{r}"]; x.value = txt; x.font = font(10, True)
    for c in "BCDEFG":
        wv[f"{c}{r}"].fill = SEC
    r += 1
    for c, txt in zip("CDFG", ["2030", "2050", "2030", "2050"]):
        x = wv[f"{c}{r}"]; x.value = txt; x.font = font(9, True, MUT); x.alignment = Alignment(horizontal="center")
    r += 1
    pyv = lambda p, yr: f'INDEX({rng(p + "_" + str(yr))},m_row)'
    rows = [
        ("GDP vs 2024 / today's path, with the programme (%)", lambda yr: f'=IFERROR({pyv("pyg", yr)}+0,"")',
         lambda yr: f"={at('gdp_vs_2024_P', yr)}", "0.00"),
        ("GDP, programme vs no-programme baseline (%)",
         lambda yr: f'=IFERROR(((1+{pyv("pyg", yr)}/100)/(1+{pyv("pygb", yr)}/100)-1)*100,"")',
         lambda yr: f"={at('gdp_dev_prog', yr)}", "0.00"),
        ("Public debt, with the programme (% GDP)", lambda yr: f'=IFERROR({pyv("pyd", yr)}+0,"")',
         lambda yr: f"={at('P_ratio', yr)}*100", "0.0"),
        ("Public debt, baseline (% GDP)", lambda yr: f'=IFERROR({pyv("pydb", yr)}+0,"")',
         lambda yr: f"={at('B_ratio', yr)}*100", "0.0"),
        ("Debt added by the programme (pp of GDP)", lambda yr: f'=IFERROR({pyv("pyd", yr)}-{pyv("pydb", yr)},"")',
         lambda yr: f"={at('debt_diff_pp', yr)}", "0.0"),
        ("Transfers, programme vs baseline: the fiscal rule (pp of GDP)", None, lambda yr: f"={at('tr_diff', yr)}", "0.00"),
        ("Hand-to-mouth consumption, programme vs baseline (%)", None, lambda yr: f"={at('chtm_dev_prog', yr)}", "0.00"),
        ("Fossil rents, resource block (% GDP)", None, lambda yr: f"={at('rents', yr)}*100", "0.0"),
    ]
    for lab, fpy, fsm, fmt in rows:
        wv[f"B{r}"] = lab; wv[f"B{r}"].font = font(9); wv[f"B{r}"].border = BOX
        for cc, yr in (("C", 2030), ("D", 2050)):
            x = wv[f"{cc}{r}"]; x.value = fpy(yr) if fpy else "-"; x.number_format = fmt
            x.font = font(10, True); x.border = BOX; x.alignment = Alignment(horizontal="center")
        for cc, yr in (("F", 2030), ("G", 2050)):
            x = wv[f"{cc}{r}"]; x.value = f'=IF(ISNA(m_row),"",{fsm(yr)[1:]})'; x.number_format = fmt
            x.font = font(10, True); x.border = BOX; x.alignment = Alignment(horizontal="center")
        r += 1
    r += 1
    for c in "BCDEFG":
        wv[f"{c}{r}"].fill = SEC
    wv[f"B{r}"] = "Simple model: calibration"; wv[f"B{r}"].font = font(10, True); r += 1
    for k, lab, fmt in [("s", "Public investment efficiency, s (from PEFA PI-11 or income group)", "0.00"),
                        ("Rg", "Return on public capital", "0.00"), ("ig", "Public investment at the start (% GDP / 100)", "0.000"),
                        ("res", "Resource block on? (1 = DIGNAR)", "0"), ("rents0", "Rents at the start, share of non-resource GDP", "0.000"),
                        ("rstar", "Real commercial rate", "0.000"), ("conc", "Concessional share of debt", "0.00"),
                        ("lam_f", "Fiscal rule speed (from the integrated run)", "0.00")]:
        wv[f"B{r}"] = lab; wv[f"B{r}"].font = font(9); wv[f"B{r}"].border = BOX
        x = wv[f"F{r}"]; x.value = f'=IFERROR(c_{k},"")'; x.number_format = fmt; x.font = font(10, True); x.border = BOX
        x.alignment = Alignment(horizontal="center"); r += 1
    r += 1
    notes = [
        "How the two differ. The integrated run solves the whole path at once with forward-looking savers and lets the consumption tax and transfers respond; "
        "GIMF-lite (most upper-middle and high-income countries) also has energy savings, private investment and the exchange rate. The simple model is DIGNAD-type "
        "for every country: each year follows from the last, private capital closes p_lambda_k of the gap to its desired level, one fiscal instrument (transfers).",
        "Scenarios. Baseline: no programme; fossil rents still fall (a global transition). Programme: NMIM's incremental budget need (capital builds public capital; "
        "land-use and non-CO2 spending does not), financed by grants (p_sh_grant), concessional loans (p_sh_conc) and commercial debt, less earmarked pricing revenue (p_pricing).",
        "Try: p_dis_year = 2030 for a disaster; p_pricing = 0 to drop the pricing revenue; p_sh_grant = 0.5 for more grants; p_programme = 0 to switch the programme off.",
        "Four inputs are income-group defaults until they are in the dashboard: public investment, return on public capital, hand-to-mouth share, grants (Macro_inputs).",
    ]
    for t_ in notes:
        wv[f"B{r}"] = t_; wv[f"B{r}"].font = font(8, italic=True, color=MUT)
        wv[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="top"); wv.merge_cells(f"B{r}:G{r}")
        wv.row_dimensions[r].height = 40; r += 1

    def chart(title, cols_, anchor, ytitle):
        ch = LineChart(); ch.title = title; ch.height = 7.5; ch.width = 17; ch.y_axis.title = ytitle
        ch.legend.position = "b"; ch.varyColors = False
        for c in cols_:
            j = MC.index(c) + 2
            ch.add_data(Reference(wm, min_col=j, max_col=j, min_row=H0, max_row=RL), titles_from_data=True)
        ch.set_categories(Reference(wm, min_col=2, max_col=2, min_row=R0, max_row=RL))
        ch.x_axis.delete = False; ch.y_axis.delete = False
        wv.add_chart(ch, anchor)
    chart("GDP (2024 = 100)", ["ch_gdp_P", "ch_gdp_B", "ch_gdp_py", "ch_gdp_pyb"], "I2", "index")
    chart("Public debt, % of GDP", ["ch_debt_P", "ch_debt_B", "ch_debt_py", "ch_debt_pyb"], "I18", "% GDP")
    chart("Simple model: transfers under the fiscal rule, % of GDP", ["ch_tr_P", "ch_tr_B"], "I34", "% GDP")
    chart("Simple model: programme, pricing revenue and rents, % of GDP", ["ch_prog", "ch_rev", "ch_rents"], "I50", "% GDP")

    extra = [wb[n] for n in ("Macro_assumptions", "Macro_model", "Macro_inputs")]
    rest = [ws for ws in wb._sheets if ws not in extra]
    k = rest.index(wv) + 1
    wb._sheets = rest[:k] + extra + rest[k:]
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    print(f"macro_live: Macro sheets for {len(wide)} countries")
    return M, R0


if __name__ == "__main__":
    if "--check" in sys.argv:
        isos = sys.argv[sys.argv.index("--check") + 1:]
        wide, long, _ = load(isos)
        res = dr.python_check(wide, long)
        for iso, sc in res.items():
            for t in (6, 26):
                b, p = sc["B"][t], sc["P"][t]
                print(iso, Y0 + t, f"gdp_vs_2024_P={100 * (p['GDP'] / sc['P'][0]['GDP'] - 1):.6f}",
                      f"gdp_dev_prog={100 * (p['GDP'] / b['GDP'] - 1):.6f}",
                      f"debt_P={100 * p['ratio']:.6f}", f"debt_B={100 * b['ratio']:.6f}")
