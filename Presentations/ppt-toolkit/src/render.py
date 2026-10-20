"""Render a JSON deck (matching schema/deck_schema.json) into a .pptx file
using either the white or WB-branded template.

Usage:
    python -m src.render deck.json --template white -o out.pptx
    python -m src.render deck.json --template wb -o out.pptx
"""

from __future__ import annotations

import argparse
import json
import os
import sys

from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from layouts import LAYOUTS  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATES_DIR = os.path.normpath(os.path.join(HERE, "..", "templates"))

TEMPLATE_PATHS = {
    "white": os.path.join(TEMPLATES_DIR, "white_template.pptx"),
    "wb": os.path.join(TEMPLATES_DIR, "wb_branded_template.pptx"),
}

CHART_TYPE_MAP = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
}


def load_schema_validator():
    try:
        import jsonschema
    except ImportError:
        return None
    schema_path = os.path.join(HERE, "..", "schema", "deck_schema.json")
    with open(schema_path, encoding="utf-8") as fh:
        schema = json.load(fh)
    return lambda deck: jsonschema.validate(instance=deck, schema=schema)


def get_layout_by_name(prs: Presentation, name: str):
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            if layout.name == name:
                return layout
    raise KeyError(f"Template does not contain a layout named {name!r}")


def _find_placeholder(slide, ph_type_str: str):
    """Find a placeholder on `slide` by our schema ph type string
    ('title'|'subTitle'|'body'|'pic'), matching python-pptx's
    PP_PLACEHOLDER_TYPE enum names loosely."""
    type_map = {
        "title": {"TITLE", "CENTER_TITLE"},
        "subTitle": {"SUBTITLE"},
        "body": {"BODY", "OBJECT"},
        "pic": {"PICTURE", "OBJECT"},
    }
    wanted = type_map.get(ph_type_str, set())
    for ph in slide.placeholders:
        if ph.placeholder_format.type is not None and ph.placeholder_format.type.name in wanted:
            return ph
    return None


def _set_bullets(text_frame, bullets):
    text_frame.clear()
    first = True
    for item in bullets:
        if isinstance(item, str):
            text, level = item, 0
        else:
            text, level = item.get("text", ""), item.get("level", 0)
        p = text_frame.paragraphs[0] if first else text_frame.add_paragraph()
        first = False
        p.text = text
        p.level = max(0, min(4, level))


def _add_chart(slide, placeholder, chart_spec):
    chart_data = CategoryChartData()
    chart_data.categories = chart_spec["categories"]
    for s in chart_spec["series"]:
        chart_data.add_series(s["name"], s["values"])
    chart_type = CHART_TYPE_MAP.get(chart_spec["type"], XL_CHART_TYPE.COLUMN_CLUSTERED)
    left, top, width, height = placeholder.left, placeholder.top, placeholder.width, placeholder.height
    placeholder._element.getparent().remove(placeholder._element)
    slide.shapes.add_chart(chart_type, left, top, width, height, chart_data)


def _add_table(slide, placeholder, table_spec):
    headers = table_spec["headers"]
    rows = table_spec["rows"]
    n_rows = len(rows) + 1
    n_cols = len(headers)
    left, top, width, height = placeholder.left, placeholder.top, placeholder.width, placeholder.height
    placeholder._element.getparent().remove(placeholder._element)
    graphic_frame = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
    table = graphic_frame.table
    for c, header in enumerate(headers):
        table.cell(0, c).text = str(header)
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            if c < n_cols:
                table.cell(r, c).text = str(value)


def render_slide(prs: Presentation, slide_spec: dict, base_dir: str):
    layout_name = slide_spec["layout"]
    if layout_name not in LAYOUTS:
        raise ValueError(f"Unknown layout {layout_name!r}; must be one of {list(LAYOUTS)}")
    layout = get_layout_by_name(prs, layout_name)
    slide = prs.slides.add_slide(layout)

    if "title" in slide_spec:
        ph = _find_placeholder(slide, "title")
        if ph is not None:
            ph.text_frame.text = slide_spec["title"]

    if "subtitle" in slide_spec:
        ph = _find_placeholder(slide, "subTitle")
        if ph is not None:
            ph.text_frame.text = slide_spec["subtitle"]

    body_ph = _find_placeholder(slide, "body")

    if slide_spec.get("chart") and body_ph is not None:
        _add_chart(slide, body_ph, slide_spec["chart"])
    elif slide_spec.get("table") and body_ph is not None:
        _add_table(slide, body_ph, slide_spec["table"])
    elif "bullets" in slide_spec and body_ph is not None:
        bullets = list(slide_spec["bullets"])
        if slide_spec.get("source"):
            bullets.append({"text": f"Source: {slide_spec['source']}", "level": 0})
        _set_bullets(body_ph.text_frame, bullets)
    elif "message" in slide_spec and body_ph is not None:
        lines = slide_spec["message"].split("\n")
        _set_bullets(body_ph.text_frame, lines)
        for p in body_ph.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER

    if slide_spec.get("image"):
        pic_ph = _find_placeholder(slide, "pic")
        image_path = slide_spec["image"]["path"]
        if not os.path.isabs(image_path):
            image_path = os.path.normpath(os.path.join(base_dir, image_path))
        if pic_ph is not None and os.path.exists(image_path):
            pic_ph.insert_picture(image_path)

    if slide_spec.get("notes"):
        slide.notes_slide.notes_text_frame.text = slide_spec["notes"]

    return slide


def render_deck(deck: dict, template_key: str, base_dir: str) -> Presentation:
    if template_key not in TEMPLATE_PATHS:
        raise ValueError(f"--template must be one of {list(TEMPLATE_PATHS)}")
    prs = Presentation(TEMPLATE_PATHS[template_key])
    # Templates ship with zero slides, but defensively clear sldIdLst in case
    # a future template revision includes example/demo slides.
    xml_slides = prs.slides._sldIdLst
    for sld_id in list(xml_slides):
        xml_slides.remove(sld_id)

    for slide_spec in deck["slides"]:
        render_slide(prs, slide_spec, base_dir)
    return prs


def main(argv=None):
    parser = argparse.ArgumentParser(description="Render a JSON deck into a .pptx file.")
    parser.add_argument("deck_json", help="Path to a JSON file matching schema/deck_schema.json")
    parser.add_argument("--template", choices=["white", "wb"], required=True)
    parser.add_argument("-o", "--output", required=True, help="Output .pptx path")
    parser.add_argument("--skip-validate", action="store_true", help="Skip JSON Schema validation")
    args = parser.parse_args(argv)

    with open(args.deck_json, encoding="utf-8") as fh:
        deck = json.load(fh)

    if not args.skip_validate:
        validator = load_schema_validator()
        if validator is not None:
            validator(deck)

    base_dir = os.path.dirname(os.path.abspath(args.deck_json))
    prs = render_deck(deck, args.template, base_dir)
    os.makedirs(os.path.dirname(os.path.abspath(args.output)) or ".", exist_ok=True)
    prs.save(args.output)
    print(f"wrote {args.output} ({len(prs.slides._sldIdLst)} slides, template={args.template})")


if __name__ == "__main__":
    main()
