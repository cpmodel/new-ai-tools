"""Extract an existing .pptx's content into the deck JSON schema format
(best-effort layout-name mapping) so old decks can be re-rendered with
render.py in either template.

Usage:
    python -m src.harvest old.pptx -o extracted.json
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER, MSO_SHAPE_TYPE

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from layouts import LAYOUT_ORDER  # noqa: E402

# Fraction of slide height below which a text shape is treated as a running
# page footer/page-number (e.g. "Deck name -- 12") rather than real content.
FOOTER_TOP_FRACTION = 0.85

# A trailing "-- <anything with a digit>" tail on a footer-ish shape (e.g.
# "Source: ... -- Deck name -- 12") is assumed to be the page-number suffix
# glued onto a real source citation; strip it off if present.
_FOOTER_SUFFIX_RE = re.compile(
    r"^(?P<body>.*?)\s*[\u2013\u2014\u00b7-]\s*[^\d]*?\s*[\u2013\u2014\u00b7-]\s*\d+\s*$"
)


def _guess_layout(slide) -> str:
    """Best-effort mapping of an arbitrary source slide to one of our 5
    canonical layout names.

    If the slide already uses one of our canonical layout names (e.g. it was
    rendered by this same toolkit, or a source deck happens to reuse the
    name), trust that directly -- this makes round-tripping our own decks
    exact. Otherwise fall back to placeholder-type heuristics (little/no
    text => divider/closing style slide) for genuinely foreign decks.
    """
    layout_name = slide.slide_layout.name
    if layout_name in LAYOUT_ORDER:
        return layout_name

    types = set()
    for ph in slide.placeholders:
        if ph.placeholder_format.type is not None:
            types.add(ph.placeholder_format.type)

    has_subtitle = PP_PLACEHOLDER.SUBTITLE in types
    has_body = PP_PLACEHOLDER.BODY in types or PP_PLACEHOLDER.OBJECT in types
    has_picture = PP_PLACEHOLDER.PICTURE in types or any(
        shape.shape_type is not None and shape.shape_type == 13 for shape in slide.shapes  # MSO_SHAPE_TYPE.PICTURE
    )
    has_title = PP_PLACEHOLDER.TITLE in types or PP_PLACEHOLDER.CENTER_TITLE in types

    text_len = sum(len(_shape_text(s)) for s in slide.shapes)

    if has_subtitle:
        return "Title Slide"
    if has_body and has_picture:
        return "Half Half"
    if has_body:
        return "Title and Bullets"
    if has_title and text_len < 80:
        return "Section Divider"
    if has_title:
        return "Closing"
    return "Title and Bullets"


def _shape_text(shape) -> str:
    if not shape.has_text_frame:
        return ""
    return "\n".join(p.text for p in shape.text_frame.paragraphs if p.text)


def _extract_bullets(shape) -> list[dict]:
    bullets = []
    if not shape.has_text_frame:
        return bullets
    for p in shape.text_frame.paragraphs:
        if p.text.strip():
            bullets.append({"text": p.text, "level": p.level or 0})
    return bullets


def _is_header_text(text: str) -> bool:
    """Heuristic: short, punchy text (a card/section heading) vs. a longer
    descriptive sentence (its detail). Used only for decks with no
    placeholders, where we have to reconstruct title/bullet structure from
    freeform text boxes positioned like a card layout."""
    text = text.strip()
    if not text:
        return False
    if text == text.upper() and any(c.isalpha() for c in text):
        return True
    return len(text) <= 40 and not text.rstrip().endswith((".", ":", ";"))


def _strip_footer_suffix(text: str) -> str | None:
    """If a bottom-of-slide text shape looks like 'Source: ... -- Deck name --
    12' (a real citation with the running page footer glued on), return just
    the citation part. If it looks like a bare page footer/number with no
    real citation, return None (drop it entirely)."""
    text = text.strip()
    if not text:
        return None
    if "source" not in text.lower():
        return None
    m = _FOOTER_SUFFIX_RE.match(text)
    body = m.group("body").strip() if m else text
    return body or None


def _save_picture(shape, images_dir: str, slide_index: int, pic_index: int) -> str | None:
    """Extract an embedded picture shape's bytes to `images_dir` and return a
    path relative to `images_dir`'s parent (suitable for the deck JSON's
    `image.path`, resolved relative to the JSON file itself)."""
    try:
        image = shape.image
    except (AttributeError, ValueError):
        return None
    os.makedirs(images_dir, exist_ok=True)
    filename = f"slide{slide_index + 1:02d}_pic{pic_index}.{image.ext}"
    with open(os.path.join(images_dir, filename), "wb") as fh:
        fh.write(image.blob)
    return os.path.join(os.path.basename(images_dir), filename).replace(os.sep, "/")


def _harvest_freeform_slide(slide, slide_index: int, slide_height: int, images_dir: str) -> dict:
    """Best-effort extraction for slides built entirely from freeform text
    boxes/shapes/tables (no slide-layout placeholders) -- common in
    hand-designed consulting decks. Reconstructs a title (+ optional
    eyebrow/section label folded in), a table or picture if present, and the
    remaining text as a bulleted list, using position + simple text
    heuristics rather than placeholder types."""
    text_shapes = []
    table_shape = None
    picture_shapes = []
    for shape in sorted(slide.shapes, key=lambda s: (s.left or 0, s.top or 0)):
        if getattr(shape, "has_table", False):
            table_shape = table_shape or shape
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            picture_shapes.append(shape)
            continue
        if shape.has_text_frame and shape.text_frame.text.strip():
            text_shapes.append(shape)

    source_text = None
    content_shapes = []
    for shape in text_shapes:
        text = shape.text_frame.text
        if slide_height and (shape.top or 0) >= FOOTER_TOP_FRACTION * slide_height:
            captured = _strip_footer_suffix(text)
            if captured:
                source_text = captured
            continue
        content_shapes.append(shape)

    spec: dict = {}

    if slide_index == 0 and table_shape is None and not picture_shapes:
        # First slide with no table/picture: treat as the deck cover.
        texts = [s.text_frame.text.strip() for s in content_shapes if s.text_frame.text.strip()]
        spec["layout"] = "Title Slide"
        if texts:
            spec["title"] = texts[1] if len(texts) > 1 and _is_header_text(texts[0]) else texts[0]
            rest = [t for t in texts if t != spec["title"]]
            if rest:
                spec["subtitle"] = "\n".join(rest)
        if source_text:
            spec["source"] = source_text
        return spec

    # Fold a short leading "eyebrow" label (e.g. "PART 2 -- TWO PRODUCTS
    # COMPARED") into the title, since our schema has no separate field.
    title_parts = []
    body_shapes = list(content_shapes)
    if body_shapes and _is_header_text(body_shapes[0].text_frame.text.strip()):
        title_parts.append(body_shapes.pop(0).text_frame.text.strip())
    if body_shapes:
        title_parts.append(body_shapes.pop(0).text_frame.text.strip())
    if title_parts:
        spec["title"] = " \u2014 ".join(title_parts)

    if table_shape is not None:
        table = table_shape.table
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if rows:
            spec["layout"] = "Title and Bullets"
            spec["table"] = {"headers": rows[0], "rows": rows[1:]}
            leftover = [s.text_frame.text.strip() for s in body_shapes if s.text_frame.text.strip()]
            if leftover:
                spec["bullets"] = [{"text": t, "level": 0} for t in leftover]
    elif picture_shapes:
        spec["layout"] = "Half Half"
        image_path = _save_picture(picture_shapes[0], images_dir, slide_index, 0)
        if image_path:
            spec["image"] = {"path": image_path}
        bullets = []
        for shape in body_shapes:
            text = shape.text_frame.text.strip()
            if text:
                bullets.append({"text": text, "level": 0 if _is_header_text(text) else 1})
        if bullets:
            spec["bullets"] = bullets
    else:
        spec["layout"] = "Title and Bullets"
        bullets = []
        for shape in body_shapes:
            text = shape.text_frame.text.strip()
            if text:
                bullets.append({"text": text, "level": 0 if _is_header_text(text) else 1})
        if bullets:
            spec["bullets"] = bullets

    if source_text:
        spec["source"] = source_text

    return spec


def harvest_slide(slide, slide_index: int = 0, slide_height: int = 0, images_dir: str = "") -> dict:
    if not list(slide.placeholders):
        spec = _harvest_freeform_slide(slide, slide_index, slide_height, images_dir)
    else:
        layout_name = _guess_layout(slide)
        spec = {"layout": layout_name}

        title_ph = None
        subtitle_ph = None
        body_ph = None
        for ph in slide.placeholders:
            t = ph.placeholder_format.type
            if t in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
                title_ph = ph
            elif t == PP_PLACEHOLDER.SUBTITLE:
                subtitle_ph = ph
            elif t in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT):
                body_ph = body_ph or ph

        if title_ph is not None and _shape_text(title_ph).strip():
            spec["title"] = _shape_text(title_ph)

        if layout_name == "Title Slide" and subtitle_ph is not None:
            text = _shape_text(subtitle_ph)
            if text.strip():
                spec["subtitle"] = text

        if body_ph is not None:
            bullets = _extract_bullets(body_ph)
            if bullets:
                if layout_name == "Closing":
                    spec["message"] = "\n".join(b["text"] for b in bullets)
                else:
                    spec["bullets"] = bullets

    if slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame.text
        if notes.strip():
            spec["notes"] = notes

    return spec


def harvest(pptx_path: str, images_dir: str | None = None) -> dict:
    prs = Presentation(pptx_path)
    if images_dir is None:
        images_dir = os.path.join(os.path.dirname(os.path.abspath(pptx_path)), "harvested_images")
    deck = {
        "title": os.path.splitext(os.path.basename(pptx_path))[0],
        "slides": [
            harvest_slide(slide, i, prs.slide_height, images_dir) for i, slide in enumerate(prs.slides)
        ],
    }
    return deck


def main(argv=None):
    parser = argparse.ArgumentParser(description="Extract a .pptx's content into deck_schema JSON.")
    parser.add_argument("pptx_path")
    parser.add_argument("-o", "--output", required=True)
    args = parser.parse_args(argv)

    images_dir = os.path.join(os.path.dirname(os.path.abspath(args.output)) or ".", "images")
    deck = harvest(args.pptx_path, images_dir=images_dir)
    os.makedirs(os.path.dirname(os.path.abspath(args.output)) or ".", exist_ok=True)
    with open(args.output, "w", encoding="utf-8") as fh:
        json.dump(deck, fh, indent=2)
    layout_counts = {}
    for s in deck["slides"]:
        layout_counts[s["layout"]] = layout_counts.get(s["layout"], 0) + 1
    print(f"wrote {args.output} ({len(deck['slides'])} slides) layout mix: {layout_counts}")


if __name__ == "__main__":
    main()
