"""Slide kit in the style of AIClimate-FiscalTools_v2.5.pptx (navy title slides, warm off-white content
slides, Cambria headings, Calibri body, orange eyebrows, cards with a coloured top rule, a navy 'So:' bar).

Every builder takes plain data, so a deck is a list of calls (see build_v2.py)."""
import copy
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NAVY, PAPER, MIST, ORANGE = "14213D", "F7F6F1", "EAEEF2", "A8521A"
BLUE, TEAL, GOLD, BODY, MUTED = "1F5A96", "3D7A73", "E8B04A", "3F4A5A", "5B6573"
RULE, WHITE, PALE = "E3E1DA", "FFFFFF", "C9D3E0"
GROUP = {"engine": NAVY, "fiscal": ORANGE, "climate": TEAL, "pfr": BLUE, "other": GOLD}

W, H = 12192000, 6858000
M = 812800                      # side margin
CW = W - 2 * M                  # content width
HEAD, SERIF = "Calibri", "Cambria"


def nlines(text, pt, width, bold=True):
    """Rough line count for text at pt points in a box width EMU wide."""
    import math
    cw = (0.56 if bold else 0.5) * pt * 12700
    return max(1, math.ceil(len(text) * cw / max(width, 1)))


def rgb(h):
    return RGBColor.from_string(h)


class Deck:
    def __init__(self, template, footer):
        self.prs = Presentation(template)
        sld = self.prs.slides._sldIdLst                 # keep the theme and size, drop the slides
        for s in list(sld):
            self.prs.part.drop_rel(s.rId)
            sld.remove(s)
        self.footer = footer
        self.n = 0

    # ---- primitives -------------------------------------------------------------------------
    def _slide(self, bg):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[0])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(bg)
        self.n += 1
        return s

    def text(self, s, x, y, w, h, runs, size=12, color=BODY, font=HEAD, bold=False, align=None,
             anchor=MSO_ANCHOR.TOP, spacing=None):
        """runs: str, or list of paragraphs; a paragraph is str or list of (text, {opts})."""
        tb = s.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        paras = runs if isinstance(runs, list) else [runs]
        for i, p in enumerate(paras):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            if align:
                para.alignment = align
            if spacing:
                para.space_after = Pt(spacing)
            for t, o in (p if isinstance(p, list) else [(p, {})]):
                r = para.add_run()
                r.text = t
                f = r.font
                f.name, f.size = o.get("font", font), Pt(o.get("size", size))
                f.bold = o.get("bold", bold)
                f.italic = o.get("italic", False)
                f.color.rgb = rgb(o.get("color", color))
        return tb

    def box(self, s, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
        b = s.shapes.add_shape(shape, Emu(x), Emu(y), Emu(w), Emu(h))
        b.fill.solid()
        b.fill.fore_color.rgb = rgb(fill)
        if line:
            b.line.color.rgb = rgb(line)
            b.line.width = Pt(0.75)
        else:
            b.line.fill.background()
        b.shadow.inherit = False
        return b

    def bullets(self, s, x, y, w, h, items, size=12, color=BODY, mark="•", spacing=4):
        paras = []
        for it in items:
            if isinstance(it, tuple):           # (bold lead, rest)
                paras.append([(f"{mark}  ", {"color": ORANGE}), (it[0], {"bold": True, "color": NAVY}), (" " + it[1], {})])
            else:
                paras.append([(f"{mark}  ", {"color": ORANGE}), (it, {})])
        return self.text(s, x, y, w, h, paras, size=size, color=color, spacing=spacing)

    def arrow(self, s, x1, y1, x2, y2, color=MUTED):
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(x1), Emu(y1), Emu(x2), Emu(y2))
        c.line.color.rgb = rgb(color)
        c.line.width = Pt(1.5)
        ln = c.line._get_or_add_ln()
        tail = ln.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd", {"type": "triangle"})
        ln.append(tail)
        return c

    # ---- slide frames -----------------------------------------------------------------------
    def content(self, eyebrow, title, bg=PAPER, source=None, eyecolor=ORANGE):
        s = self._slide(bg)
        self.text(s, M, M, CW, 213320, eyebrow.upper(), size=10, bold=True, color=eyecolor)
        self.text(s, M, 1127720, CW, 467320, title, size=26, bold=True, color=NAVY, font=SERIF)
        foot = f"{self.footer} · {self.n}"
        if source:
            foot = f"{source} · {foot}"
        self.text(s, M, H - 520000, CW, 200000, foot, size=9, color="7A8494")
        return s

    def title(self, eyebrow, title, subtitle, byline):
        s = self._slide(NAVY)
        self.text(s, M, 1650000, CW * 0.8, 220000, eyebrow.upper(), size=11, bold=True, color="E3A36B")
        self.text(s, M, 2000000, CW * 0.75, 1500000, title, size=40, bold=True, color=WHITE, font=SERIF)
        self.text(s, M, 2000000 + 640000 * nlines(title, 40, CW * 0.75) + 250000, CW * 0.8, 600000, subtitle, size=16, color=PALE)
        self.text(s, M, H - 900000, CW, 200000, byline, size=10, color="9AA8BC")
        return s

    def divider(self, eyebrow, title, subtitle, accent="E3A36B"):
        s = self._slide(NAVY)
        self.text(s, M, 2250000, CW, 220000, eyebrow.upper(), size=11, bold=True, color=accent)
        self.text(s, M, 2550000, CW, 900000, title, size=44, bold=True, color=WHITE, font=SERIF)
        self.text(s, M, 2550000 + 700000 * nlines(title, 44, CW) + 250000, CW * 0.8, 800000, subtitle, size=15, color=PALE)
        self.text(s, M, H - 520000, CW, 200000, f"{self.footer} · {self.n}", size=9, color="9AA8BC")
        return s

    def so(self, s, text, y=None):
        y = y or H - 1180000
        self.box(s, M, y, CW, 400000, NAVY)
        self.text(s, M + 200000, y, CW - 400000, 400000, [[("So: ", {"bold": True, "color": "E3A36B"}), (text, {"color": WHITE})]],
                  size=13, anchor=MSO_ANCHOR.MIDDLE)

    # ---- layouts ----------------------------------------------------------------------------
    def cards(self, s, items, y=1800000, h=3450000, gap=180000, size=11):
        """items: dict(tag, color, title, bullets, pct, status, foot, big)"""
        n = len(items)
        w = (CW - gap * (n - 1)) // n
        pad = 190000
        for i, it in enumerate(items):
            x = M + i * (w + gap)
            col = it.get("color", ORANGE)
            self.box(s, x, y, w, h, WHITE)
            self.box(s, x, y, w, 50000, col)
            yy = y + 190000
            if it.get("tag"):
                self.text(s, x + pad, yy, w - 2 * pad, 200000, it["tag"].upper(), size=9, bold=True, color=col)
                yy += 220000
            if it.get("big"):
                self.text(s, x + pad, yy, w - 2 * pad, 500000, it["big"], size=30, bold=True, color=col, font=SERIF)
                yy += 520000
            self.text(s, x + pad, yy, w - 2 * pad, 480000, it["title"], size=14, bold=True, color=NAVY)
            yy += 140000 + 240000 * nlines(it["title"], 14, w - 2 * pad)
            if it.get("bullets"):
                bh = h - (yy - y) - (700000 if it.get("pct") is not None or it.get("foot") else 120000)
                self.bullets(s, x + pad, yy, w - 2 * pad, max(bh, 250000), it["bullets"], size=size)
            if it.get("pct") is not None:
                by = y + h - 760000
                self.text(s, x + pad, by - 330000, w - 2 * pad, 330000,
                          [[(f"{it['pct']}%", {"size": 22, "bold": True, "color": col, "font": SERIF}),
                            ("  " + it.get("status", "complete"), {"size": 10, "color": MUTED})]])
                self.box(s, x + pad, by + 40000, w - 2 * pad, 90000, RULE)
                if it["pct"]:
                    self.box(s, x + pad, by + 40000, int((w - 2 * pad) * it["pct"] / 100), 90000, col)
            if it.get("foot"):
                self.text(s, x + pad, y + h - 520000, w - 2 * pad, 400000, it["foot"], size=10, color=MUTED)

    def table(self, s, headers, rows, widths=None, y=1800000, size=11, h_row=330000, head_fill=NAVY):
        t = s.shapes.add_table(len(rows) + 1, len(headers), Emu(M), Emu(y), Emu(CW), Emu(h_row * (len(rows) + 1))).table
        widths = widths or [1 / len(headers)] * len(headers)
        for j, wv in enumerate(widths):
            t.columns[j].width = Emu(int(CW * wv))
        for i, row in enumerate([headers] + rows):
            for j, val in enumerate(row):
                c = t.cell(i, j)
                c.fill.solid()
                c.fill.fore_color.rgb = rgb(head_fill if i == 0 else (WHITE if i % 2 else "F1F0EB"))
                c.margin_left = c.margin_right = Emu(90000)
                c.margin_top = c.margin_bottom = Emu(40000)
                tf = c.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                r = p.add_run()
                r.text = str(val)
                r.font.name, r.font.size = HEAD, Pt(size)
                r.font.bold = i == 0 or (j == 0 and len(headers) > 1)
                r.font.color.rgb = rgb(WHITE if i == 0 else (NAVY if j == 0 else BODY))
        return t

    def flow(self, s, steps, y=1900000, h=1150000, last_dark=True, size=11):
        """steps: list of (title, text, color or None)."""
        n = len(steps)
        gap = 330000
        w = (CW - gap * (n - 1)) // n
        for i, (ttl, txt, col) in enumerate(steps):
            x = M + i * (w + gap)
            dark = last_dark and i == n - 1
            fill = NAVY if dark else (col or WHITE)
            self.box(s, x, y, w, h, fill, line=None if dark or col else "C9D3E0")
            tc = WHITE if dark or (col and col not in (WHITE, MIST, "DCE8F5", "FBE3CE", "E1EFEC", "FBF0D9")) else NAVY
            tl = nlines(ttl, size + 2, w - 240000)
            self.text(s, x + 120000, y + 120000, w - 240000, 300000 * tl, ttl, size=size + 2, bold=True, color=tc)
            by = y + 170000 + 250000 * tl
            self.text(s, x + 120000, by, w - 240000, h - (by - y) - 60000, txt, size=size - 1,
                      color=(PALE if tc == WHITE else BODY))
            if i < n - 1:
                self.arrow(s, x + w + 40000, y + h // 2, x + w + gap - 40000, y + h // 2)

    def numbered(self, s, items, y=1800000, step=560000, size=12, col=ORANGE):
        """items: (title, text)"""
        for i, (ttl, txt) in enumerate(items):
            yy = y + i * step
            cs = min(330000, step - 40000)
            c = self.box(s, M, yy, cs, cs, col, shape=MSO_SHAPE.OVAL)
            self.text(s, M, yy, cs, cs, str(i + 1), size=12 if cs > 250000 else 9, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
                      anchor=MSO_ANCHOR.MIDDLE)
            self.text(s, M + 500000, yy + 30000, 3000000, 400000, ttl, size=size + 1, bold=True, color=NAVY)
            self.text(s, M + 3700000, yy + 30000, CW - 3700000, step - 60000, txt, size=size - 1, color=BODY)

    def gantt(self, s, rows, weeks, y=1750000, label_w=3300000, row_h=300000, today=None, size=10, legend=None):
        """rows: (label, sub, start_idx, end_idx, color, milestone_idx or None); weeks: header labels."""
        gx = M + label_w
        gw = CW - label_w
        n = len(weeks)
        cw = gw / n
        for i, wk in enumerate(weeks):
            if wk:
                self.text(s, int(gx + i * cw), y, int(cw * 2), 200000, wk, size=8, color=MUTED)
        top = y + 250000
        self.box(s, gx, top - 20000, gw, 12000, RULE)
        for r, (lab, sub, a, b, col, ms) in enumerate(rows):
            yy = top + r * row_h
            self.text(s, M, yy + 10000, label_w - 100000, row_h, [[(lab, {"bold": True, "color": NAVY, "size": size}),
                      ("  " + sub if sub else "", {"color": MUTED, "size": size - 1})]])
            if b is not None:
                self.box(s, int(gx + a * cw), yy + 70000, max(int((b - a + 1) * cw) - 20000, 60000), row_h - 140000, col)
            if ms is not None:
                d = self.box(s, int(gx + ms * cw + cw / 2 - 80000), yy + 70000, 160000, 160000, GOLD, shape=MSO_SHAPE.DIAMOND)
        if today is not None:
            self.box(s, int(gx + today * cw), top - 60000, 15000, row_h * len(rows) + 60000, ORANGE)
            self.text(s, int(gx + today * cw) + 40000, top - 250000, 900000, 200000, "today", size=8, bold=True, color=ORANGE)
        if legend:
            ly = top + len(rows) * row_h + 90000
            lx = M
            for name, col in legend:
                self.box(s, lx, ly + 30000, 160000, 120000, col, shape=MSO_SHAPE.DIAMOND if col == GOLD else MSO_SHAPE.RECTANGLE)
                self.text(s, lx + 220000, ly, 2200000, 200000, name, size=9, color=MUTED)
                lx += 220000 + 120000 * 0 + len(name) * 62000 + 300000

    def kpis(self, s, items, y=1800000, h=1100000, gap=180000):
        """items: (number, label, color)"""
        n = len(items)
        w = (CW - gap * (n - 1)) // n
        for i, (num, lab, col) in enumerate(items):
            x = M + i * (w + gap)
            self.box(s, x, y, w, h, WHITE)
            self.box(s, x, y, 50000, h, col)
            self.text(s, x + 200000, y + 130000, w - 300000, 500000, num, size=28, bold=True, color=col, font=SERIF)
            self.text(s, x + 200000, y + 640000, w - 300000, h - 680000, lab, size=11, color=BODY)

    def notes(self, s, text):
        s.notes_slide.notes_text_frame.text = text

    def save(self, path):
        self.prs.save(path)
