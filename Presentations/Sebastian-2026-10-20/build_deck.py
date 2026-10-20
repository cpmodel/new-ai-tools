"""Build the Sebastian deck (20 Oct 2026): one combined deck plus one file per product.

Every product slide uses the same table: What / Platform / Status / Next / Timeline / Ask or risk.
Run from anywhere:  python3 build_deck.py   (needs the ppt-toolkit, clean white template)
"""
import json
import subprocess
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
TOOLKIT = HERE.parent / "ppt-toolkit"
ROWS = ["What it is", "Platform", "Status (20 Oct)", "Next", "Timeline", "Ask / risk"]
SOURCE = "CDPM = the model formerly called SiSePuede. Sources: FY27 work plan v2.6; repos CPAT-AI, CPAT-CDPM, new-ai-tools; task tracker (20 Oct 2026)"


def product(key, title, cells, notes=""):
    """One product slide: the same six rows every time."""
    return key, {"layout": "Title and Bullets", "title": title,
                 "table": {"headers": ["", "Status and timeline"], "rows": [[r, c] for r, c in zip(ROWS, cells)]},
                 "source": SOURCE, "notes": notes}


PRODUCTS = [
    product("CPAT-AI", "CPAT-AI (Excel-AI): mitigation MVP", [
        "Cut-down CPAT in Excel, rebuilt with AI from legacy CPAT; every figure traceable to equation, parameter and source",
        "Reference model for both platforms; feeds the WB-IMF CPAT process",
        "v6.2 (20 Oct): mitigation, power and distribution built; legacy bugs fixed, not copied; MVP 70%",
        "Equations list, traceability, known-gaps note; Kacper holds the bug list while I travel",
        "MVP 23 Oct; power 2-13 Nov; remaining mitigation 16-27 Nov; final 4 Dec",
        "Ownership with the IMF/CPAT team; one Excel model, not two; data requests (GAINS, VAT)",
    ]),
    product("Coded-CPAT", "Coded CPAT (Python)", [
        "Python reimplementation of CPAT (mitigation + distribution), same codes as the Excel; Shiny app on top",
        "The core of BOTH platforms: climate pathways and fiscal",
        "Runs end to end; 57 tests; exports back to the legacy layout. Final Python 10%",
        "Port the validated Excel MVP logic; Python-vs-Excel comparison on the same inputs",
        "Enhanced version to Edmundo early Nov; final Python 11 Dec; Shiny app 5 Feb",
        "Open fixes: distribution revenue units, national totals, full electricity price build-up",
    ]),
    product("CPAT-CDPM", "Pathways: CPAT-CDPM integration (Edmundo)", [
        "Coded CPAT inside CDPM (formerly SiSePuede; Edmundo Molina, Tec de Monterrey; with Dave Groves and Sonya Sampson): one pathways tool",
        "Climate platform (pathways project)",
        "Edmundo meeting 19 Oct went well: on track. CPAT sent; CPAT-CDPM repo created. Egypt proof of concept close to CPAT (-7.3% vs -8.0% emissions)",
        "Integrate early (no wait for a final product); one user flow; methodology note on subsidies vs public investment",
        "User-flow workshop early Nov; integration 4-8 Jan; pilots (Libya, Egypt, Morocco) 11-15 Jan; end Jan",
        "Agree the methodology split (running subsidies vs upfront public investment) before integration",
    ]),
    product("FTT", "FTT national (Zoe)", [
        "National-level data so FTT technology diffusion (power, cars, freight, heat, steel) runs on CPAT's countries",
        "Climate platform; with MFMod, the bridge to the fiscal platform",
        "Zoe delivering; I review. Progress % not yet reported",
        "Ask Zoe for a review date; then FTT + MFMod on one Python platform",
        "FTT project 27 Nov; FTT + MFMod working 12 Feb",
        "Decision: MFMod in or out for Feb (FTT-only is doable in the week)",
    ]),
    product("InvestmentNeeds", "Investment needs (NMIM)", [
        "National mitigation investment model: 186 countries, 2024-2050, built to CPAT norms; who pays (public, utility, private)",
        "Both: investment for pathways; budget need for the fiscal platform",
        "v0.2 (20 Oct). 2030: about USD 5.1tn total, 1.7tn incremental; direct budget need ~14% (85% in low-income countries)",
        "Financing and debt layer linked to MFMod; load cost-of-capital data; add a price response",
        "Proposal stage: fits the pathways timeline (Jan) and the fiscal platform (Feb)",
        "The natural hook for the macro unit: investment -> debt and fiscal space in MFMod",
    ]),
    product("CBAM", "CBAM policy tool", [
        "Excel tool: what a domestic carbon price does to heavy industry and the EU CBAM charge; 6 policy scenarios",
        "Fiscal platform (revenue, credits); process emissions feed the investment model",
        "Egypt v3.0 (7 CBAM goods) and a global MVP v1.0 (9 Oct)",
        "CBAM modelling inside CPAT (to discuss); port to Python with coded CPAT",
        "To agree: after the Excel-AI final (4 Dec)",
        "No trade model; elasticities by judgement. Ask: is CBAM in CPAT a priority?",
    ]),
    product("Excise-Dashboard", "Excise Diagnostic and Tax & Macro-Fiscal Dashboard", [
        "Excise Diagnostic (Sierra Leone, Uganda, Liberia) and a dashboard: 53 indicators, 217 countries",
        "Fiscal platform: tax and macro-fiscal data layer",
        "Done: both handed to Dirk early (week of 5 Oct); all 4 Objective 3 milestones complete",
        "Owners use them; one week of follow-up support",
        "Support 8-12 Mar 2027",
        "Owners not yet named; tools not validated",
    ]),
]

TIMELINE = {"layout": "Title and Bullets", "title": "Timeline: FY27 milestones",
            "table": {"headers": ["Date", "Milestone", "Platform"], "rows": [
                ["19 Oct", "Edmundo on track (done)", "Climate"],
                ["23 Oct", "Excel-AI mitigation MVP", "Both (reference)"],
                ["27 Nov", "Zoe's FTT national project", "Climate"],
                ["4 Dec", "Final Excel-AI", "Both (reference)"],
                ["11 Dec", "Final coded CPAT (Python)", "Both"],
                ["15 Jan", "Pathways tool: integration and pilots", "Climate"],
                ["5 Feb", "Shiny app over coded CPAT", "Both"],
                ["12 Feb", "FTT + MFMod working", "Fiscal (macro)"],
                ["5 Mar", "GHG accounting, crediting, FFI", "Fiscal"],
                ["20 Mar", "Handover; contract ends", "-"]]},
            "source": "FY27 work plan v2.6 (15 Oct); away 21 Oct to about 2 Nov",
            "notes": "Slack is about 3 days before 4 Dec; power was built early, which covers the trip."}

DECK = [
    {"layout": "Title Slide", "title": "Climate-fiscal tools: two Python platforms",
     "subtitle": "Status, timeline and asks\nStephen Stretton, meeting with Sebastian, 20 Oct 2026"},
    {"layout": "Title and Bullets", "title": "Summary: one code base, two platforms",
     "bullets": [
         "Climate policy platform (the pathways project): coded CPAT + CDPM (formerly SiSePuede) + FTT",
         "Fiscal policy platform: coded CPAT + MFMod-Python (macro unit) + investment needs + CBAM",
         {"text": "Coded CPAT sits in both; the Excel-AI model is the checked reference", "level": 1},
         "Goal of one Python platform: work with the macro unit on a shared code base",
         "On track: Edmundo meeting went well (19 Oct); Excel-AI MVP due 23 Oct",
         "Today: agree the two platforms, the macro-unit link, and the look-ahead topics"],
     "source": SOURCE},
    {"layout": "Title and Bullets", "title": "The products at a glance",
     "table": {"headers": ["Product", "Platform", "Status", "Next date"], "rows": [
         ["CPAT-AI (Excel-AI)", "Reference", "MVP 70%, v6.2", "23 Oct"],
         ["Coded CPAT (Python)", "Both", "Runs; final 10%", "11 Dec"],
         ["CPAT-CDPM pathways (Edmundo)", "Climate", "On track", "15 Jan"],
         ["FTT national (Zoe)", "Climate", "Zoe delivering", "27 Nov"],
         ["Investment needs (NMIM)", "Both", "v0.2", "tbc"],
         ["CBAM policy tool", "Fiscal", "Egypt v3.0, global v1.0", "tbc"],
         ["Excise + Macro-Fiscal Dashboard", "Fiscal", "Handed to Dirk", "8-12 Mar"]]},
     "source": SOURCE},
    {"layout": "Section Divider", "title": "Two platforms"},
    {"layout": "Title and Bullets", "title": "A. Climate policy platform: the pathways project",
     "bullets": [
         "What: one Python tool for country decarbonisation pathways and their costs",
         "Parts: coded CPAT, CDPM (formerly SiSePuede; Edmundo), FTT national data (Zoe), investment needs",
         "Users: one integrated flow, not several tools; pilots Libya, Egypt, Morocco",
         "Status: Edmundo on track; Egypt proof of concept close to CPAT",
         "Timeline: integration 4-8 Jan; pilots 11-15 Jan; full integration end Jan",
         "Open: separate running subsidies from upfront public investment"],
     "source": SOURCE},
    {"layout": "Title and Bullets", "title": "B. Fiscal policy platform: with the macro unit",
     "bullets": [
         "What: one Python platform for climate-related fiscal policy, shared with the macro unit",
         "Parts: coded CPAT (revenue, prices, distribution), MFMod-Python (ModelFlow), FTT, investment needs, CBAM",
         "Why: MFMod handles investment, debt and growth; CPAT handles energy, prices and revenue",
         "Already linked: CPAT's GDP block uses MFMod fiscal multipliers",
         "Timeline: FTT + MFMod working 12 Feb; GHG accounting, crediting, FFI 5 Mar",
         "Discussed with Charl; next step: a first call with the macro unit"],
     "source": SOURCE},
    {"layout": "Section Divider", "title": "Products"},
]
DECK += [s for _, s in PRODUCTS]
DECK += [
    TIMELINE,
    {"layout": "Title and Bullets", "title": "Looking ahead: topics to agree",
     "bullets": [
         "Expenditure policy: PFRs lack an expenditure-policy chapter; the platforms can supply one",
         "Energy fiscal risks: what the finance ministry carries (subsidies, utilities, fuel revenue); scope to agree",
         "Investment process: who pays (public, utility, private) and the budget need, from NMIM",
         "Macro unit: shared Python code base; MFMod for debt and growth"],
     "source": SOURCE},
    {"layout": "Title and Bullets", "title": "Asks for today",
     "bullets": [
         "Agree the two platforms, with coded CPAT in both",
         "An introduction to the macro unit (MFMod-Python) for a first call in November",
         "Steer on the look-ahead: expenditure policy, energy fiscal risks, the investment process",
         "Decide: MFMod in or out for the February FTT milestone",
         "Resources: Claude Code ($200); widening the prototype team",
         "CBAM in CPAT: priority or not?"],
     "source": SOURCE},
    {"layout": "Closing", "title": "Thank you",
     "message": "Away 21 Oct to about 2 Nov; Kacper covers CPAT\nStephen Stretton"},
]


def bulletize(slide):
    """The white template has no bullet markers: add them in the text. The source goes in a footer."""
    s = {k: v for k, v in slide.items() if k != "source"}
    if "bullets" in s:
        s["bullets"] = [{"text": ("\u2013 " if isinstance(b, dict) and b.get("level") else "\u2022 ") + (b["text"] if isinstance(b, dict) else b),
                         "level": b.get("level", 0) if isinstance(b, dict) else 0} for b in s["bullets"]]
    return s


def polish(out, slides):
    """Consistent look: larger titles, readable tables, a small source line at the foot."""
    from pptx import Presentation
    from pptx.util import Pt, Inches
    from pptx.dml.color import RGBColor
    prs = Presentation(out)
    for sl, spec in zip(prs.slides, slides):
        for sh in sl.shapes:
            if sh.is_placeholder and sh.placeholder_format.type in (1, 3) and spec["layout"] not in ("Section Divider", "Closing"):
                for para in sh.text_frame.paragraphs:
                    for r in para.runs:
                        r.font.size, r.font.bold, r.font.color.rgb = Pt(26), True, RGBColor(0x1F, 0x3A, 0x5F)
            if sh.has_table:
                t = sh.table
                sh.left, sh.width = Inches(0.5), prs.slide_width - Inches(1)
                n = len(t.columns)
                widths = [Inches(2.2), sh.width - Inches(2.2)] if n == 2 else [sh.width // n] * n
                for col, w in zip(t.columns, widths):
                    col.width = w
                for row in t.rows:
                    for c in row.cells:
                        for para in c.text_frame.paragraphs:
                            para.font.size = Pt(14)
                            for r in para.runs:
                                r.font.size = Pt(14)
                if len(t.columns) == 2:
                    for row in list(t.rows)[1:]:
                        for r in row.cells[0].text_frame.paragraphs[0].runs:
                            r.font.bold = True
            elif sh.has_text_frame and not sh.is_placeholder:
                pass
            elif sh.has_text_frame and sh.placeholder_format.type not in (1, 3):
                for para in sh.text_frame.paragraphs:
                    para.space_after = Pt(8)
                    for r in para.runs:
                        r.font.size = Pt(18)
        if spec.get("source"):
            tb = sl.shapes.add_textbox(Inches(0.5), prs.slide_height - Inches(0.45), prs.slide_width - Inches(1), Inches(0.3))
            r = tb.text_frame.paragraphs[0].add_run()
            r.text, r.font.size, r.font.italic = spec["source"], Pt(9), True
            r.font.color.rgb = RGBColor(0x70, 0x70, 0x70)
    prs.save(out)


def render(slides, out, title):
    js = out.with_suffix(".json")
    js.write_text(json.dumps({"title": title, "author": "Stephen Stretton", "slides": [bulletize(x) for x in slides]}, indent=1))
    subprocess.run([sys.executable, "-m", "src.cli", "render", str(js), "--template", "white", "-o", str(out)],
                   cwd=TOOLKIT, check=True)
    polish(out, slides)


if __name__ == "__main__":
    (HERE / "products").mkdir(exist_ok=True)
    render(DECK, HERE / "Sebastian-TwoPlatforms-20Oct.pptx", "Climate-fiscal tools: two Python platforms")
    for key, slide in PRODUCTS:
        render([slide], HERE / "products" / f"{key}.pptx", slide["title"])
