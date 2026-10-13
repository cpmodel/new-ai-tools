"""v2.2: imports build_v2_2 / check_v2_2 (previous version in Old/; method unchanged).

Compare CPAT-AI-Mitigation-MVP (scenario 1, baseline) with the cached baseline results in the legacy workbook.

The legacy file CPAT 1.0pre_456_NoPropData.xlsb keeps the cached values of its last full run (Egypt, real data):
fuel use and after-tax prices by subsector and fuel, CO2 by subsector. This script
  1. extracts them by output code (egy.mit.ener.<sub>.<fuel>.e.1, egy.mit.atp.<sub>.<fuel>.e.1,
     egy.mit.co2.<grp>.<sub>.1) -> data/legacy_cached_baseline.csv;
  2. recalculates the MVP twice in LibreOffice: as shipped (last historical price year 2024) and with the last
     historical price year set to 2022 (Settings C10; the LAMBDA formulas are copied over whole rows, so the
     history / projection switch follows the setting);
  3. writes legacy_comparison_v2_2.md: totals by sector, selected prices, and the fit of each variant.

    python legacy_compare_v2_2.py
"""
import csv
import os

import openpyxl
from openpyxl.utils import get_column_letter as L

import build_v2_2 as B
import check_v2_2 as C

HERE = os.path.dirname(os.path.abspath(__file__))
LEGACY = os.path.join(HERE, '..', '..', 'cpat_excel_original', 'CPAT 1.0pre_456_NoPropData.xlsb')
CSV_OUT = os.path.join(HERE, 'data', 'legacy_cached_baseline.csv')
REPORT = os.path.join(HERE, 'legacy_comparison_v2_2.md')
YEARS = list(range(2022, 2041))
SUBS = [s[0] for s in B.SUBSECTORS]
GRP = {s[0]: s[2] for s in B.SUBSECTORS}
LEG_GRP = {'tra': 'tra', 'bld': 'bld', 'ind': 'ind', 'oen': 'ind'}    # legacy co2 codes: oen sits in 'oth'
FUELS = [f[0] for f in B.FUELS]


def extract():
    from pyxlsb import open_workbook
    want = {f'egy.mit.ener.{s}.{f}.e.1' for s in SUBS for f in FUELS} | \
        {f'egy.mit.atp.{s}.{f}.e.1' for s in SUBS for f in FUELS} | \
        {f'egy.mit.co2.{LEG_GRP[GRP[s]]}.{s}.1' for s in SUBS if s != 'oen'} | {'egy.mit.co2.oth.1'}
    rows = {}
    with open_workbook(LEGACY) as wb, wb.get_sheet('Mitigation') as sh:
        for i, row in enumerate(sh.rows()):
            vals = [c.v for c in row]
            for j, v in enumerate(vals[:12]):
                if isinstance(v, str) and v in want and v not in rows:
                    rows[v] = (i + 1, [vals[11 + k] if 11 + k < len(vals) else None for k in range(len(YEARS))])
                    break
    with open(CSV_OUT, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f, lineterminator='\n')
        w.writerow(['code', 'legacy_row'] + YEARS + ['source'])
        for k in sorted(rows):
            w.writerow([k, rows[k][0]] + rows[k][1] + ['CPAT 1.0pre_456_NoPropData.xlsb Mitigation (cached values)'])
    missing = sorted(want - set(rows))
    return {k: v[1] for k, v in rows.items()}, missing


def model(last_hist):
    wb = openpyxl.load_workbook(B.OUT)
    if last_hist != B.HIST_YEARS[-1]:
        wb['Settings']['C10'] = last_hist
        path = C.variant(wb, 'everywhere')       # LAMBDA over whole rows, expanded: the switch follows C10
    else:
        path = C.variant(wb, 'expand')
    ws = C.recalc_wb(path)['Mitigation']
    code, base, years = B.group_cols(1)
    vals = {}
    for r in range(1, ws.max_row + 1):
        k = ws.cell(r, code).value
        if isinstance(k, str):
            vals[k] = [ws.cell(r, c).value for c in [base] + years]
    return vals


def main():
    leg, missing = extract()
    ours = {2024: model(2024), 2022: model(2022)}
    units = {r['fuel']: float(r['gj_per_price_unit']) for r in csv.DictReader(open(os.path.join(HERE, 'data',
                                                                                                'fuel_units.csv')))}
    yi = {y: k for k, y in enumerate(YEARS)}
    show = [2022, 2023, 2024, 2025, 2027, 2030]

    def lsum(var, grp):
        out = []
        for y in show:
            if var == 'ener':
                out.append(sum(leg[f'egy.mit.ener.{s}.{f}.e.1'][yi[y]] or 0 for s in SUBS if GRP[s] == grp
                               for f in FUELS))
            else:
                ks = [f'egy.mit.co2.{LEG_GRP[GRP[s]]}.{s}.1' for s in SUBS if GRP[s] == grp and s != 'oen'] + \
                    (['egy.mit.co2.oth.1'] if grp == 'oen' else [])
                out.append(sum(leg[k][yi[y]] or 0 for k in ks))
        return out

    def osum(v, var, grp):
        sfx = '.e' if var == 'ener' else '.e'
        return [v[f'egy.mit.{var}.{grp}.all{sfx}.1'][yi[y]] for y in show]

    lines = ['# Legacy CPAT cached baseline vs CPAT-AI-Mitigation-MVP (scenario 1)', '',
             'Legacy: cached values of the last full run stored in `CPAT 1.0pre_456_NoPropData.xlsb` (Egypt, real '
             f'data), extracted to `data/legacy_cached_baseline.csv`. MVP: v{B.VERSION} scenario 1 recalculated in '
             'LibreOffice, (a) as shipped, last historical price year 2024; (b) last historical price year 2022 '
             '(Settings C10, LAMBDAs copied over whole rows). Fuel use ktoe, CO2 MtCO2.', '',
             f'Legacy codes not found: {", ".join(missing) or "none"}.', '']
    for var, unit in (('ener', 'ktoe'), ('co2', 'MtCO2')):
        lines += [f'## {"Fuel use" if var == "ener" else "CO2 from fuel combustion"} by sector ({unit})', '',
                  '| Sector | Run | ' + ' | '.join(map(str, show)) + ' |', '|---|---|' + '---|' * len(show)]
        for _n, grp, title in B.SECTIONS:
            fmt = (lambda x: f'{x:,.0f}') if var == 'ener' else (lambda x: f'{x:,.1f}')
            lines.append(f'| {title} | legacy | ' + ' | '.join(fmt(x) for x in lsum(var, grp)) + ' |')
            for lh in (2024, 2022):
                lines.append(f'| | MVP, prices to {lh} | ' + ' | '.join(fmt(x) for x in osum(ours[lh], var, grp))
                             + ' |')
        lines.append('')
    lines += ['## After-tax prices, scenario 1 (legacy units: $/liter for liquids, $/GJ otherwise; real 2026 USD)', '',
              '| Row | Run | ' + ' | '.join(map(str, show)) + ' |', '|---|---|' + '---|' * len(show)]
    for s, f in (('rod', 'gso'), ('rod', 'die'), ('res', 'lpg'), ('res', 'nga'), ('cem', 'coa'), ('omn', 'oop')):
        k = f'egy.mit.atp.{s}.{f}.e.1'
        u = units[f] if f in ('gso', 'die', 'lpg', 'ker') else (1.0 if f != 'oop' else units['oop'])
        lines.append(f'| {s} {f} | legacy | ' + ' | '.join(f'{leg[k][yi[y]]:.3f}' for y in show) + ' |')
        for lh in (2024, 2022):
            lines.append(f'| | MVP, prices to {lh} | ' + ' | '.join(f'{ours[lh][k][yi[y]] * u:.3f}' for y in show)
                         + ' |')
    lines.append('')
    # fit: mean absolute % difference of subsector x fuel fuel use, 2023-2030, rows with legacy > 1 ktoe
    lines += ['## Fit: fuel use by subsector and fuel, 2023-2030', '',
              '| Run | Mean abs. % difference vs legacy | Total fuel use 2030 vs legacy |', '|---|---|---|']
    for lh in (2024, 2022):
        diffs, tl, to = [], 0.0, 0.0
        for s in SUBS:
            for f in FUELS:
                k = f'egy.mit.ener.{s}.{f}.e.1'
                for y in range(2023, 2031):
                    a, b = leg[k][yi[y]] or 0, ours[lh][k][yi[y]]
                    if a > 1:
                        diffs.append(abs(b / a - 1))
                tl += leg[k][yi[2030]] or 0
                to += ours[lh][k][yi[2030]]
        lines.append(f'| MVP, prices to {lh} | {sum(diffs) / len(diffs):.1%} | {to / tl - 1:+.1%} |')
    lines.append('')
    with open(REPORT, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines) + '\n')
    print('\n'.join(lines))


if __name__ == '__main__':
    main()
