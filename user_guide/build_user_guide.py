"""Build the AI Tools Quick Start: README.md, .docx, .pptx and the kit .zip.

One source (CONTENT below) -> four outputs, so they never drift apart.
Run from anywhere:  python user_guide/build_user_guide.py
Needs: python-docx, python-pptx.
"""
import os
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm
from pptx import Presentation
from pptx.dml.color import RGBColor as PRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt as PPt

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
NAME = "AI_Tools_Quick_Start"
DATE = "October 2026"

NAVY, BLUE, GOLD, GREY, LIGHT = "1F3864", "2E75B6", "C9A227", "595959", "EEF3F9"

# --------------------------------------------------------------------------- content
TITLE = "AI Tools Quick Start"
SUBTITLE = "Two worked examples: an excise diagnostic for Indonesia (Claude) and new tax indicators in the dashboard (GitHub Copilot)"

TOOLS = [
    ("Claude", "claude.ai", "Deep research and writing. Finds laws and budget papers on the web, reads them, and writes the report and workbook. Use it for Example 1."),
    ("GitHub Copilot", "VS Code or the Copilot app", "Changes files in our repo: code, config, rebuilds the dashboard, pushes to GitHub. Use it for Example 2."),
    (".md process files", "in the repo", "Plain-text instructions the AI follows step by step. You give them to the AI. You do not need to edit them."),
]

BEFORE = [
    "A Claude account that has Projects and Research (deep research).",
    "GitHub Copilot set up, and the repo cloned to C:\\Users\\<you>\\Repos\\new-ai-tools (see Presentations/GitHub_Copilot_AI_User_Guide.pptx).",
    "This kit unzipped. Folder 1 is for Example 1, folder 2 for Example 2.",
]

EX1_TITLE = "Example 1 - Excise diagnostic for Indonesia (Claude)"
EX1_WHAT = "A structured snapshot of Indonesia's excise schedule at 30 Sept 2026: every excise line, quoted from the law, standardized, with gaps and erosion flagged. Output: a Word report and an Excel workbook."
EX1_STEPS = [
    ("Create a Project", "On claude.ai: Projects > Create project. Name it \"Excise - Indonesia\"."),
    ("Add 3 files", "Project knowledge > Add: excise_diagnostic_process.md, Excise_Diagnostic_CrossCountry.xlsx, INDONESIA_INPUTS.md (all in kit folder 1)."),
    ("Start a chat", "New chat inside the Project. Turn on Research and Web search. Pick the most capable model."),
    ("Paste Prompt 0, then A1", "From PROMPTS.md. Prompt 0 checks Claude understood. A1 builds the legal history."),
    ("If Claude asks for files", "It gives you a links list and a .bat file. Run the .bat (or click the links), upload the PDFs to the Project, type \"continue\"."),
    ("Check, then next prompt", "Every rate quoted with its provision and date, every source linked. Then send A2, A3 ... A6, one per message."),
    ("Save the outputs", "Download the Word report and Excel workbook into ExciseDiagnostic/country_results/Indonesia/."),
]
EX1_PROMPTS = [
    ("A1", "Legislative history, 2020 to cut-off, rates as printed"),
    ("A2", "Current schedule, every line, rate trajectory"),
    ("A3", "Parameters, standard schedule, workbook"),
    ("A4", "Findings (Level 1), every number traced"),
    ("A5", "Peer-review register, corrections applied"),
    ("A6", "Final Word report and Excel workbook"),
]
EX1_SUBSIDY = "Want a fossil fuel SUBSIDY diagnostic instead? Use the variant prompt at the end of PROMPTS.md. There is no subsidy process file yet, so treat the output as a first draft."
EX1_SHORT = "Short on time? Add excise_part1_fuel.md to the Project and run Part 1 (fuel) only, with the prompts in that file."

EX2_TITLE = "Example 2 - Add tax indicators to the dashboard (GitHub Copilot)"
EX2_WHAT = "Add four rows under Tax revenue (% GDP) in the Global Fiscal Indicators dashboard: personal income tax, corporate income tax, VAT and excises, all % of GDP. The data are already in the repo (IMF WoRLD 2026), so nothing to download."
EX2_STEPS = [
    ("Open the repo", "VS Code > File > Open Folder > Repos\\new-ai-tools. Pull (Sync) first."),
    ("Open Copilot Chat", "Choose Agent mode and the model Claude Sonnet."),
    ("Prompt 1 - plan", "Copilot reads the README and lists what it will change. Nothing is edited yet."),
    ("Prompt 2 - edit", "It edits 3 files: the Python build, the R build and config/indicators.csv. Click Keep."),
    ("Prompt 3 - build and check", "It rebuilds and shows Indonesia's rows. Compare with the numbers below."),
    ("Prompt 4 - new version", "It ships v1.4 of the workbook and updates VersionNotes.md."),
    ("Prompt 5 - push", "It commits and pushes to GitHub (or double-click SyncRepo.bat)."),
]
EX2_CHECK = [
    ("Tax revenue", "10.08"), ("Personal income tax", "1.79"), ("Corporate income tax", "2.65"),
    ("VAT / general sales tax", "3.52"), ("Excises", "1.11"),
]

RULES = [
    "One prompt per message. Wait for the answer.",
    "Check each answer before the next prompt. You are the reviewer.",
    "No number without a source. Missing data stays missing (data_gap).",
    "Ask for a plan first (\"do not edit yet\"), then say \"go ahead\".",
    "Save outputs in the repo folder, then push, so the team has them.",
]

KIT = [
    (f"{NAME}.docx / .pptx / README.md", "This guide"),
    ("1_Excise_Indonesia_Claude/PROMPTS.md", "Copy-paste prompts A1-A6, plus the subsidy variant"),
    ("1_Excise_Indonesia_Claude/INDONESIA_INPUTS.md", "Country, currency, cut-off, the laws to look for"),
    ("1_Excise_Indonesia_Claude/process_files/", "The six .md process files (master + 5 parts)"),
    ("1_Excise_Indonesia_Claude/Excise_Diagnostic_CrossCountry.xlsx", "Workbook layout for Claude to follow"),
    ("1_Excise_Indonesia_Claude/Example_SLE_Excise_Diagnostic.docx", "What a finished report looks like (Sierra Leone)"),
    ("2_Tax_Indicators_Copilot/PROMPTS.md", "Copy-paste prompts 1-5 and the checks"),
]

# Files copied into the zip from the repo (repo path -> path inside the zip)
ZIP_FROM_REPO = {
    "ExciseDiagnostic/Excise_Diagnostic_CrossCountry.xlsx": "1_Excise_Indonesia_Claude/Excise_Diagnostic_CrossCountry.xlsx",
    "ExciseDiagnostic/country_results/SierraLeone/Example_SLE_Excise_Diagnostic.docx": "1_Excise_Indonesia_Claude/Example_SLE_Excise_Diagnostic.docx",
}
for f in ["excise_diagnostic_process.md", "excise_part1_fuel.md", "excise_part2_vehicles.md",
          "excise_part3_health.md", "excise_part4_other.md", "excise_part5_overarching.md"]:
    ZIP_FROM_REPO[f"ExciseDiagnostic/process_files/{f}"] = f"1_Excise_Indonesia_Claude/process_files/{f}"


# --------------------------------------------------------------------------- markdown
def build_md():
    L = [f"# {TITLE}", "", f"_{SUBTITLE}. {DATE}._", "",
         f"**Download everything:** [`{NAME}_Kit.zip`]({NAME}_Kit.zip) - this guide (Word, PowerPoint), "
         "the copy-paste prompts and every file the AI needs.", "",
         "## Which AI for what", "", "| Tool | Where | What it does |", "|---|---|---|"]
    L += [f"| **{a}** | {b} | {c} |" for a, b, c in TOOLS]
    L += ["", "## Before you start (once)", ""] + [f"- {x}" for x in BEFORE]
    L += ["", f"## {EX1_TITLE}", "", EX1_WHAT, ""]
    L += [f"{i}. **{a}.** {b}" for i, (a, b) in enumerate(EX1_STEPS, 1)]
    L += ["", "The six prompts (full text in [`kit/1_Excise_Indonesia_Claude/PROMPTS.md`](kit/1_Excise_Indonesia_Claude/PROMPTS.md)):", "",
          "| Prompt | What you get |", "|---|---|"] + [f"| {a} | {b} |" for a, b in EX1_PROMPTS]
    L += ["", f"> {EX1_SHORT}", ">", f"> {EX1_SUBSIDY}"]
    L += ["", f"## {EX2_TITLE}", "", EX2_WHAT, ""]
    L += [f"{i}. **{a}.** {b}" for i, (a, b) in enumerate(EX2_STEPS, 1)]
    L += ["", "Full prompts: [`kit/2_Tax_Indicators_Copilot/PROMPTS.md`](kit/2_Tax_Indicators_Copilot/PROMPTS.md). "
          "Check: Indonesia, 2024, % of GDP:", "", "| Row | Expected |", "|---|---|"]
    L += [f"| {a} | {b} |" for a, b in EX2_CHECK]
    L += ["", "## Five golden rules", ""] + [f"{i}. {x}" for i, x in enumerate(RULES, 1)]
    L += ["", "## What is in the zip", "", "| File | What it is |", "|---|---|"]
    L += [f"| `{a}` | {b} |" for a, b in KIT]
    L += ["", "Rebuild this guide, the Word and PowerPoint versions and the zip with "
          "`python user_guide/build_user_guide.py`.", ""]
    with open(os.path.join(HERE, "README.md"), "w", encoding="utf-8") as fh:
        fh.write("\n".join(L))


# --------------------------------------------------------------------------- word
def _shade(cell_or_par, hexfill):
    el = cell_or_par._tc if hasattr(cell_or_par, "_tc") else cell_or_par._p
    pr = el.get_or_add_tcPr() if hasattr(cell_or_par, "_tc") else el.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexfill); pr.append(shd)


def _table(doc, header, rows, widths):
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True
        r.font.color.rgb = RGBColor.from_string("FFFFFF"); _shade(c, NAVY)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; run = cells[i].paragraphs[0].add_run(v)
            if i == 0: run.bold = True
    for row in t.rows:
        for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    doc.add_paragraph()


def _box(doc, text, fill=LIGHT):
    p = doc.add_paragraph(); _shade(p, fill)
    p.paragraph_format.left_indent = Cm(0.2); p.add_run(text).italic = True


def build_docx(path):
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    for s in ["Title", "Heading 1", "Heading 2"]:
        doc.styles[s].font.name = "Arial"; doc.styles[s].font.color.rgb = RGBColor.from_string(NAVY)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2); sec.top_margin = sec.bottom_margin = Cm(1.8)

    doc.add_paragraph(TITLE, style="Title")
    p = doc.add_paragraph(); r = p.add_run(f"{SUBTITLE}. {DATE}."); r.italic = True
    r.font.color.rgb = RGBColor.from_string(GREY)

    doc.add_heading("Which AI for what", 1)
    _table(doc, ["Tool", "Where", "What it does"], TOOLS, [3.2, 3.6, 10.2])
    doc.add_heading("Before you start (once)", 1)
    for x in BEFORE: doc.add_paragraph(x, style="List Bullet")

    for title, what, steps in [(EX1_TITLE, EX1_WHAT, EX1_STEPS), (EX2_TITLE, EX2_WHAT, EX2_STEPS)]:
        doc.add_heading(title, 1)
        doc.add_paragraph(what)
        for i, (a, b) in enumerate(steps, 1):                 # typed numbers: Word's list style
            p = doc.add_paragraph()                           # would run on from Example 1
            p.paragraph_format.left_indent = Cm(0.7); p.paragraph_format.first_line_indent = Cm(-0.7)
            p.add_run(f"{i}.\t{a}. ").bold = True; p.add_run(b)
        if title == EX1_TITLE:
            doc.add_heading("The six prompts (full text in PROMPTS.md, kit folder 1)", 2)
            _table(doc, ["Prompt", "What you get"], EX1_PROMPTS, [2.5, 14.5])
            _box(doc, EX1_SHORT); _box(doc, EX1_SUBSIDY)
        else:
            doc.add_heading("Check: Indonesia, 2024, % of GDP (full prompts in PROMPTS.md, kit folder 2)", 2)
            _table(doc, ["Row", "Expected"], EX2_CHECK, [8, 4])

    doc.add_heading("Five golden rules", 1)
    for i, x in enumerate(RULES, 1):
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.first_line_indent = Cm(-0.7); p.add_run(f"{i}.\t").bold = True; p.add_run(x)
    doc.add_heading("What is in the kit", 1)
    _table(doc, ["File", "What it is"], KIT, [8.5, 8.5])
    doc.save(path)


# --------------------------------------------------------------------------- powerpoint
def _rgb(h): return PRGB.from_string(h)


def _text(slide, x, y, w, h, text, size=14, bold=False, color=NAVY, align=PP_ALIGN.LEFT,
          anchor=MSO_ANCHOR.TOP, font="Arial"):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align
        r = p.add_run(); r.text = ln; r.font.size = PPt(size); r.font.bold = bold
        r.font.color.rgb = _rgb(color); r.font.name = font
    return tb


def _rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = _rgb(fill)
    if line: s.line.color.rgb = _rgb(line); s.line.width = PPt(1)
    else: s.line.fill.background()
    s.shadow.inherit = False
    return s


def _header(prs, kicker, title):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(s, 0, 0, 13.333, 0.12, NAVY)
    _text(s, 0.6, 0.35, 12, 0.35, kicker.upper(), 11, True, GOLD)
    _text(s, 0.6, 0.65, 12.2, 0.8, title, 26, True, NAVY)
    _text(s, 0.6, 7.05, 12.2, 0.3, f"{TITLE}  ·  {DATE}  ·  Screens and model names change over time", 9, False, GREY)
    return s


def _steps(s, steps, x, y, w, row_h):
    for i, (a, b) in enumerate(steps):
        yy = y + i * row_h
        c = _rect(s, x, yy + 0.05, 0.42, 0.42, BLUE, shape=MSO_SHAPE.OVAL)
        tf = c.text_frame; tf.margin_left = tf.margin_right = 0
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = str(i + 1)
        r.font.size = PPt(13); r.font.bold = True; r.font.color.rgb = _rgb("FFFFFF"); r.font.name = "Arial"
        _text(s, x + 0.55, yy, w - 0.55, 0.3, a, 13, True, NAVY)
        _text(s, x + 0.55, yy + 0.27, w - 0.55, row_h - 0.27, b, 10, False, GREY)


def _panel(s, x, y, w, h, heading, rows, size=11):
    _rect(s, x, y, w, h, LIGHT)
    _rect(s, x, y, w, 0.08, GOLD)
    _text(s, x + 0.2, y + 0.18, w - 0.4, 0.35, heading, 13, True, NAVY)
    yy = y + 0.6
    for a, b in rows:
        _text(s, x + 0.2, yy, 1.4 if len(a) < 6 else w * 0.55, 0.3, a, size, True, BLUE)
        _text(s, x + (1.2 if len(a) < 6 else w * 0.6), yy, w - (1.4 if len(a) < 6 else w * 0.62), 0.45, b, size, False, GREY)
        yy += 0.42 if len(a) < 6 else 0.36


def build_pptx(path):
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)

    # 1 title
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(s, 0, 0, 13.333, 7.5, NAVY); _rect(s, 0.8, 2.2, 1.2, 0.08, GOLD)
    _text(s, 0.8, 0.9, 11, 0.4, "USER GUIDE  ·  START HERE", 13, True, GOLD)
    _text(s, 0.8, 2.45, 11.5, 1.0, TITLE, 44, True, "FFFFFF")
    _text(s, 0.8, 3.55, 11, 1.2, SUBTITLE, 18, False, "D9E2F3")
    _text(s, 0.8, 6.3, 11, 0.4, f"{DATE}  ·  Everything you need is in {NAME}_Kit.zip", 12, False, "D9E2F3")

    # 2 which AI for what + before you start
    s = _header(prs, "The tools", "Which AI for what")
    for i, (a, b, c) in enumerate(TOOLS):
        x = 0.6 + i * 4.1
        _rect(s, x, 1.7, 3.85, 2.9, LIGHT); _rect(s, x, 1.7, 3.85, 0.08, [BLUE, NAVY, GOLD][i])
        _text(s, x + 0.25, 1.95, 3.4, 0.45, a, 20, True, NAVY)
        _text(s, x + 0.25, 2.45, 3.4, 0.3, b, 11, True, BLUE)
        _text(s, x + 0.25, 2.85, 3.4, 1.7, c, 12, False, GREY)
    _text(s, 0.6, 4.95, 12, 0.4, "Before you start (once)", 15, True, NAVY)
    _text(s, 0.6, 5.35, 12.1, 1.6, [f"✓  {x}" for x in BEFORE], 12, False, GREY)

    # 3 example 1 steps
    s = _header(prs, "Example 1 · Claude · deep research", "Excise diagnostic for Indonesia in seven steps")
    _text(s, 0.6, 1.45, 12, 0.6, EX1_WHAT, 12, False, GREY)
    _steps(s, EX1_STEPS, 0.6, 2.15, 7.4, 0.68)
    _panel(s, 8.4, 2.15, 4.35, 3.4, "The six prompts", EX1_PROMPTS, 11)
    _rect(s, 8.4, 5.75, 4.35, 1.15, "FFF6DB")
    _text(s, 8.55, 5.8, 4.1, 1.1, [EX1_SHORT, "Subsidy diagnostic: see the variant prompt in PROMPTS.md (first draft; no subsidy process file yet)."], 9.5, False, GREY)

    # 4 example 1 prompt A1 in full
    s = _header(prs, "Example 1 · the first prompt", "What you paste: prompt A1 (the rest are in PROMPTS.md)")
    a1 = open(os.path.join(HERE, "kit", "1_Excise_Indonesia_Claude", "PROMPTS.md"), encoding="utf-8").read()
    a1 = a1.split("## Prompt A1")[1].split("```text\n")[1].split("```")[0].strip()
    _rect(s, 0.6, 1.6, 7.6, 5.2, "F2F2F2", line="BFBFBF")
    _text(s, 0.8, 1.75, 7.25, 5.0, " ".join(a1.split()), 12, False, "262626", font="Consolas")
    _panel(s, 8.5, 1.6, 4.25, 5.2, "What to check in the answer", [
        ("✓", "Every rate quoted as printed, with its article and date"),
        ("✓", "A working link for every source"),
        ("✓", "Pending bills listed, not applied"),
        ("✓", "Unreadable tables marked \"not legible\""),
        ("✓", "Gaps marked data_gap, never guessed"),
        ("→", "Then paste A2. If not right, tell Claude what to fix first."),
    ], 11)

    # 5 example 2 steps
    s = _header(prs, "Example 2 · GitHub Copilot · edits the repo", "Add four tax indicators to the dashboard in seven steps")
    _text(s, 0.6, 1.45, 12, 0.6, EX2_WHAT, 12, False, GREY)
    _steps(s, EX2_STEPS, 0.6, 2.15, 7.4, 0.68)
    _panel(s, 8.4, 2.15, 4.35, 3.0, "Check: Indonesia 2024, % of GDP", EX2_CHECK, 12)
    _rect(s, 8.4, 5.35, 4.35, 1.55, "FFF6DB")
    _text(s, 8.55, 5.42, 4.1, 1.45, ["The four parts add to less than total tax (9.07 vs 10.08): trade, property and other taxes are the rest.",
                                    "git diff should show only the files Copilot listed in its plan."], 10, False, GREY)

    # 6 example 2 prompts 1-2
    s = _header(prs, "Example 2 · the first prompt", "What you paste: prompt 1 asks for a plan, prompt 2 says go ahead")
    p2 = open(os.path.join(HERE, "kit", "2_Tax_Indicators_Copilot", "PROMPTS.md"), encoding="utf-8").read()
    p1 = p2.split("## Prompt 1")[1].split("```text\n")[1].split("```")[0].strip()
    _rect(s, 0.6, 1.6, 7.6, 5.2, "F2F2F2", line="BFBFBF")
    head, rest = p1.split("\npit_", 1)                    # intro | four mapping lines + last line
    lines = ("pit_" + rest).split("\n")
    _text(s, 0.8, 1.75, 7.25, 5.0, [" ".join(head.split())] + lines[:4] + [" ".join(" ".join(lines[4:]).split())],
          12, False, "262626", font="Consolas")
    _panel(s, 8.5, 1.6, 4.25, 5.2, "Then, one per message", [
        ("2", "Go ahead: edit Python, R and config"),
        ("3", "Build and show Indonesia's rows"),
        ("4", "Ship v1.4, update VersionNotes"),
        ("5", "Commit and push"),
        ("!", "If anything else changes, say \"undo that\""),
    ], 11)

    # 7 golden rules + kit
    s = _header(prs, "Remember", "Five golden rules, and what is in the kit")
    for i, x in enumerate(RULES):
        yy = 1.7 + i * 0.95
        c = _rect(s, 0.6, yy, 0.55, 0.55, GOLD, shape=MSO_SHAPE.OVAL)
        p = c.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = str(i + 1)
        r.font.size = PPt(16); r.font.bold = True; r.font.color.rgb = _rgb("FFFFFF"); r.font.name = "Arial"
        _text(s, 1.35, yy + 0.08, 4.9, 0.8, x, 13, False, NAVY)
    _rect(s, 6.6, 1.6, 6.15, 5.2, LIGHT); _rect(s, 6.6, 1.6, 6.15, 0.08, BLUE)
    _text(s, 6.8, 1.8, 5.8, 0.4, f"{NAME}_Kit.zip", 14, True, NAVY)
    yy = 2.3
    for a, b in KIT:
        _text(s, 6.8, yy, 5.8, 0.3, a, 9.5, True, BLUE, font="Consolas")
        _text(s, 6.8, yy + 0.25, 5.8, 0.3, b, 10, False, GREY); yy += 0.62
    prs.save(path)


# --------------------------------------------------------------------------- zip
def build_zip(path, docx, pptx):
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        root = NAME + "_Kit/"
        for f in [docx, pptx]:
            z.write(f, root + os.path.basename(f))
        md = open(os.path.join(HERE, "README.md"), encoding="utf-8").read()
        z.writestr(root + "README.md", md.replace("kit/", ""))   # kit folders sit at the zip root
        for d, _, files in os.walk(os.path.join(HERE, "kit")):
            for f in files:
                full = os.path.join(d, f); z.write(full, root + os.path.relpath(full, os.path.join(HERE, "kit")))
        for src, dst in ZIP_FROM_REPO.items():
            z.write(os.path.join(REPO, src), root + dst)


if __name__ == "__main__":
    build_md()
    docx = os.path.join(HERE, NAME + ".docx"); pptx = os.path.join(HERE, NAME + ".pptx")
    build_docx(docx); build_pptx(pptx)
    zp = os.path.join(HERE, NAME + "_Kit.zip"); build_zip(zp, docx, pptx)
    print("wrote", docx, pptx, zp, sep="\n  ")
